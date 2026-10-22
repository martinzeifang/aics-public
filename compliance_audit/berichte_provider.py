"""Compliance-Audit-Adapter für das geteilte Berichts-Center (#1726).

Liefert Katalog, Render-Callable und den Kontext für die KI-Zusammenfassung — nach
demselben Muster wie CRA/NIS2/DSGVO/WiBA. Die Typ-IDs sind modulübergreifend
identisch (#1502/#1504): ``gesamt``, ``anforderungen``, ``massnahmen``. Nur die
Labels sind fachlich: Im Audit heißen „Anforderungen" **Fragen**.

Grundlage jedes Berichts ist der **berechnete** Compliance-Stand aus
:mod:`compliance_audit.status` — derselbe, den die Wasserstandsmeldung ausweist.
So können Bericht und Meldung nicht auseinanderlaufen, und die Stufe im Bericht
ist ebenso reproduzierbar wie dort.

Liegt ein KI-Entwurf vor (``gutachten_drafts``, erzeugt über
``/bericht/{prompt,import}``), übernimmt der Gesamtbericht dessen Formulierungen
(Zusammenfassung, Einordnung je Framework, Maßnahmen). Ohne Entwurf entsteht der
Bericht rein aus den Daten — er ist also nie von einem KI-Lauf abhängig.
"""
from __future__ import annotations

import json
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

from shared.reports.core import ReportSpec

DB_PATH = Path('data/db/gutachten.sqlite')

CATALOG: list[ReportSpec] = [
    ReportSpec('gesamt', 'Auditbericht (Gesamt)', 'Framework-abhängig',
               'Vollständiger Auditbericht: Compliance-Stand, Einordnung je Framework, '
               'Feststellungen und Maßnahmen. Nutzt den KI-Entwurf, wenn vorhanden.'),
    ReportSpec('anforderungen', 'Fragen-Status', 'Framework-abhängig',
               'Alle Prüffragen je Framework mit Antwort und Erfüllung — ohne '
               'Maßnahmenteil.'),
    ReportSpec('massnahmen', 'Maßnahmenplan', 'Framework-abhängig',
               'Lücken (nicht erfüllt, teilweise, offen) als Maßnahmenplan, '
               'nach Dringlichkeit sortiert.'),
]

#: Anzeigetexte der Erfüllung im Bericht — Quelle ist das Status-Modul (#1730),
#: damit Bericht und Word-Vorlage nicht auseinanderlaufen.
from compliance_audit.status import ERFUELLUNG_LABELS as _ERF_LABEL  # noqa: E402


# ── Farbgebung (#1727) ───────────────────────────────────────────────────────
# python-docx kann Zellen-/Absatzschattierung nicht ueber die hohe Schnittstelle,
# deshalb direkt als w:shd-Element. Die Palette ist dieselbe wie in der Oberflaeche,
# damit Bericht und Bildschirm zusammenpassen.

FARBEN = {
    'kopf':        '1565C0',   # Blau — Tabellenkopf, Ueberschriften-Akzent
    'vollstaendig':'2E7D32',   # Gruen
    'weitgehend':  'F9A825',   # Gelb
    'teilweise':   'EF6C00',   # Orange
    'nicht':       'C62828',   # Rot
    'unbewertet':  '90A4AE',   # Grau
}

#: Zellenhintergrund je Erfuellung — hell, damit der Text lesbar bleibt.
ERF_FARBE = {
    'erfuellt':       'E8F5E9',
    'teilweise':      'FFF3E0',
    'nicht_erfuellt': 'FFEBEE',
    'nicht_anwendbar':'F5F5F5',
    'offen':          'E3F2FD',
}

PRIO_FARBE = {'hoch': 'FFEBEE', 'mittel': 'FFF8E1', 'niedrig': 'E8F5E9'}


def _schattieren(zelle, hex_farbe: str) -> None:
    """Zellenhintergrund setzen (w:shd — kein High-Level-API in python-docx)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_farbe)
    zelle._tc.get_or_add_tcPr().append(shd)


def _weiss_fett(zelle, text: str) -> None:
    """Kopfzellen-Text: weiss auf Blau."""
    from docx.shared import RGBColor

    run = zelle.paragraphs[0].add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _tabelle(doc, kopf: list[str], zeilen: int):
    """Tabelle mit farbigem Kopf anlegen."""
    tbl = doc.add_table(rows=zeilen + 1, cols=len(kopf))
    tbl.style = 'Table Grid'
    for ci, titel in enumerate(kopf):
        z = tbl.cell(0, ci)
        _schattieren(z, FARBEN['kopf'])
        _weiss_fett(z, titel)
    return tbl


def _ueberschrift(doc, text: str, level: int = 1):
    """Ueberschrift mit farbigem Akzent — Word-Blau wirkt sonst beliebig."""
    from docx.shared import RGBColor

    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)
    return h


def _stand_banner(doc, stand: dict[str, Any]) -> None:
    """Farbiges Band mit Stufe und Quote — das Erste, was man sehen soll."""
    from docx.shared import Pt, RGBColor

    stufe = str(stand.get('stufe') or 'unbewertet')
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    z = tbl.cell(0, 0)
    _schattieren(z, FARBEN.get(stufe, FARBEN['unbewertet']))
    para = z.paragraphs[0]
    run = para.add_run(f"{stand.get('label', '')}")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    zweite = z.add_paragraph()
    r2 = zweite.add_run(
        f"Erfüllungsquote {stand.get('quote_prozent')} % · "
        f"Abdeckung {stand.get('abdeckung_prozent')} %")
    r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r2.font.size = Pt(10)
    doc.add_paragraph()


def _balken(prozent: Any, breite: int = 20) -> str:
    """Einfacher Textbalken — funktioniert in DOCX und PDF ohne Grafik."""
    try:
        anteil = max(0.0, min(100.0, float(prozent))) / 100.0
    except (TypeError, ValueError):
        return ''
    voll = int(round(anteil * breite))
    return '█' * voll + '░' * (breite - voll)


def catalog() -> list[ReportSpec]:
    return CATALOG


# ── Datenbeschaffung ─────────────────────────────────────────────────────────

def _sammeln(projekt: str) -> dict[str, Any]:
    """Projekt, Fragen, berechneter Stand und KI-Entwurf in einem Zug."""
    from gutachten.db import load_project, load_questions, load_gutachten_draft

    from compliance_audit.status import build_status, erfuellung_of

    proj = load_project(DB_PATH, projekt) or {}
    fragen = load_questions(DB_PATH, projekt)
    stand = build_status(fragen)

    entwurf: dict[str, Any] = {}
    roh = load_gutachten_draft(DB_PATH, projekt)
    if roh:
        try:
            geladen = json.loads(roh)
            if isinstance(geladen, dict):
                entwurf = geladen
        except (ValueError, TypeError):
            entwurf = {}   # kaputter Entwurf darf den Bericht nicht verhindern

    # Fragen je Framework, Reihenfolge wie im Projekt
    nach_fw: dict[str, list[dict[str, Any]]] = {}
    for q in fragen:
        nach_fw.setdefault(str(q.get('framework') or '(ohne Framework)'), []).append(q)
    for liste in nach_fw.values():
        liste.sort(key=lambda q: (str(q.get('section_ref') or ''),
                                  q.get('question_num') or 0))

    return {'projekt': proj, 'fragen': fragen, 'stand': stand, 'entwurf': entwurf,
            'nach_framework': nach_fw, 'erfuellung_of': erfuellung_of}


def _dokument(titel: str, proj: dict[str, Any], stand: dict[str, Any]):
    """Neues DOCX mit Kopf, A4 und Basis-Schrift (#1617)."""
    from docx import Document

    from shared.reports.docx_pdf import apply_base_font, set_a4

    doc = Document()
    set_a4(doc)
    apply_base_font(doc)

    from docx.shared import RGBColor

    h = doc.add_heading(f"{titel} — {proj.get('name', '')}", level=0)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)
    doc.add_paragraph(f"Erstellt: {date.today().strftime('%d.%m.%Y')}")

    # Stand als farbiges Band direkt unter dem Titel.
    _stand_banner(doc, stand)

    meta = _tabelle(doc, ['Angabe', 'Wert'], 3)
    for ri, (k, v) in enumerate((
        ('Unternehmen', proj.get('unternehmen') or '—'),
        ('Frameworks', ', '.join(proj.get('frameworks') or []) or '—'),
        ('Prüfungsfokus', str(proj.get('pruefungsfokus') or '—')[:1500]),
    ), start=1):
        meta.cell(ri, 0).paragraphs[0].add_run(k).bold = True
        meta.cell(ri, 1).text = str(v)
    doc.add_paragraph()
    return doc


def _stand_tabelle(doc, stand: dict[str, Any]) -> None:
    """Kennzahlen + Stufe je Framework."""
    f = stand.get('fragen') or {}
    _ueberschrift(doc, 'Compliance-Stand', level=1)

    # Verteilung als farbige Tabelle — auf einen Blick erkennbar.
    vert = _tabelle(doc, ['Erfüllung', 'Anzahl'], 5)
    for ri, (schluessel, label) in enumerate((
        ('erfuellt', 'erfüllt'), ('teilweise', 'teilweise erfüllt'),
        ('nicht_erfuellt', 'nicht erfüllt'), ('offen', 'offen'),
        ('nicht_anwendbar', 'nicht anwendbar'),
    ), start=1):
        z = vert.cell(ri, 0)
        _schattieren(z, ERF_FARBE[schluessel])
        z.text = label
        vert.cell(ri, 1).text = str(f.get(schluessel, 0))
    doc.add_paragraph(
        f"{f.get('gesamt', 0)} Fragen gesamt. Offene Fragen zählen in der Quote als "
        'nicht nachgewiesen; nur die bewerteten Fragen ergäben '
        f"{stand.get('quote_bewertet_prozent')} %.")

    fws = stand.get('frameworks') or []
    if not fws:
        return
    tbl = _tabelle(doc, ['Framework', 'Stufe', 'Quote', 'Verlauf', 'offen / anwendbar'],
                   len(fws))
    for ri, fw in enumerate(fws, start=1):
        ff = fw.get('fragen') or {}
        tbl.cell(ri, 0).text = str(fw.get('framework', ''))
        z = tbl.cell(ri, 1)
        _schattieren(z, ERF_FARBE.get(
            {'vollstaendig': 'erfuellt', 'weitgehend': 'teilweise',
             'teilweise': 'offen', 'nicht': 'nicht_erfuellt'}.get(
                str(fw.get('stufe')), 'nicht_anwendbar'), 'F5F5F5'))
        z.text = str(fw.get('label', ''))
        tbl.cell(ri, 2).text = f"{fw.get('quote_prozent')} %"
        tbl.cell(ri, 3).text = _balken(fw.get('quote_prozent'))
        tbl.cell(ri, 4).text = f"{ff.get('offen', 0)} / {ff.get('anwendbar', 0)}"


# ── Berichtstypen ────────────────────────────────────────────────────────────

def _fragen_docx(projekt: str) -> bytes:
    d = _sammeln(projekt)
    doc = _dokument('Fragen-Status', d['projekt'], d['stand'])
    _stand_tabelle(doc, d['stand'])

    for fw, fragen in d['nach_framework'].items():
        _ueberschrift(doc, fw, level=1)
        tbl = _tabelle(doc, ['Fundstelle', 'Frage', 'Antwort', 'Erfüllung'], len(fragen))
        for ri, q in enumerate(fragen, start=1):
            tbl.cell(ri, 0).text = str(q.get('section_ref') or '')[:60]
            tbl.cell(ri, 1).text = str(q.get('frage') or '')[:500]
            tbl.cell(ri, 2).text = str(q.get('antwort') or '')[:800] or '—'
            z = tbl.cell(ri, 3)
            _schattieren(z, ERF_FARBE.get(d['erfuellung_of'](q), 'FFFFFF'))
            z.text = _ERF_LABEL.get(d['erfuellung_of'](q), 'offen')

    from shared.reports.docx_pdf import fit_tables
    fit_tables(doc)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _massnahmen_docx(projekt: str) -> bytes:
    d = _sammeln(projekt)
    doc = _dokument('Maßnahmenplan', d['projekt'], d['stand'])

    # Zuerst die KI-Maßnahmen aus dem Entwurf, falls vorhanden — sie sind
    # priorisiert und begruendet.
    empf = d['entwurf'].get('handlungsempfehlungen') or []
    if empf:
        _ueberschrift(doc, 'Empfohlene Maßnahmen', level=1)
        tbl = _tabelle(doc, ['Maßnahme', 'Priorität', 'Aufwand', 'Begründung'], len(empf))
        for ri, m in enumerate(empf, start=1):
            if not isinstance(m, dict):
                m = {'massnahme': str(m)}
            tbl.cell(ri, 0).text = str(m.get('massnahme', ''))[:400]
            prio = str(m.get('prioritaet', '') or '').lower()
            z = tbl.cell(ri, 1)
            if prio in PRIO_FARBE:
                _schattieren(z, PRIO_FARBE[prio])
            z.text = prio or '—'
            tbl.cell(ri, 2).text = str(m.get('aufwand', '') or '—')
            tbl.cell(ri, 3).text = str(m.get('begruendung', ''))[:600]

    luecken = d['stand'].get('luecken') or []
    _ueberschrift(doc, 'Offene Punkte aus dem Fragenkatalog', level=1)
    if not luecken:
        doc.add_paragraph('Keine Lücken — alle anwendbaren Fragen sind erfüllt.')
    else:
        doc.add_paragraph(
            f"{len(luecken)} Punkte, sortiert nach Dringlichkeit "
            '(nicht erfüllt, dann teilweise, dann offen).')
        tbl = _tabelle(doc, ['Framework', 'Fundstelle', 'Frage', 'Status'], len(luecken))
        for ri, lk in enumerate(luecken, start=1):
            tbl.cell(ri, 0).text = str(lk.get('framework', ''))
            tbl.cell(ri, 1).text = str(lk.get('section_ref', ''))[:60]
            tbl.cell(ri, 2).text = str(lk.get('frage', ''))[:500]
            z = tbl.cell(ri, 3)
            _schattieren(z, ERF_FARBE.get(lk.get('erfuellung'), 'FFFFFF'))
            z.text = _ERF_LABEL.get(lk.get('erfuellung'), '')

    from shared.reports.docx_pdf import fit_tables
    fit_tables(doc)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _gesamt_docx(projekt: str) -> bytes:
    d = _sammeln(projekt)
    e = d['entwurf']
    doc = _dokument('Auditbericht', d['projekt'], d['stand'])

    if e.get('gesamtbewertung'):
        _ueberschrift(doc, 'Gesamtbewertung', level=1)
        doc.add_paragraph(str(e['gesamtbewertung']))
    if e.get('zusammenfassung'):
        _ueberschrift(doc, 'Zusammenfassung', level=1)
        doc.add_paragraph(str(e['zusammenfassung']))
        doc.add_paragraph('🤖 KI-generiert — fachlich zu prüfen.').italic = True
    elif not e:
        _ueberschrift(doc, 'Zusammenfassung', level=1)
        doc.add_paragraph(
            'Für dieses Audit liegt noch kein ausformulierter Entwurf vor. Der Bericht '
            'enthält den berechneten Stand und die Datenlage; die Einordnung lässt sich '
            'im Bericht-Tab erzeugen („Auditbericht erzeugen").')

    _stand_tabelle(doc, d['stand'])

    for fw in e.get('framework_bewertungen') or []:
        if not isinstance(fw, dict):
            continue
        _ueberschrift(doc, str(fw.get('framework', '')), level=1)
        if fw.get('bewertung'):
            doc.add_paragraph(str(fw['bewertung']))
        for titel, schluessel in (('Stärken', 'staerken'), ('Lücken', 'luecken')):
            werte = fw.get(schluessel) or []
            if werte:
                _ueberschrift(doc, titel, level=2)
                for w in werte:
                    doc.add_paragraph(str(w), style='List Bullet')

    empf = e.get('handlungsempfehlungen') or []
    if empf:
        _ueberschrift(doc, 'Handlungsempfehlungen', level=1)
        tbl = _tabelle(doc, ['Maßnahme', 'Priorität', 'Aufwand', 'Begründung'], len(empf))
        for ri, m in enumerate(empf, start=1):
            if not isinstance(m, dict):
                m = {'massnahme': str(m)}
            tbl.cell(ri, 0).text = str(m.get('massnahme', ''))[:400]
            prio = str(m.get('prioritaet', '') or '').lower()
            z = tbl.cell(ri, 1)
            if prio in PRIO_FARBE:
                _schattieren(z, PRIO_FARBE[prio])
            z.text = prio or '—'
            tbl.cell(ri, 2).text = str(m.get('aufwand', '') or '—')
            tbl.cell(ri, 3).text = str(m.get('begruendung', ''))[:600]

    # Datenlage immer anhaengen — der Bericht soll ohne Rueckfragen pruefbar sein.
    _ueberschrift(doc, 'Anhang: Fragen und Antworten', level=1)
    for fw, fragen in d['nach_framework'].items():
        _ueberschrift(doc, fw, level=2)
        tbl = _tabelle(doc, ['Fundstelle', 'Frage', 'Antwort', 'Erfüllung'], len(fragen))
        for ri, q in enumerate(fragen, start=1):
            tbl.cell(ri, 0).text = str(q.get('section_ref') or '')[:60]
            tbl.cell(ri, 1).text = str(q.get('frage') or '')[:500]
            tbl.cell(ri, 2).text = str(q.get('antwort') or '')[:800] or '—'
            z = tbl.cell(ri, 3)
            _schattieren(z, ERF_FARBE.get(d['erfuellung_of'](q), 'FFFFFF'))
            z.text = _ERF_LABEL.get(d['erfuellung_of'](q), 'offen')

    if e.get('fazit'):
        _ueberschrift(doc, 'Fazit', level=1)
        doc.add_paragraph(str(e['fazit']))

    from shared.reports.docx_pdf import fit_tables
    fit_tables(doc)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


_BAUER = {
    'gesamt': _gesamt_docx,
    'anforderungen': _fragen_docx,
    'massnahmen': _massnahmen_docx,
}


def render(typ: str, fmt: str, *, projekt: str, von: str | None = None,
           bis: str | None = None, **_ctx: Any) -> bytes:
    """Bericht als DOCX- oder PDF-Bytes (kein Zeitraum — ein Audit ist ein Stand)."""
    bauer = _BAUER.get(typ)
    if bauer is None:
        raise ValueError(f'Unbekannter Berichtstyp: {typ}')
    docx_bytes = bauer(projekt)
    if fmt == 'pdf':
        from shared.reports.docx_pdf import docx_bytes_to_pdf
        return docx_bytes_to_pdf(docx_bytes)
    return docx_bytes


def summary_context(projekt: str) -> dict[str, Any]:
    """Kontext für die geteilte KI-Zusammenfassung des Berichts-Centers.

    Bewusst knapp und ohne Freitext-Antworten: Die Zusammenfassung soll den Stand
    einordnen, nicht die Rohdaten wiederholen.
    """
    d = _sammeln(projekt)
    stand = d['stand']
    luecken = stand.get('luecken') or []
    return {
        # 'reifegrad'/'offene' sind der Vertrag von shared.reports.summary — nur so
        # landen die Zahlen im Prompt (#1725).
        'reifegrad': {
            'stufe': stand.get('label'),
            'erfuellungsquote_prozent': stand.get('quote_prozent'),
            'abdeckung_prozent': stand.get('abdeckung_prozent'),
            # 'fragen' ist die Zählung je Erfüllungsgrad (gesamt/bewertet/erfuellt/…).
            'fragen': stand.get('fragen'),
            'je_framework': [
                {'framework': fw.get('framework'), 'stufe': fw.get('label'),
                 'erfuellungsquote_prozent': fw.get('quote_prozent'),
                 'fragen': fw.get('fragen')}
                for fw in (stand.get('frameworks') or [])
            ],
        },
        'offene': [
            # Ohne 'bewertung': Das Audit bewertet nicht 0–5, sondern Erfüllungsgrade.
            {'id': ' '.join(x for x in (lk.get('framework'), lk.get('section_ref')) if x),
             'titel': lk.get('frage'),
             'status': _ERF_LABEL.get(str(lk.get('erfuellung')), str(lk.get('erfuellung')))}
            for lk in luecken[:40]
        ],
        'unternehmen': d['projekt'].get('unternehmen', ''),
        'frameworks': d['projekt'].get('frameworks') or [],
        'pruefungsfokus': d['projekt'].get('pruefungsfokus', ''),
        'luecken_anzahl': len(luecken),
    }
