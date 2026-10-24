"""Word-Vorlagen-Adapter des Compliance-Audits (#1730).

Der Adapter speist Kundenvorlagen (`{{ jinja }}`) über die zentrale
Vorlagen-Engine (#988). Zwei Eigenschaften entscheiden darüber, ob das in der
Praxis trägt, und beide sind hier festgehalten:

1. **Die Variablenliste muss stimmen.** Sie ist die einzige Quelle, aus der ein
   Anwender die Platzhalter abschreibt (#1092). Ein dokumentierter Schlüssel, den
   der Kontext nicht liefert, führt zu einem leeren Feld im Kundendokument — ohne
   Fehlermeldung. Der Test löst deshalb **jeden** dokumentierten Schlüssel im
   echten Kontext auf.
2. **Kein ``None``.** Jinja druckt ``None`` wörtlich ins Dokument. Der Kontext
   muss auch für ein unbekanntes Projekt vollständig und druckbar sein.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from compliance_audit import status, template_context as tc


# ── Fixture ──────────────────────────────────────────────────────────────────

@pytest.fixture
def audit_db():
    """Audit-Projekt mit gemischten Erfüllungsgraden (erzeugt echte Lücken)."""
    from compliance_audit import db as adb
    from gutachten.db import save_project, save_questions

    p = Path('data/db/pytest_audit_1730.sqlite')
    p.parent.mkdir(parents=True, exist_ok=True)
    adb.ensure_db(p)
    con = adb._connect(p)
    try:
        for t in ('audit_wasserstand', 'gutachten_questions', 'gutachten_projects',
                  'gutachten_drafts'):
            con.execute(f'DELETE FROM {t}')  # nosec B608 — Konstanten
        con.commit()
    finally:
        con.close()

    save_project(p, 'VorlageAudit', frameworks=['NIS2', 'CRA'],
                 pruefungsfokus='Lieferkette', unternehmen='Acme GmbH')
    # Q2 traegt die Erfuellung bewusst nur im Freitext ``bewertung`` — so kommen
    # importierte Altbestaende an (``save_questions`` schreibt die Spalte
    # ``erfuellung`` nicht; dafuer gibt es ``set_erfuellung`` und den Fallback
    # ``normalize_erfuellung``). Der Adapter muss beide Wege lesen.
    save_questions(p, 'VorlageAudit', [
        {'framework': 'NIS2', 'section_ref': 'Art. 21(2)a',
         'thema': 'Risikomanagement', 'frage': 'Gibt es ein ISMS?',
         'antwort': 'Ja, zertifiziert.', 'kommentar': 'Zertifikat lag vor.'},
        {'framework': 'NIS2', 'section_ref': 'Art. 21(2)d',
         'thema': 'Lieferkette', 'frage': 'Lieferantenverzeichnis?',
         'antwort': 'Teilweise.', 'bewertung': 'teilweise erfüllt'},
        {'framework': 'CRA', 'section_ref': 'Art. 13(1)',
         'thema': 'Schwachstellen', 'frage': 'Schwachstellenprozess?',
         'antwort': 'Nein.', 'kommentar': 'Offen.'},
        {'framework': 'CRA', 'section_ref': 'Art. 13(2)',
         'thema': 'SBOM', 'frage': 'Liegt eine SBOM vor?', 'antwort': ''},
    ])

    from gutachten.db import load_questions
    nach_nr = {q['question_num']: q['id'] for q in load_questions(p, 'VorlageAudit')}
    assert adb.set_erfuellung(p, nach_nr[1], 'erfuellt')
    assert adb.set_erfuellung(p, nach_nr[3], 'nicht_erfuellt')
    return p


def _pfad(ctx: dict, key: str):
    """Punkt-Pfad im Kontext auflösen; KeyError, wenn er nicht existiert."""
    cur = ctx
    for teil in key.split('.'):
        if not isinstance(cur, dict) or teil not in cur:
            raise KeyError(key)
        cur = cur[teil]
    return cur


def _none_pfade(obj, pfad='') -> list[str]:
    """Alle Stellen mit ``None`` — für eine aussagekräftige Fehlermeldung."""
    treffer = []
    if obj is None:
        treffer.append(pfad or '<wurzel>')
    elif isinstance(obj, dict):
        for k, v in obj.items():
            treffer += _none_pfade(v, f'{pfad}.{k}' if pfad else str(k))
    elif isinstance(obj, (list, tuple)):
        for i, v in enumerate(obj):
            treffer += _none_pfade(v, f'{pfad}[{i}]')
    return treffer


# ── Variablenliste ───────────────────────────────────────────────────────────

def test_jede_dokumentierte_variable_existiert_im_kontext(audit_db):
    """Der Kern: Was die Liste verspricht, muss der Kontext liefern.

    Sonst schreibt ein Anwender den Platzhalter ab und erhält ein stillschweigend
    leeres Feld im Kundendokument.
    """
    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    fehlend = []
    for v in tc.COMPLIANCE_AUDIT_VARIABLES:
        try:
            _pfad(ctx, v['key'])
        except KeyError:
            fehlend.append(v['key'])
    assert not fehlend, f'dokumentiert, aber nicht im Kontext: {fehlend}'


def test_variablenliste_ist_vollstaendig_beschrieben():
    """Die Liste ist Anwender-Dokumentation, nicht nur Technik."""
    for v in tc.COMPLIANCE_AUDIT_VARIABLES:
        assert set(v) >= {'key', 'typ', 'beschreibung', 'pflicht'}, v
        assert v['typ'] in ('text', 'zahl', 'liste', 'objekt'), v
        assert len(v['beschreibung']) >= 15, f'Beschreibung zu dünn: {v["key"]}'
        assert isinstance(v['pflicht'], bool)


def test_variablen_sind_eindeutig():
    keys = [v['key'] for v in tc.COMPLIANCE_AUDIT_VARIABLES]
    assert len(keys) == len(set(keys)), 'doppelte Variablen-Keys'


def test_engine_kennt_den_adapter():
    """Ohne Registrierung in der Engine antwortet der Render-Weg mit 501."""
    from shared.templates import schema
    assert 'compliance_audit' in schema.MODULE
    assert schema.get_context_builder('compliance_audit') is not None
    assert len(schema.get_variables('compliance_audit')) == \
        len(tc.COMPLIANCE_AUDIT_VARIABLES)


# ── Jinja-Robustheit ─────────────────────────────────────────────────────────

def test_kontext_enthaelt_kein_none(audit_db):
    """Jinja würde ``None`` wörtlich ins Kundendokument drucken."""
    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    assert not _none_pfade(ctx), f'None im Kontext: {_none_pfade(ctx)}'


def test_unbekanntes_projekt_liefert_druckbaren_kontext(audit_db):
    """Ein fehlendes Projekt darf keinen Export abbrechen."""
    ctx = tc.build_compliance_audit_context(audit_db, 'Gibt-Es-Nicht')
    assert not _none_pfade(ctx)
    assert ctx['projekt']['name'] == 'Gibt-Es-Nicht', 'Name muss erhalten bleiben'
    assert ctx['fragen'] == [] and ctx['luecken'] == []
    assert ctx['meta']['fragen']['gesamt'] == 0
    assert ctx['bericht']['zusammenfassung'] == ''


def test_kaputter_ki_entwurf_bricht_nicht(audit_db):
    """Ein unvollständiger KI-Entwurf darf den Export nicht verhindern."""
    from gutachten.db import save_gutachten_draft

    save_gutachten_draft(audit_db, 'VorlageAudit', '{das ist kein JSON')
    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    assert ctx['bericht']['gesamtbewertung'] == ''
    assert ctx['bericht']['handlungsempfehlungen'] == []
    assert not _none_pfade(ctx)


def test_ki_entwurf_wird_normalisiert(audit_db):
    """Fehlende Felder im Entwurf werden ergänzt, nicht durchgereicht."""
    from gutachten.db import save_gutachten_draft

    save_gutachten_draft(audit_db, 'VorlageAudit', json.dumps({
        'gesamtbewertung': 'Solide Basis.',
        'framework_bewertungen': [{'framework': 'NIS2', 'staerken': ['ISMS']}],
        'handlungsempfehlungen': [{'massnahme': 'Prozess einführen'}],
    }))
    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    b = ctx['bericht']
    assert b['gesamtbewertung'] == 'Solide Basis.'
    assert b['fazit'] == '', 'fehlendes Feld muss leer, nicht None sein'
    assert b['framework_bewertungen'][0]['luecken'] == []
    assert b['handlungsempfehlungen'][0]['prioritaet'] == ''
    assert not _none_pfade(ctx)


# ── Inhaltliche Treue zum gerechneten Stand ──────────────────────────────────

def test_stand_stimmt_mit_der_berechnung(audit_db):
    """Die Vorlage darf keine andere Zahl drucken als die App anzeigt.

    1 erfüllt + 1 teilweise + 1 nicht erfüllt + 1 offen über 4 anwendbare Fragen
    → (1 + 0,5) / 4 = 37,5 %.
    """
    from gutachten.db import load_questions

    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    erwartet = status.build_status(load_questions(audit_db, 'VorlageAudit'))
    assert ctx['meta']['quote_prozent'] == erwartet['quote_prozent'] == 37.5
    assert ctx['meta']['stufe_label'] == erwartet['label']
    assert ctx['meta']['stufe'] == erwartet['stufe']
    assert ctx['meta']['bedeutung'], 'die Stufe muss erklärt werden'


def test_offene_frage_senkt_die_quote_nicht_die_abdeckung(audit_db):
    """Die kontraintuitive Regel aus #1713 muss auch in der Vorlage gelten."""
    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    assert ctx['meta']['fragen']['offen'] == 1
    assert ctx['meta']['abdeckung_prozent'] == 75.0
    # Nur über die bewerteten Fragen gerechnet faellt die Quote besser aus.
    assert ctx['meta']['quote_bewertet_prozent'] > ctx['meta']['quote_prozent']


def test_je_framework_getrennt(audit_db):
    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    nach = {f['framework']: f for f in ctx['frameworks']}
    assert set(nach) == {'NIS2', 'CRA'}
    assert nach['NIS2']['quote_prozent'] == 75.0
    assert nach['CRA']['quote_prozent'] == 0.0
    assert nach['NIS2']['fragen']['gesamt'] == 2


def test_luecken_nicht_erfuellt_zuerst(audit_db):
    """Reihenfolge ist eine Aussage: zuerst, was fehlt."""
    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    folge = [lk['erfuellung'] for lk in ctx['luecken']]
    assert folge == ['nicht_erfuellt', 'teilweise', 'offen'], folge
    assert ctx['meta']['anzahl_luecken'] == 3
    assert ctx['luecken'][0]['fundstelle'] == 'Art. 13(1)'
    assert ctx['luecken'][0]['erfuellung_label'] == 'nicht erfüllt'


def test_fragen_tragen_lesbare_erfuellung(audit_db):
    """Im Dokument steht „nicht erfüllt“, nicht ``nicht_erfuellt``."""
    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    labels = {f['erfuellung_label'] for f in ctx['fragen']}
    assert 'nicht_erfuellt' not in labels, 'Schlüsselwort statt Klartext im Dokument'
    assert {'erfüllt', 'teilweise erfüllt', 'nicht erfüllt', 'offen'} == labels


def test_projekt_stammdaten(audit_db):
    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    assert ctx['projekt']['unternehmen'] == 'Acme GmbH'
    assert ctx['projekt']['pruefungsfokus'] == 'Lieferkette'
    assert ctx['projekt']['frameworks'] == ['NIS2', 'CRA']
    # Deckblatt/Fusszeile brauchen einen Text, keine Liste.
    assert ctx['projekt']['frameworks_text'] == 'NIS2, CRA'


def test_erfuellungs_labels_haben_eine_quelle():
    """Bericht und Vorlage dürfen nicht eigene Beschriftungen führen."""
    from compliance_audit import berichte_provider as prov
    assert prov._ERF_LABEL is status.ERFUELLUNG_LABELS
    assert set(status.ERFUELLUNG_LABELS) == set(status.ERFUELLUNG)


# ── Ende zu Ende durch die echte Engine ──────────────────────────────────────

def test_kundenvorlage_rendert_durch_die_engine(audit_db, tmp_path):
    """Der eigentliche Zweck: eine Kundenvorlage mit Jinja füllen.

    Geprüft wird der ganze Weg — Kontext-Builder, Sandbox-Umgebung von
    ``shared.templates.engine`` und die Schleifen, die ein Anwender typischerweise
    schreibt (Lücken-Tabelle, Maßnahmen). Die Einzeltests oben sichern den
    Kontext; dieser sichert, dass er in einem echten Dokument ankommt.
    """
    from docx import Document

    from shared.templates.engine import render_docx_from_path

    vorlage = tmp_path / 'kunde.docx'
    d = Document()
    d.add_heading('Auditbericht {{ projekt.name }}', 1)
    d.add_paragraph('Organisation: {{ projekt.unternehmen }}')
    d.add_paragraph('Frameworks: {{ projekt.frameworks_text }}')
    d.add_paragraph('Stand: {{ meta.stufe_label }} ({{ meta.quote_prozent }} %)')
    d.add_paragraph('Erhebungsstand: {{ meta.abdeckung_prozent }} %')
    d.add_paragraph('{% for lk in luecken %}LUECKE {{ lk.fundstelle }}: '
                    '{{ lk.erfuellung_label }}{% endfor %}')
    d.add_paragraph('{% for f in frameworks %}FW {{ f.framework }}='
                    '{{ f.quote_prozent }}{% endfor %}')
    d.add_paragraph('Fazit: {{ bericht.fazit }}')
    d.save(vorlage)

    ctx = tc.build_compliance_audit_context(audit_db, 'VorlageAudit')
    raus = render_docx_from_path(vorlage, ctx)
    text = '\n'.join(p.text for p in Document(__import__('io').BytesIO(raus)).paragraphs)

    assert 'Auditbericht VorlageAudit' in text
    assert 'Acme GmbH' in text
    assert 'NIS2, CRA' in text
    # 37,5 % liegt unter der 40-%-Schwelle → "nicht compliant" (SCHWELLEN).
    assert 'nicht compliant (37.5 %)' in text
    assert 'LUECKE Art. 13(1): nicht erfüllt' in text
    assert 'FW NIS2=75.0' in text and 'FW CRA=0.0' in text
    # Der KI-Entwurf fehlt in diesem Projekt — das darf kein "None" drucken.
    assert 'None' not in text, f'None im gerenderten Dokument:\n{text}'
    assert '{{' not in text and '{%' not in text, 'Platzhalter nicht ersetzt'
