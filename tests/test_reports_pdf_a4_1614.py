"""#1614: Berichts-PDFs sind A4 und vektoriell (Textebene), kein Rasterbild."""
import shutil
from io import BytesIO

import pytest


def test_set_a4():
    from docx import Document
    from shared.reports.docx_pdf import set_a4
    d = Document()
    set_a4(d)
    s = d.sections[0]
    # EMU → mm (1 mm = 36000 EMU)
    assert round(s.page_width / 36000) == 210
    assert round(s.page_height / 36000) == 297


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice nicht verfügbar")
def test_docx_to_pdf_vektor_a4(monkeypatch):
    """DOCX→PDF ergibt A4 mit durchsuchbarer Textebene (nicht gerastert)."""
    # Gotenberg bewusst ausschalten → soffice-Fallback
    monkeypatch.setenv("GOTENBERG_URL", "http://127.0.0.1:1")
    from docx import Document
    from shared.reports.docx_pdf import set_a4, docx_bytes_to_pdf
    d = Document(); set_a4(d)
    d.add_heading("Konformitaetsnachweis Pruefung", 0)
    d.add_paragraph("Vektorieller Text 1234567890 ABCDEF.")
    bio = BytesIO(); d.save(bio)
    pdf = docx_bytes_to_pdf(bio.getvalue(), timeout_s=90)
    import pdfplumber
    with pdfplumber.open(BytesIO(pdf)) as p:
        pg = p.pages[0]
        assert 590 < pg.width < 600 and 835 < pg.height < 848   # A4 ~595x842 pt
        assert "Pruefung" in (pg.extract_text() or "")          # echte Textebene
