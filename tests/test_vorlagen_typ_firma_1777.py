"""Vorlagen je Berichtstyp und je Mandant, Nachweis im Berichts-Center
(#1776/#1777/#1778).

Der Anlass kam aus dem Betrieb: Im CRA-Berichts-Center fehlte der
**Konformitätsnachweis** ganz — er hing allein am Knopf im Releases-Tab. Wer ihn
dort suchte, klickte den nächstliegenden Typ und bekam den Readiness-Bericht.
Und weil eine als Standard markierte Word-Vorlage **modulweit** galt, lieferten
mit gesetztem Standard ohnehin alle Karten dasselbe Papier: Die Brücke ruft
nicht den Typ-Renderer, sondern baut einmal den Modul-Kontext und rendert die
Vorlage.

Vier Zusagen, für jede ein Test:

1. Der Nachweis ist ein Berichtstyp, und jeder Typ erzeugt sein eigenes Dokument.
2. Eine Vorlage gilt für **einen** Bericht (oder ausdrücklich für alle), und ein
   Standard ist je Bericht und Mandant exklusiv — nicht je Modul.
3. Ein Dokument mit vorgeschriebener Struktur (Erklärung, Nachweis) bekommt
   keine modulweite Vorlage von selbst.
4. Eine Kundenvorlage erzeugt nur Dokumente **ihres** Mandanten, und wer
   Pflichtangaben auslässt, wird nicht Standard.

Die DB-abhängigen Tests hängen am ``pg``-Fixture (tests/conftest.py).
"""
from __future__ import annotations

import io
import uuid
from pathlib import Path

import pytest
from docx import Document

from shared.reports import vorlagen as bruecke
from shared.templates import db as tdb
from shared.templates import schema, storage, validate


# ──────────────────────────────────────────────────────────────────────────────
# Hilfsmittel
# ──────────────────────────────────────────────────────────────────────────────
def _vorlage(pfad: Path, zeilen: list[str]) -> Path:
    pfad.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    for z in zeilen:
        d.add_paragraph(z)
    d.save(str(pfad))
    return pfad


def _eintrag(tid: int, *, typ: str = "", firma: int | None = None,
             default: bool = True) -> dict:
    """Ein Listeneintrag, wie ``bruecke.vorlagen`` ihn liefert."""
    return {"id": tid, "name": f"v{tid}", "version": 1, "ist_default": default,
            "bericht_typ": typ, "firmen_id": firma}


@pytest.fixture
def rest(pg, request):
    """Testclient + Admin-Header, aber erst nach dem Postgres-Check (wie #1741)."""
    return request.getfixturevalue("client"), request.getfixturevalue("auth_headers")


@pytest.fixture
def upload(rest, tmp_path):
    """Lädt eine CRA-Vorlage hoch; liefert den Antwort-Body."""
    client, headers = rest

    def _tun(zeilen: list[str], **form) -> dict:
        pfad = _vorlage(tmp_path / f"{uuid.uuid4().hex}.docx", zeilen)
        daten = {"file": (io.BytesIO(pfad.read_bytes()), "v.docx"),
                 "modul": "cra", "name": f"pytest-1777-{uuid.uuid4().hex[:8]}"}
        daten.update({k: str(v) for k, v in form.items()})
        res = client.post("/api/templates", data=daten,
                          content_type="multipart/form-data", headers=headers)
        assert res.status_code == 201, res.get_data(as_text=True)
        return res.get_json()

    return _tun


# ──────────────────────────────────────────────────────────────────────────────
# 1 · Der Nachweis ist ein Berichtstyp
# ──────────────────────────────────────────────────────────────────────────────
def test_nachweis_steht_im_katalog_und_ist_vorgeschrieben():
    from cra import berichte_provider as bp

    spec = next((s for s in bp.CATALOG if s.key == "konformitaetsnachweis"), None)
    assert spec is not None, "Konformitätsnachweis fehlt im Berichts-Center"
    assert spec.vorgeschrieben is True
    # Die Erklärung ebenso — beide Dokumente folgen dem Rechtstext, nicht uns.
    erkl = next(s for s in bp.CATALOG if s.key == "konformitaetserklaerung")
    assert erkl.vorgeschrieben is True
    # Ein Bericht ist kein vorgeschriebenes Dokument.
    assert next(s for s in bp.CATALOG if s.key == "gesamt").vorgeschrieben is False


def test_jeder_katalog_typ_erzeugt_ein_eigenes_dokument(pg):
    """Die eigentliche Fehlanzeige: zwei Karten, dasselbe Dokument.

    Geprüft wird über die erste Überschrift — sie ist das, was der Anwender
    sieht, wenn er die Datei öffnet.
    """
    from cra.db import ensure_db
    from cra import berichte_provider as bp

    ensure_db(bp.DB_PATH)
    projekt = f"pytest-1777-{uuid.uuid4().hex[:8]}"

    ueberschriften: dict[str, str] = {}
    for spec in bp.CATALOG:
        if spec.key in ("gesamt", "anforderungen", "massnahmen"):
            continue    # dieselbe Readiness-Vorlage mit anderen Abschnitten
        daten = bp.render(spec.key, "docx", projekt=projekt)
        doc = Document(io.BytesIO(daten))
        erste = next((p.text.strip() for p in doc.paragraphs if p.text.strip()), "")
        ueberschriften[spec.key] = erste

    assert len(set(ueberschriften.values())) == len(ueberschriften), ueberschriften
    assert "Konformitätsnachweis" in ueberschriften["konformitaetsnachweis"]
    assert "Konformitätserklärung" in ueberschriften["konformitaetserklaerung"]


def test_nachweis_laeuft_ueber_dieselbe_funktion_wie_der_releases_knopf(pg):
    """Ein zweiter Renderer wären zwei Konformitätsaussagen zum selben Release."""
    from cra.db import ensure_db
    from cra import berichte_provider as bp
    from cra import konformitaet_report as kr

    ensure_db(bp.DB_PATH)
    projekt = f"pytest-1777-{uuid.uuid4().hex[:8]}"

    def _text(data: bytes) -> list[str]:
        d = Document(io.BytesIO(data))
        return [p.text.strip() for p in d.paragraphs if p.text.strip()]

    assert _text(bp.render("konformitaetsnachweis", "docx", projekt=projekt)) == \
        _text(kr.build_docx(projekt))


# ──────────────────────────────────────────────────────────────────────────────
# 2 · Kontext: der Nachweis ist für eine Vorlage überhaupt füllbar
# ──────────────────────────────────────────────────────────────────────────────
def test_nachweis_variablen_sind_vollstaendig_aufloesbar(pg):
    """Ohne diese Variablen konnte eine Vorlage den Nachweis nicht erzeugen.

    Geprüft an einem **unbekannten** Projekt: Eine frische Installation ist der
    Fall, in dem ein fehlender Schlüssel durchrutscht.
    """
    from cra.db import ensure_db
    from cra.template_context import build_cra_context

    db = schema.MODULE_DB["cra"]
    ensure_db(db)
    ctx = build_cra_context(db, f"pytest-1777-{uuid.uuid4().hex[:8]}")

    assert "nachweis" in ctx
    deklariert = [v["key"] for v in schema.get_variables("cra")
                  if v["key"].startswith("nachweis.")]
    assert deklariert, "keine nachweis-Variablen deklariert"
    for key in deklariert:
        teil = key.split(".", 1)[1]
        assert teil in ctx["nachweis"], f"{key} deklariert, aber nicht geliefert"
        assert ctx["nachweis"][teil] is not None, f"{key} ist None — Jinja druckt das"


def test_builder_nimmt_version_nur_wenn_gefragt(pg):
    """Die sieben übrigen Adapter kennen ``version`` nicht.

    Würde die Brücke sie unbedingt durchreichen, brächen sie mit ``TypeError``.
    """
    from cra.db import ensure_db

    ensure_db(schema.MODULE_DB["cra"])
    projekt = f"pytest-1777-{uuid.uuid4().hex[:8]}"
    for modul in ("cra", "nis2"):
        from importlib import import_module
        pfad = {"cra": "cra.db", "nis2": "nis2.db"}[modul]
        import_module(pfad).ensure_db(schema.MODULE_DB[modul])
        ctx = bruecke.kontext(modul, projekt, typ="gesamt", version="1.2.3")
        assert isinstance(ctx, dict) and ctx


# ──────────────────────────────────────────────────────────────────────────────
# 3 · Vorauswahl: die Regel, die den Fehler verursacht hat
# ──────────────────────────────────────────────────────────────────────────────
def test_modulweite_vorlage_gilt_fuer_gewoehnliche_berichte():
    """Das Verhalten bis v8.0.3 bleibt für Berichte erhalten."""
    alle = [_eintrag(1)]
    assert bruecke.vorauswahl("cra", "gesamt", alle=alle) == 1


def test_modulweite_vorlage_greift_nicht_bei_vorgeschriebenem_dokument():
    """Der eigentliche Fehler: ein Haus-Standard ersetzte lautlos ein Rechtsdokument."""
    alle = [_eintrag(1)]
    assert bruecke.vorauswahl("cra", "konformitaetsnachweis", alle=alle,
                              vorgeschrieben=True) == 0
    assert bruecke.vorauswahl("cra", "konformitaetserklaerung", alle=alle,
                              vorgeschrieben=True) == 0


def test_typ_vorlage_gewinnt_gegen_modulweite():
    alle = [_eintrag(1), _eintrag(2, typ="konformitaetsnachweis")]
    assert bruecke.vorauswahl("cra", "konformitaetsnachweis", alle=alle,
                              vorgeschrieben=True) == 2
    assert bruecke.vorauswahl("cra", "gesamt", alle=alle) == 1


def test_mandantenvorlage_gewinnt_gegen_hausvorlage():
    alle = [_eintrag(1, typ="konformitaetsnachweis"),
            _eintrag(2, typ="konformitaetsnachweis", firma=7)]
    assert bruecke.vorauswahl("cra", "konformitaetsnachweis", firmen_id=7, alle=alle,
                              vorgeschrieben=True) == 2
    # Ein anderer Mandant fällt auf die Haus-Vorlage zurück.
    assert bruecke.vorauswahl("cra", "konformitaetsnachweis", firmen_id=8, alle=alle,
                              vorgeschrieben=True) == 1


def test_ohne_standard_wird_nichts_vorgewaehlt():
    """Eine hochgeladene Vorlage wirkt erst, wenn jemand sie zum Standard macht."""
    alle = [_eintrag(1, typ="konformitaetsnachweis", default=False)]
    assert bruecke.vorauswahl("cra", "konformitaetsnachweis", alle=alle,
                              vorgeschrieben=True) == 0


def test_gilt_fuer_typ():
    assert bruecke.gilt_fuer_typ({"bericht_typ": ""}, "gesamt") is True
    assert bruecke.gilt_fuer_typ({"bericht_typ": "gesamt"}, "gesamt") is True
    assert bruecke.gilt_fuer_typ({"bericht_typ": "gesamt"}, "massnahmen") is False


# ──────────────────────────────────────────────────────────────────────────────
# 4 · Registry: Standard je Bericht und Mandant
# ──────────────────────────────────────────────────────────────────────────────
def _anlegen(tmp_path: Path, **felder) -> dict:
    pfad = _vorlage(tmp_path / f"{uuid.uuid4().hex}.docx", ["{{ projekt.name }}"])
    return storage.upload_template(pfad, modul="cra",
                                   name=f"t-{uuid.uuid4().hex[:8]}",
                                   db_path=tdb.DEFAULT_DB_PATH,
                                   storage_root=tmp_path / "store", **felder)


def test_standard_ist_je_bericht_und_mandant_exklusiv(pg, tmp_path):
    """Vorher verlor **jede** andere Vorlage des Moduls das Flag.

    Damit war nicht sagbar „diese Vorlage ist unsere Erklärung **und** jene
    unser Gesamtbericht".
    """
    a = _anlegen(tmp_path, bericht_typ="gesamt")
    b = _anlegen(tmp_path, bericht_typ="konformitaetsnachweis")
    c = _anlegen(tmp_path, bericht_typ="konformitaetsnachweis", firmen_id=7)

    for rec in (a, b, c):
        tdb.set_default(tdb.DEFAULT_DB_PATH, rec["id"])

    aktuell = {r["id"]: r for r in tdb.list_templates(tdb.DEFAULT_DB_PATH, "cra")}
    assert aktuell[a["id"]]["ist_default"]
    assert aktuell[b["id"]]["ist_default"]
    assert aktuell[c["id"]]["ist_default"]

    # Zweite Vorlage für denselben Bericht desselben Mandanten verdrängt die erste.
    d = _anlegen(tmp_path, bericht_typ="konformitaetsnachweis", firmen_id=7)
    tdb.set_default(tdb.DEFAULT_DB_PATH, d["id"])
    aktuell = {r["id"]: r for r in tdb.list_templates(tdb.DEFAULT_DB_PATH, "cra")}
    assert not aktuell[c["id"]]["ist_default"]
    assert aktuell[d["id"]]["ist_default"]
    assert aktuell[b["id"]]["ist_default"], "Haus-Vorlage darf nicht mitfallen"


def test_liste_mit_mandant_zeigt_eigene_und_hausvorlagen(pg, tmp_path):
    haus = _anlegen(tmp_path)
    mein = _anlegen(tmp_path, firmen_id=7)
    fremd = _anlegen(tmp_path, firmen_id=8)

    ids = {r["id"] for r in tdb.list_templates(tdb.DEFAULT_DB_PATH, "cra", firmen_id=7)}
    assert haus["id"] in ids and mein["id"] in ids
    assert fremd["id"] not in ids, "fremde Mandantenvorlage in der Auswahl"

    nur_haus = {r["id"] for r in tdb.list_templates(tdb.DEFAULT_DB_PATH, "cra",
                                                    nur_haus=True)}
    assert nur_haus == {haus["id"]}


def test_firmen_id_null_bedeutet_hausvorlage():
    """0, '' und Müll dürfen keine Vorlage erzeugen, die zu keinem Mandanten gehört."""
    for wert in (0, "", "null", None, "abc", -3):
        assert tdb._firmen_id(wert) is None
    assert tdb._firmen_id("7") == 7


# ──────────────────────────────────────────────────────────────────────────────
# 5 · Abgrenzung: fremder Mandant, fremder Bericht
# ──────────────────────────────────────────────────────────────────────────────
def test_vorlage_eines_anderen_berichts_wird_abgelehnt(pg, tmp_path):
    rec = _anlegen(tmp_path, bericht_typ="gesamt")
    with pytest.raises(bruecke.VorlageFehlt) as fehler:
        bruecke.pruefe("cra", rec["id"], typ="konformitaetsnachweis")
    assert "gesamt" in str(fehler.value)
    # Für ihren eigenen Bericht bleibt sie gültig.
    assert bruecke.pruefe("cra", rec["id"], typ="gesamt")


def test_vorlage_eines_anderen_mandanten_wird_abgelehnt(pg, tmp_path, monkeypatch):
    rec = _anlegen(tmp_path, firmen_id=8)
    monkeypatch.setattr(bruecke, "projekt_firma", lambda modul, projekt: 7)
    with pytest.raises(bruecke.VorlageFehlt) as fehler:
        bruecke.pruefe("cra", rec["id"], projekt="egal")
    assert "Mandanten" in str(fehler.value)


def test_hausvorlage_gilt_fuer_jeden_mandanten(pg, tmp_path, monkeypatch):
    rec = _anlegen(tmp_path)
    monkeypatch.setattr(bruecke, "projekt_firma", lambda modul, projekt: 7)
    assert bruecke.pruefe("cra", rec["id"], projekt="egal")


# ──────────────────────────────────────────────────────────────────────────────
# 6 · Pflichtangaben des Nachweises
# ──────────────────────────────────────────────────────────────────────────────
def test_nachweis_vorlage_ohne_release_wird_benannt():
    """Ein Nachweis ohne Release belegt nichts Bestimmtes.

    Das ist keine erfundene Rechtspflicht, sondern der Zweck des Dokuments:
    Es ist der versionsgebundene Nachweis.
    """
    erg = validate.pruefe_namen("cra", ["nachweis"], inhalt=(
        {"nachweis", "nachweis.produkt", "nachweis.rechtsgrundlage",
         "nachweis.ampel_label"}, ""))
    luecken = [f["bezeichnung"] for f in erg.fehlende_pflicht]
    assert any("Release" in l for l in luecken), erg.fehlende_pflicht
    assert not erg.ok
    # Ohne Fundstelle keine Nummer — „Nr. None" behauptete eine.
    assert "Nr. None" not in erg.meldung


def test_vollstaendige_nachweis_vorlage_ist_in_ordnung():
    erg = validate.pruefe_namen("cra", ["nachweis"], inhalt=(
        {"nachweis", "nachweis.produkt", "nachweis.version",
         "nachweis.rechtsgrundlage", "nachweis.ampel_label"}, ""))
    assert erg.fehlende_pflicht == []
    assert erg.ok


def test_gesamtbericht_vorlage_bleibt_von_der_pflichtpruefung_unberuehrt():
    """Bei einem Bericht ist die Vorlage Geschmacksfrage (#1753-Grenze)."""
    erg = validate.pruefe_namen("cra", ["projekt", "meta"],
                                inhalt=({"projekt.name", "meta.datum"}, ""))
    assert erg.ok and erg.fehlende_pflicht == []


def test_beide_pflichtquellen_werden_gelesen():
    """Erklärung und Nachweis bringen ihre Pflichtliste je selbst mit (#1731)."""
    dokumente = {s["dokument"] for s in validate.pflicht_vorlagen("cra")}
    assert {"EU-Konformitätserklärung", "Konformitätsnachweis"} <= dokumente


# ──────────────────────────────────────────────────────────────────────────────
# 7 · REST
# ──────────────────────────────────────────────────────────────────────────────
def test_upload_mit_bericht_typ_und_mandant(upload):
    body = upload(["{{ projekt.name }}"], bericht_typ="gesamt", firmen_id=7)
    assert body["bericht_typ"] == "gesamt"
    assert body["firmen_id"] == 7


def test_upload_mit_unbekanntem_bericht_typ_wird_abgewiesen(rest, tmp_path):
    client, headers = rest
    pfad = _vorlage(tmp_path / "x.docx", ["{{ projekt.name }}"])
    res = client.post("/api/templates",
                      data={"file": (io.BytesIO(pfad.read_bytes()), "v.docx"),
                            "modul": "cra", "name": "pytest-1777-falsch",
                            "bericht_typ": "gibt-es-nicht"},
                      content_type="multipart/form-data", headers=headers)
    assert res.status_code == 400
    assert "Berichtstyp" in res.get_json()["error"]


def test_berichtstypen_endpunkt_nennt_die_karten(rest):
    client, headers = rest
    res = client.get("/api/templates/berichte/cra", headers=headers)
    assert res.status_code == 200
    keys = [b["key"] for b in res.get_json()["berichte"]]
    assert "konformitaetsnachweis" in keys and "gesamt" in keys


def test_berichts_center_nennt_die_vorgewaehlte_vorlage(rest, upload, pg):
    """Die Oberfläche soll die Regel nicht nachbauen (Lehre aus #1703/S9)."""
    from cra.db import ensure_db, save_projekt

    client, headers = rest
    ensure_db(Path("data/db/cra.sqlite"))
    projekt = f"pytest-1777-{uuid.uuid4().hex[:8]}"
    save_projekt(Path("data/db/cra.sqlite"), name=projekt, unternehmen="Test GmbH",
                 produkt="P", produktklasse="default", beschreibung="", berater="")

    body = upload(["{{ projekt.name }}"], bericht_typ="gesamt")
    res = client.put(f"/api/templates/{body['id']}/default", headers=headers)
    assert res.status_code == 200, res.get_data(as_text=True)

    res = client.get(f"/api/cra/projekte/{projekt}/berichte", headers=headers)
    assert res.status_code == 200
    karten = {r["key"]: r for r in res.get_json()["reports"]}
    assert karten["gesamt"]["vorlage_id"] == body["id"]
    # Und genau das, was der Anwender gemeldet hat: der Nachweis bleibt frei.
    assert karten["konformitaetsnachweis"]["vorlage_id"] == 0


def test_export_mit_falscher_vorlage_antwortet_404_statt_500(rest, upload, pg):
    client, headers = rest
    body = upload(["{{ projekt.name }}"], bericht_typ="gesamt")
    res = client.get(
        f"/api/cra/projekte/Egal/berichte/konformitaetsnachweis"
        f"?format=docx&template_id={body['id']}", headers=headers)
    assert res.status_code == 404, res.get_data(as_text=True)
