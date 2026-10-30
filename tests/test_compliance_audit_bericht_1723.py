"""Auditbericht, KI-Parität und der 405-Fehler (#1723).

Drei Befunde aus dem Praxistest des neuen Moduls:

1. **Bericht-Tab, HTTP 500.** ``/bericht/prompt`` rief
   ``gutachten.prompts._build_gutachten_prompt`` mit Argumenten auf, die dessen
   Signatur nicht kennt (``TypeError: unexpected keyword argument 'pruefungsfokus'``).
   Fachlich war es ohnehin falsch — hier entsteht ein **Auditbericht**, kein Gutachten.

2. **Wasserstandsmeldung, HTTP 405 von nginx.** ``KiStreamView`` bekam die URL
   ``/compliance-audit/…/wasserstand/stream`` **ohne** ``/api``-Präfix.
   ``useSseRun.runSse`` nutzt rohes ``fetch()`` ohne die axios-``baseURL``, der POST
   landete also in nginx' ``location /`` mit ``try_files`` — und ein POST auf eine
   statische Datei ergibt genau diesen 405er. Der Test unten hält alle
   ``KiStreamView``-URLs auf ``/api/`` fest.

3. **Alter KI-Weg.** Fragen und Bericht liefen nur über Copy/Paste in zwei
   Eigenbau-Dialogen. Alle anderen Module nutzen ``WizardPromptModal`` mit
   „⚡ Direkt mit KI ausführen" über ``/api/ai/run-stream``.
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest

from compliance_audit import prompts, status

REPO = Path(__file__).resolve().parent.parent
FRONTEND = REPO / 'frontend/src'


# ── Auditbericht-Prompt ──────────────────────────────────────────────────────

def _projekt():
    return {'name': 'A1', 'unternehmen': 'Acme', 'frameworks': ['NIS2', 'CRA'],
            'pruefungsfokus': 'Lieferkette'}


def _fragen(n: int = 3):
    return [{'framework': 'NIS2', 'section_ref': f'Art. {i}', 'frage': f'Frage {i}',
             'antwort': f'Antwort {i}', 'erfuellung': 'erfuellt'} for i in range(n)]


def test_prompt_ist_ein_auditbericht_kein_gutachten():
    p = prompts.build_bericht_prompt(_projekt(), _fragen())
    assert 'Auditbericht' in p or 'Compliance-Audit' in p
    assert 'Gutachten' not in p, 'der Bericht darf sich nicht als Gutachten ausgeben'


def test_prompt_fordert_das_zielschema():
    """Das Schema muss zu ``/bericht/import`` und zum DOCX-Export passen."""
    p = prompts.build_bericht_prompt(_projekt(), _fragen())
    for feld in prompts.BERICHT_FELDER:
        assert feld in p, f'Pflichtfeld {feld} fehlt im Prompt'


def test_prompt_enthaelt_den_berechneten_stand():
    """Bericht und Wasserstandsmeldung müssen dieselbe Aussage treffen."""
    st = status.build_status([{'framework': 'NIS2', 'erfuellung': 'erfuellt', 'frage': 'A'}] * 9
                             + [{'framework': 'NIS2', 'erfuellung': 'nicht_erfuellt', 'frage': 'B'}])
    p = prompts.build_bericht_prompt(_projekt(), _fragen(), st)
    assert 'weitgehend compliant' in p
    assert 'nicht neu bewerten' in p


def test_prompt_kuerzt_lange_fragenlisten():
    p = prompts.build_bericht_prompt(_projekt(), _fragen(200), max_fragen=10)
    assert 'Frage 9' in p
    assert 'Frage 150' not in p
    assert 'weitere Fragen' in p, 'die Kürzung muss ausgewiesen werden'


# ── Parser ───────────────────────────────────────────────────────────────────

VOLL = ('{"gesamtbewertung":"solide","zusammenfassung":"Z","fazit":"F",'
        '"framework_bewertungen":[{"framework":"NIS2","bewertung":"ok",'
        '"staerken":["A"],"luecken":["B"]}],'
        '"handlungsempfehlungen":[{"massnahme":"M","begruendung":"B",'
        '"prioritaet":"Hoch","aufwand":"Mittel"}]}')


def test_parser_normalisiert():
    out = prompts.parse_bericht_response(VOLL)
    assert out['gesamtbewertung'] == 'solide'
    assert out['framework_bewertungen'][0]['staerken'] == ['A']
    assert out['handlungsempfehlungen'][0]['prioritaet'] == 'hoch'  # kleingeschrieben
    assert out['handlungsempfehlungen'][0]['aufwand'] == 'mittel'


def test_parser_vertraegt_rahmen_und_beitext():
    roh = f'Gerne, hier der Bericht:\n```json\n{VOLL}\n```\nViel Erfolg!'
    assert prompts.parse_bericht_response(roh)['fazit'] == 'F'


def test_parser_vertraegt_massnahmen_als_strings():
    """Manche Antworten liefern nur Strings — das soll nicht scheitern."""
    roh = ('{"gesamtbewertung":"g","zusammenfassung":"z","fazit":"f",'
           '"handlungsempfehlungen":["MFA einfuehren","Backups testen"]}')
    out = prompts.parse_bericht_response(roh)
    assert [m['massnahme'] for m in out['handlungsempfehlungen']] == \
        ['MFA einfuehren', 'Backups testen']


@pytest.mark.parametrize('roh', [
    'Dazu kann ich nichts sagen.',
    '{"zusammenfassung":"nur eins"}',
    '[]',
])
def test_parser_lehnt_unbrauchbares_ab(roh):
    with pytest.raises(ValueError):
        prompts.parse_bericht_response(roh)


# ── API ──────────────────────────────────────────────────────────────────────

@pytest.fixture
def audit_projekt(monkeypatch):
    from compliance_audit import db as adb
    from gutachten.db import save_project, save_questions

    p = Path('data/db/pytest_audit_1723.sqlite')
    p.parent.mkdir(parents=True, exist_ok=True)
    adb.ensure_db(p)
    con = adb._connect(p)
    try:
        for t in ('audit_wasserstand', 'gutachten_questions', 'gutachten_projects',
                  'gutachten_drafts'):
            con.execute(f'DELETE FROM {t}')  # nosec B608 — Konstanten
        con.commit()
    finally:
        con.close()
    save_project(p, 'BerichtAudit', frameworks=['NIS2'], pruefungsfokus='Fokus',
                 unternehmen='Acme')
    save_questions(p, 'BerichtAudit', [
        {'question_num': 1, 'framework': 'NIS2', 'frage': 'ISMS vorhanden?',
         'antwort': 'Ja', 'bewertung': 'erfüllt'}])

    import server.api.compliance_audit as bp
    monkeypatch.setattr(bp, 'DB_PATH', p)
    return p


def test_api_bericht_prompt_liefert_200(client, auth_headers, audit_projekt):
    """Vorher: HTTP 500 durch die Signatur-Abweichung."""
    r = client.post('/api/compliance-audit/BerichtAudit/bericht/prompt',
                    json={}, headers=auth_headers)
    assert r.status_code == 200, r.data[:400]
    assert r.json['questions_count'] == 1
    assert 'Gutachten' not in r.json['prompt']


def test_api_bericht_prompt_ohne_fragen_400(client, auth_headers, audit_projekt):
    from gutachten.db import save_project
    save_project(audit_projekt, 'Leer', frameworks=['NIS2'], pruefungsfokus='',
                 unternehmen='')
    r = client.post('/api/compliance-audit/Leer/bericht/prompt',
                    json={}, headers=auth_headers)
    assert r.status_code == 400


def test_api_bericht_import_und_entwurf(client, auth_headers, audit_projekt):
    r = client.post('/api/compliance-audit/BerichtAudit/bericht/import',
                    json={'raw': VOLL}, headers=auth_headers)
    assert r.status_code == 200, r.data[:300]
    assert sorted(r.json['draft']) == sorted(prompts.BERICHT_FELDER)
    r = client.get('/api/compliance-audit/BerichtAudit/bericht/draft',
                   headers=auth_headers)
    assert r.status_code == 200
    assert r.json['draft']['gesamtbewertung'] == 'solide'


def test_api_bericht_import_lehnt_muell_ab(client, auth_headers, audit_projekt):
    r = client.post('/api/compliance-audit/BerichtAudit/bericht/import',
                    json={'raw': 'kein json'}, headers=auth_headers)
    assert r.status_code == 400
    assert 'nicht lesbar' in r.json['error']


# ── Frontend-Wächter ─────────────────────────────────────────────────────────

def test_ki_stream_urls_haben_api_praefix():
    """Der 405er: ``runSse`` nutzt rohes ``fetch()`` ohne axios-``baseURL``.

    Eine URL ohne ``/api`` landet in nginx' ``location /`` mit ``try_files``; ein
    POST auf eine statische Datei ergibt „405 Not Allowed" — im Browser kaum als
    Programmierfehler zu erkennen.

    Gesucht wird im ``<KiStreamView>``-Tag. Steht dort eine **Variable**
    (``:url="streamUrl"``), wird deren Definition in derselben Datei nachgesehen —
    sonst wäre der Test entweder blind oder voller Falschmeldungen.
    """
    # Zwei Einstiege: KiStreamView direkt und AutoBewertungDialog, das die URL
    # als Prop durchreicht — dort liegt die Verantwortung beim Aufrufer.
    tag = re.compile(r'<(?:KiStreamView|AutoBewertungDialog)\b[^>]*', re.DOTALL)
    attr = re.compile(r':(?:url|stream-url)\s*=\s*"([^"]+)"')
    bezeichner = re.compile(r'^[A-Za-z_$][\w$]*$')
    schlecht: list[tuple[str, str, str]] = []
    geprueft = 0

    for datei in FRONTEND.rglob('*.vue'):
        txt = datei.read_text(encoding='utf-8', errors='replace')
        for t in tag.finditer(txt):
            for m in attr.finditer(t.group(0)):
                roh = m.group(1).strip()
                geprueft += 1
                if '/api/' in roh or '/api$' in roh or '/api`' in roh:
                    continue
                if roh == 'streamUrl' and datei.name == 'AutoBewertungDialog.vue':
                    continue  # Prop — geprueft wird an den Aufrufstellen
                if '.streamUrl(' in roh:
                    # Adapter-Helfer (utils/massAdapters.ts): dort steht die URL.
                    adapter = (FRONTEND / 'utils/massAdapters.ts').read_text()
                    # Nur Implementierungen pruefen (Template-Literal), nicht die
                    # Typdeklaration `streamUrl: (id) => string`.
                    impl = [z for z in adapter.splitlines()
                            if 'streamUrl:' in z and '`' in z]
                    if impl and all('/api' in z for z in impl):
                        continue
                    schlecht.append((str(datei.relative_to(FRONTEND)), roh,
                                     'Adapter liefert URL ohne /api'))
                    continue
                if bezeichner.match(roh):
                    # Variable: Definition in derselben Datei pruefen.
                    defi = re.search(
                        rf'(?:const|let|var)\s+{re.escape(roh)}\s*=(.{{0,400}})',
                        txt, re.DOTALL)
                    if defi and '/api' in defi.group(1):
                        continue
                    schlecht.append((str(datei.relative_to(FRONTEND)), roh,
                                     'Definition ohne /api' if defi else 'Definition nicht gefunden'))
                    continue
                schlecht.append((str(datei.relative_to(FRONTEND)), roh, 'Literal ohne /api'))

    assert geprueft >= 5, f'nur {geprueft} KiStreamView-URLs gefunden — Erkennung kaputt?'
    assert not schlecht, (
        'Diese KiStreamView-URLs erreichen das Backend nicht (nginx antwortet 405):\n' +
        '\n'.join(f'  {d}: {u} — {grund}' for d, u, grund in schlecht))


def test_audit_nutzt_die_gemeinsame_prompt_komponente():
    """KI-Parität: Prompt **oder** direkt, wie in allen anderen Modulen."""
    view = (FRONTEND / 'views/compliance-audit/ComplianceAuditView.vue').read_text()
    panel = (FRONTEND / 'components/compliance-audit/WasserstandPanel.vue').read_text()
    for name, txt in (('ComplianceAuditView', view), ('WasserstandPanel', panel)):
        assert '<WizardPromptModal' in txt, \
            f'{name} nutzt nicht die gemeinsame Prompt-Komponente'
    # Der Eigenbau mit rohem KiStreamView ist im Panel verschwunden (Tag, nicht Wort).
    assert '<KiStreamView' not in panel, \
        'WasserstandPanel baut die KI-Ausfuehrung wieder selbst'


def test_bericht_tab_spricht_nicht_mehr_von_gutachten():
    """Im **Bericht-Tab** darf kein Gutachten-Wortlaut stehen.

    Geprüft wird nur dieser Tab. Der Privatgutachten-Assistent spricht
    zu Recht von Gutachten — er erzeugt eines (§ 407a ZPO).
    """
    view = (FRONTEND / 'views/compliance-audit/ComplianceAuditView.vue').read_text()
    start = view.index('<!-- TAB: Auditbericht -->')
    ende = view.index('<!-- ==== DIALOGE ==== -->', start)
    block = view[start:ende]
    sichtbar = re.findall(r'>\s*([^<>{}]*Gutachten[^<>{}]*?)\s*<', block)
    verdaechtig = [t.strip() for t in sichtbar if t.strip()]
    assert not verdaechtig, f'Gutachten-Beschriftungen im Bericht-Tab: {verdaechtig}'
    assert 'Auditbericht' in block, 'der Tab nennt den Bericht nicht beim Namen'


# ── Berichts-Center und Dashboard (#1726) ────────────────────────────────────

def test_berichts_center_katalog_und_erzeugung(client, auth_headers, audit_projekt):
    """Drei Berichtstypen, modulübergreifend gleiche IDs (#1502/#1504)."""
    r = client.get('/api/compliance-audit/projekte/BerichtAudit/berichte',
                   headers=auth_headers)
    assert r.status_code == 200, r.data[:300]
    typen = {b['key']: b['titel'] for b in r.json['reports']}
    assert set(typen) == {'gesamt', 'anforderungen', 'massnahmen'}
    # Im Audit heissen "Anforderungen" fachlich Fragen.
    assert 'Fragen' in typen['anforderungen']

    r = client.post('/api/compliance-audit/projekte/BerichtAudit/berichte/gesamt/generate',
                    json={'format': 'docx'}, headers=auth_headers)
    assert r.status_code == 201, r.data[:300]
    lauf = r.json['id']

    d = client.get(f'/api/compliance-audit/projekte/BerichtAudit/berichte/runs/{lauf}/download',
                   headers=auth_headers)
    assert d.status_code == 200
    assert d.data[:4] == b'PK\x03\x04', 'kein gueltiges DOCX'
    assert len(d.data) > 5000, 'Bericht wirkt leer'


@pytest.mark.parametrize('typ', ['gesamt', 'anforderungen', 'massnahmen'])
def test_jeder_berichtstyp_baut_ohne_ki_entwurf(typ, audit_projekt):
    """Ein Bericht darf nicht von einem KI-Lauf abhängen.

    Ohne Entwurf entsteht er rein aus den Daten; sonst wäre das Berichts-Center
    unbrauchbar, bis jemand die KI bemüht hat.
    """
    from compliance_audit import berichte_provider as prov

    prov.DB_PATH = audit_projekt   # Provider liest modulweit
    daten = prov.render(typ, 'docx', projekt='BerichtAudit')
    assert daten[:4] == b'PK\x03\x04'
    assert len(daten) > 5000


def test_bericht_nutzt_den_berechneten_stand(audit_projekt):
    """Bericht und Wasserstandsmeldung müssen dieselbe Stufe nennen."""
    from docx import Document
    from io import BytesIO

    from compliance_audit import berichte_provider as prov
    from compliance_audit.db import set_erfuellung
    from gutachten.db import load_questions

    prov.DB_PATH = audit_projekt
    q = load_questions(audit_projekt, 'BerichtAudit')[0]
    set_erfuellung(audit_projekt, q['id'], 'erfuellt')

    doc = Document(BytesIO(prov.render('gesamt', 'docx', projekt='BerichtAudit')))
    # Der Stand steht seit #1727 im farbigen Banner, also in einer Tabellenzelle.
    text = '\n'.join([p.text for p in doc.paragraphs] +
                     [z.text for t in doc.tables for r in t.rows for z in r.cells])
    assert 'vollständig compliant' in text, \
        'der Bericht nennt den berechneten Stand nicht'


def test_summary_context_ohne_freitext_antworten(audit_projekt):
    """Die KI-Zusammenfassung bekommt Kennzahlen, nicht die Rohantworten."""
    from compliance_audit import berichte_provider as prov

    prov.DB_PATH = audit_projekt
    ctx = prov.summary_context('BerichtAudit')
    # 'reifegrad'/'offene' ist der Vertrag von shared.reports.summary (#1725).
    assert 'reifegrad' in ctx and 'offene' in ctx
    r = ctx['reifegrad']
    assert 'stufe' in r and 'erfuellungsquote_prozent' in r and 'je_framework' in r
    assert 'antwort' not in json_text(ctx).lower(), \
        'Rohantworten gehoeren nicht in die Zusammenfassung'


def test_summary_prompt_enthaelt_die_kennzahlen(audit_projekt):
    """Ende zu Ende: Was der Provider liefert, muss im Prompt ankommen (#1725).

    Der gemeldete Fehler lag genau in dieser Fuge: Der Kontext war vollständig,
    der gemeinsame Prompt-Bauer verwarf ihn — die KI berichtete daraufhin, es
    lägen keine Daten vor. Ein Test auf ``summary_context`` allein hätte das nie
    bemerkt, deshalb prüft dieser die Kette.
    """
    from compliance_audit import berichte_provider as prov
    from shared.reports.summary import build_summary_prompt

    prov.DB_PATH = audit_projekt
    p = build_summary_prompt('Compliance-Audit', 'BerichtAudit',
                             prov.summary_context('BerichtAudit'))
    assert 'compliant' in p, 'die Compliance-Stufe fehlt im Prompt'
    assert 'Acme' in p, 'das Unternehmen fehlt im Prompt'
    assert 'keine Kennzahlen erfasst' not in p, \
        'Prompt behauptet fehlende Daten, obwohl ein bewertetes Projekt vorliegt'


def json_text(obj) -> str:
    import json as _j
    return _j.dumps(obj, ensure_ascii=False)


def _tab_block(view: str) -> str:
    r"""Der Inhalt von `const tabs = [...]` — mit oder ohne Typannotation.

    #1858: Seit die Liste ihren Typ selbst trägt (`const tabs: { id: TabId;
    label: string }[] = [`), griff das alte Muster `const tabs = \[` nicht
    mehr, und `re.search(...).group(1)` lief in einen AttributeError. Geprüft
    werden soll hier die **Reihenfolge der Tabs**, nicht die Schreibweise der
    Deklaration — ein Test, der an einer Typangabe zerbricht, prüft die
    Erzählung statt das Programm.
    """
    treffer = re.search(r'const tabs\s*(?::[^=]+)?=\s*\[(.*?)\n\]', view, re.DOTALL)
    assert treffer, 'Tab-Katalog `const tabs = [...]` nicht gefunden'
    return treffer.group(1)


def test_view_hat_dashboard_und_berichts_center():
    """Parität zu CRA/NIS2/AI-Act/DSGVO: Dashboard als erster Tab, Berichts-Center."""
    view = (FRONTEND / 'views/compliance-audit/ComplianceAuditView.vue').read_text()
    assert '<ModuleDashboard' in view, 'kein geteiltes Dashboard eingebunden'
    assert '<ReportCenterPanel' in view, 'kein Berichts-Center eingebunden'
    # Dashboard muss der erste Tab sein (wie in den anderen Modulen, #1250).
    tabs = _tab_block(view)
    erster = re.search(r"id:\s*'([^']+)'", tabs).group(1)
    assert erster == 'dashboard', f'erster Tab ist {erster!r}, erwartet dashboard'
    # Das Dashboard speist sich aus dem berechneten Stand, nicht aus einer zweiten Quelle.
    assert '/status' in view and 'dashStand' in view


# ── Ein Berichte-Tab + farbige Berichte (#1727) ──────────────────────────────

def _docx_xml(daten: bytes) -> str:
    import zipfile
    from io import BytesIO
    with zipfile.ZipFile(BytesIO(daten)) as z:
        return z.read('word/document.xml').decode('utf-8', 'replace')


@pytest.mark.parametrize('typ', ['gesamt', 'anforderungen', 'massnahmen'])
def test_berichte_sind_farbig(typ, audit_projekt):
    """Berichte sollen auf einen Blick lesbar sein, nicht nur korrekt.

    Geprüft wird die tatsächliche Schattierung im DOCX (``w:shd``) — python-docx
    kann das nicht über die hohe Schnittstelle, deshalb ist es leicht, es beim
    Umbau zu verlieren.
    """
    import re

    from compliance_audit import berichte_provider as prov

    prov.DB_PATH = audit_projekt
    xml = _docx_xml(prov.render(typ, 'docx', projekt='BerichtAudit'))
    farben = set(re.findall(r'w:fill="([0-9A-Fa-f]{6})"', xml))
    assert prov.FARBEN['kopf'] in farben, 'Tabellenkopf ist nicht farbig'
    # Das Stufen-Band ist in JEDEM Bericht — unabhaengig von der Datenlage.
    stufen_farben = set(prov.FARBEN.values())
    assert farben & stufen_farben, 'kein farbiges Stufen-Band im Bericht'


def test_erfuellung_wird_farblich_unterschieden(audit_projekt):
    """Erfüllt/teilweise/nicht erfüllt müssen sich farblich unterscheiden."""
    import re

    from compliance_audit import berichte_provider as prov
    from compliance_audit.db import set_erfuellung
    from gutachten.db import load_questions, save_questions

    prov.DB_PATH = audit_projekt
    save_questions(audit_projekt, 'BerichtAudit', [
        {'question_num': 1, 'framework': 'NIS2', 'frage': 'A', 'antwort': 'ja'},
        {'question_num': 2, 'framework': 'NIS2', 'frage': 'B', 'antwort': 'teils'},
        {'question_num': 3, 'framework': 'NIS2', 'frage': 'C', 'antwort': 'nein'},
    ])
    qs = load_questions(audit_projekt, 'BerichtAudit')
    for q, e in zip(qs, ['erfuellt', 'teilweise', 'nicht_erfuellt']):
        set_erfuellung(audit_projekt, q['id'], e)

    xml = _docx_xml(prov.render('anforderungen', 'docx', projekt='BerichtAudit'))
    farben = set(re.findall(r'w:fill="([0-9A-Fa-f]{6})"', xml))
    for erf in ('erfuellt', 'teilweise', 'nicht_erfuellt'):
        assert prov.ERF_FARBE[erf] in farben, f'{erf} ist nicht eingefaerbt'


def test_stand_banner_nutzt_die_stufenfarbe(audit_projekt):
    """Das Band oben zeigt die Stufe — in ihrer Farbe."""
    import re

    from compliance_audit import berichte_provider as prov
    from compliance_audit.db import set_erfuellung
    from gutachten.db import load_questions

    prov.DB_PATH = audit_projekt
    for q in load_questions(audit_projekt, 'BerichtAudit'):
        set_erfuellung(audit_projekt, q['id'], 'erfuellt')
    xml = _docx_xml(prov.render('gesamt', 'docx', projekt='BerichtAudit'))
    assert prov.FARBEN['vollstaendig'] in set(
        re.findall(r'w:fill="([0-9A-Fa-f]{6})"', xml)), \
        'Banner zeigt nicht die Farbe der erreichten Stufe'


def test_balken_visualisiert_die_quote():
    from compliance_audit import berichte_provider as prov
    assert prov._balken(100).count('█') == 20
    assert prov._balken(0).count('█') == 0
    assert prov._balken(50).count('█') == 10
    assert prov._balken(None) == '' and prov._balken('x') == ''


def test_nur_ein_berichte_tab():
    """Es gab „Bericht" UND „Berichts-Center" — beide erzeugten einen Auditbericht.

    Eigenständig war daran nur der KI-Entwurf. Zwei Renderer für dasselbe Dokument
    wären zwei Wahrheiten.
    """
    view = (FRONTEND / 'views/compliance-audit/ComplianceAuditView.vue').read_text()
    tabs = _tab_block(view)
    ids = re.findall(r"id:\s*'([^']+)'", tabs)
    assert ids == ['dashboard', 'overview', 'fragen', 'wasserstand', 'berichte'], ids
    assert "activeTab === 'gutachten'" not in view, 'alter Bericht-Tab noch vorhanden'
    # Der KI-Entwurf lebt weiter — im Berichte-Tab als Schritt 1.
    assert 'Auditbericht-Entwurf' in view
    assert '<ReportCenterPanel' in view


def test_kein_zweiter_export_weg_im_frontend():
    """Der frühere Einzel-Export ist durch das Berichts-Center ersetzt."""
    view = (FRONTEND / 'views/compliance-audit/ComplianceAuditView.vue').read_text()
    # Nur echte Aufrufe zaehlen; der Kommentar, der die Ablösung erklaert, bleibt.
    aufrufe = re.findall(r'apiClient\.\w+\(\s*`?[^`\n]*bericht/export', view)
    assert not aufrufe, \
        f'alter Export-Weg noch verdrahtet: {aufrufe} — das Berichts-Center erzeugt den Bericht'


# ── Thema-Spalte der Fragentabelle (#1728) ───────────────────────────────────

def test_thema_spalte_bricht_um_und_wird_gekuerzt():
    """Ein langes Thema darf die Tabelle nicht in die Breite schieben.

    ``max-width`` allein genügt nicht: Ein langes Wort **ohne** Leerzeichen — in
    eingelesenen Frameworks üblich — überschreitet die Breite trotzdem, weil der
    Standard-Umbruch nur an Wortgrenzen bricht. Nötig sind beide Regeln: Umbruch
    an beliebiger Stelle und eine Höhenbegrenzung, damit eine einzelne Zelle die
    Zeile nicht aufbläht.
    """
    view = (FRONTEND / 'views/compliance-audit/ComplianceAuditView.vue').read_text()
    block = re.search(r'\.thema-text\s*\{(.*?)\}', view, re.DOTALL)
    assert block, 'keine Formatierung fuer den Thema-Text'
    css = block.group(1)
    assert 'overflow-wrap: anywhere' in css, \
        'ohne anywhere sprengt ein langes Wort ohne Leerzeichen die Spalte'
    assert '-webkit-line-clamp' in css and 'overflow: hidden' in css, \
        'ohne Hoehenbegrenzung blaeht ein langes Thema die Tabellenzeile auf'


def test_thema_zelle_zeigt_den_vollen_text_als_tooltip():
    """Kürzen darf keine Information verstecken."""
    view = (FRONTEND / 'views/compliance-audit/ComplianceAuditView.vue').read_text()
    zelle = re.search(r'<td class="thema-cell">(.*?)</td>', view, re.DOTALL)
    assert zelle, 'Thema-Zelle nicht gefunden'
    assert ':title="f.thema"' in zelle.group(1), \
        'der vollstaendige Text muss als Tooltip lesbar bleiben'
    assert 'thema-text' in zelle.group(1), 'Clamp greift nur am inneren Element'


def test_thema_wird_nicht_im_markup_abgeschnitten():
    """Die Kürzung gehört in die Darstellung, nicht in die Daten.

    Ein ``slice()`` im Template würde den Text auch für Tooltip, Suche und
    Bearbeiten-Dialog verkürzen — und bei breitem Fenster unnötig kürzen.
    """
    view = (FRONTEND / 'views/compliance-audit/ComplianceAuditView.vue').read_text()
    zelle = re.search(r'<td class="thema-cell">(.*?)</td>', view, re.DOTALL).group(1)
    assert 'slice(' not in zelle and 'substring(' not in zelle, \
        'Thema wird im Markup beschnitten statt per CSS gekuerzt'
