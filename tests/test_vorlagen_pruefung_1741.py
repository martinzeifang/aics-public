"""Vorlagen-Variablen werden gegen das Modul geprüft (#1741).

Der Anlass: Eine Word-Vorlage mit den Variablen des **Compliance-Audit**-Moduls
(``wasserstand``, ``fragen``, ``luecken`` …) lag als **Standard** für ``cra``. CRA
liefert diese Variablen nicht; ``ChainableUndefined`` rendert sie leer, bis in
Zeile 11 ein ``round``-Filter darauf zugreift. Ergebnis: ``TypeError`` → HTTP 500
→ im Frontend „Bericht konnte nicht erzeugt werden". Jeder CRA-Export war kaputt,
bis jemand den Standard entfernte.

Drei Dinge fehlten, und für jedes steht hier ein Test:

1. **Niemand prüfte die Vorlage gegen das Modul.** Die Information lag vor
   (``extract_variables`` + acht Variablenlisten) und wurde nur angezeigt.
2. **Ein Standard wirkt auf jeden Export des Moduls.** Eine unpassende Vorlage
   als Standard ist ein stiller Totalausfall — Hochladen bleibt erlaubt, Standard
   nicht.
3. **Der Fehler war unbrauchbar.** Weder Variable noch Zeile noch „gehört zu
   einem anderen Modul" erreichten den Anwender.

Zum Zuschnitt: Die Prüfung muss auf dem **ersten Pfadsegment** vergleichen.
``extract_variables()`` liefert Top-Level-Namen (``meta``), die Module deklarieren
Blattpfade (``meta.datum``) — ein naiver Mengenvergleich würde in jedem Modul
Fehlalarme melden. Genau das prüft :func:`test_jedes_modul_kennt_seine_eigenen_variablen`
für alle acht Module.

Die DB-abhängigen Tests hängen am ``pg``-Fixture und werden ohne Postgres
übersprungen (tests/conftest.py).
"""
from __future__ import annotations

import io
import uuid
from pathlib import Path

import pytest
from docx import Document

from shared.templates import engine, schema, validate


# ──────────────────────────────────────────────────────────────────────────────
# Hilfsmittel
# ──────────────────────────────────────────────────────────────────────────────
def _vorlage(pfad: Path, zeilen: list[str]) -> Path:
    pfad.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    for z in zeilen:
        d.add_paragraph(z)
    d.save(str(pfad))
    return pfad


def _keys(modul: str) -> list[str]:
    return [v["key"] for v in schema.get_variables(modul)]


@pytest.fixture
def rest(pg, request):
    """Testclient + Admin-Header — aber erst nach dem Postgres-Check.

    ``client``/``auth_headers`` hängen am **session-weiten** ``app``-Fixture, und
    pytest baut höher gescopte Fixtures zuerst auf. Als Parameter deklariert würden
    sie die Flask-App also noch vor dem ``pg``-Skip erzeugen — ohne Datenbank
    ergäbe das einen Fehler statt eines Übersprungs. Daher verzögert geholt.
    """
    return request.getfixturevalue("client"), request.getfixturevalue("auth_headers")


# ──────────────────────────────────────────────────────────────────────────────
# 1 · Der Fall des Anwenders
# ──────────────────────────────────────────────────────────────────────────────
def test_audit_vorlage_unter_cra_wird_als_fremd_erkannt():
    """Die Audit-Variablen gegen ``cra`` — und der Hinweis, wohin sie gehören.

    „Unbekannte Variable" allein hilft nicht weiter; die Antwort, die gefehlt hat,
    ist „``wasserstand`` gehört zu Compliance-Audit".
    """
    erg = validate.pruefe_namen(
        "cra", ["bericht", "fragen", "frameworks", "luecken", "meta", "projekt",
                "wasserstand"])

    assert not erg.ok
    # ``meta`` und ``projekt`` kennt CRA — nur der Rest ist fremd.
    assert set(erg.namen) == {"bericht", "fragen", "frameworks", "luecken", "wasserstand"}
    for eintrag in erg.unbekannt:
        assert "compliance_audit" in eintrag["bekannt_in"], eintrag

    assert "Compliance-Audit" in erg.meldung
    assert "CRA" in erg.meldung
    assert "wasserstand" in erg.meldung


def test_voellig_fremde_variable_hat_keine_herkunft():
    """Ein Tippfehler gehört zu keinem Modul — die Meldung darf nichts erfinden."""
    erg = validate.pruefe_namen("cra", ["projektt"])
    assert not erg.ok
    assert erg.unbekannt[0]["bekannt_in"] == []
    assert "in keinem Modul deklariert" in erg.meldung


# ──────────────────────────────────────────────────────────────────────────────
# 2 · Keine Fehlalarme
# ──────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("modul", list(schema.MODULE))
def test_jedes_modul_kennt_seine_eigenen_variablen(modul):
    """Die eigenen Deklarationen eines Moduls dürfen nie als fremd gelten.

    Geprüft in beiden Schreibweisen, weil beide vorkommen: der deklarierte
    Blattpfad (``meta.datum``, sieben Module) und der Top-Level-Name, den
    ``extract_variables()`` aus der Vorlage liefert (``meta``).
    """
    keys = _keys(modul)
    assert keys, f"{modul} deklariert keine Variablen"

    dotted = validate.pruefe_namen(modul, keys)
    assert dotted.ok, dotted.meldung

    wurzeln = validate.pruefe_namen(modul, [k.split(".", 1)[0] for k in keys])
    assert wurzeln.ok, wurzeln.meldung


def test_indexzugriff_zaehlt_nicht_als_fremde_variable():
    """``{{ risiken[0].name }}`` ist dieselbe Wurzel wie ``risiken``."""
    assert validate.pruefe_namen("cra", ["risiken[0].name", "meta.kapitel"]).ok


def test_cra_deklariert_blattpfade():
    """Punkt 4 des Issues: CRA war das einzige Modul mit Objekt-Deklarationen.

    Damit war es das Modul, für das eine Vorlage am schwersten zu schreiben war —
    und das Modul, in dem der Fehler auftrat. Das Admin-UI kann die Tokens nun
    kopierbar anzeigen wie überall sonst.
    """
    keys = set(_keys("cra"))
    assert {"projekt.name", "projekt.unternehmen", "meta.datum",
            "meta.reifegrad_pct"} <= keys
    # Objekte sind aufgelöst, nicht mehr als ein Eintrag mit Prosa deklariert.
    assert "projekt" not in keys and "meta" not in keys
    # Listen bleiben ein Eintrag (ein Blattpfad je Element wäre ohne Index
    # unbenutzbar) — Muster wie NIS2.
    assert {"bewertungen", "risiken", "dokumente"} <= keys


# ──────────────────────────────────────────────────────────────────────────────
# 3 · Modulschlüssel ai_act ↔ aiact
# ──────────────────────────────────────────────────────────────────────────────
def test_ai_act_alias_wird_aufgeloest():
    """Das Berichts-Center kennt ``ai_act``, die Vorlagen-Engine ``aiact`` (#1735).

    Ohne Auflösung prüfte AI Act gegen ein Modul ohne Variablen — und jede
    Vorlage wäre „fremd".
    """
    assert validate.engine_modul("ai_act") == "aiact"
    erg = validate.pruefe_namen("ai_act", _keys("aiact"))
    assert erg.ok, erg.meldung
    assert erg.modul == "aiact"
    assert validate.deklarierte_wurzeln("ai_act") == validate.deklarierte_wurzeln("aiact")


# ──────────────────────────────────────────────────────────────────────────────
# 4 · Variablen, die der Rahmen beisteuert
# ──────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("modul", list(schema.MODULE))
def test_bruecken_variablen_gelten_in_jedem_modul(modul):
    """``bericht_typ``/``zeitraum`` stehen in keiner Modul-Liste, sind aber immer da.

    Ohne diese Ausnahme würde die Prüfung genau die Vorlagen abweisen, die das
    vorgesehene Verzweigen auf den Berichtstyp nutzen.
    """
    assert validate.pruefe_namen(modul, ["bericht_typ", "zeitraum.von"]).ok


def test_bruecke_liefert_genau_diese_zusatzvariablen(monkeypatch):
    """Hält Ausnahmeliste und Brücke zusammen.

    Legt ``shared/reports/vorlagen.py`` künftig eine weitere Variable in den
    Kontext, fiele sie hier auf — sonst würde sie in Vorlagen als „unbekannt"
    gemeldet, obwohl sie funktioniert.
    """
    from shared.reports import vorlagen as vb

    gesehen: dict = {}

    class _FakeDB:
        DEFAULT_DB_PATH = Path("/dev/null")

        def get_template(self, _db, tid):  # noqa: ANN001
            return {"id": int(tid), "modul": "cra", "name": "X", "version": 1, "aktiv": 1}

    monkeypatch.setattr(vb, "tdb", _FakeDB())
    monkeypatch.setattr(schema, "get_context_builder", lambda m: (lambda db, p: {}))
    monkeypatch.setattr(vb.engine, "render_docx",
                        lambda db, tid, ctx: gesehen.update(ctx) or b"PK")

    vb.render_vorlage("cra", 1, projekt="P", typ="gesamt", fmt="docx")
    assert set(gesehen) == set(validate.ZUSATZ_WURZELN)


# ──────────────────────────────────────────────────────────────────────────────
# 5 · Render-Fehler: Klartext statt 500
# ──────────────────────────────────────────────────────────────────────────────
def test_round_auf_fremder_variable_wird_sprechender_fehler(tmp_path):
    """Exakt der Fall aus dem Log: ``TypeError`` auf ``ChainableUndefined``.

    ``ChainableUndefined`` versprach, fehlende Variablen leer zu rendern. Das gilt
    für ``{{ a.b }}`` und fällt um, sobald ein Filter darauf zugreift — die
    Absicherung trug gerade in den Fällen nicht, für die sie gedacht war.
    """
    pfad = _vorlage(tmp_path / "audit.docx", [
        "Projekt: {{ projekt.name }}",
        "Stand: {{ wasserstand.quote_prozent | round }} %",
    ])

    with pytest.raises(engine.RenderFehler) as fehler:
        engine.render_docx_from_path(pfad, {"projekt": {"name": "X"}}, modul="cra")

    e = fehler.value
    assert "wasserstand" in str(e)
    assert "Compliance-Audit" in str(e)          # wohin die Variable gehört
    assert "CRA" in str(e)                       # wogegen geprüft wurde
    assert e.zeile is not None                   # wo es abbrach
    assert "ChainableUndefined" in e.urspruenglich
    assert [u["name"] for u in e.unbekannt] == ["wasserstand"]


def test_rechnen_auf_fremder_variable_wird_ebenfalls_gefangen(tmp_path):
    """Arithmetik wirft ``UndefinedError`` statt ``TypeError`` — gleicher Anwenderfehler.

    Beide Wege müssen in derselben Meldung landen, sonst bleibt je nach Vorlage ein
    500 übrig.
    """
    pfad = _vorlage(tmp_path / "audit2.docx", ["Fragen: {{ fragen + 1 }}"])
    with pytest.raises(engine.RenderFehler) as fehler:
        engine.render_docx_from_path(pfad, {}, modul="cra")
    assert "fragen" in str(fehler.value)
    assert "UndefinedError" in fehler.value.urspruenglich


def test_length_bricht_nicht_ab_und_liefert_still_null(tmp_path):
    """Festgehalten, weil es die Grenze der Fehlerbehandlung markiert.

    ``{{ x | length }}`` auf einer fehlenden Variable rendert ``0`` statt
    abzubrechen — eine Zahl, die im Compliance-Bericht korrekt aussieht und falsch
    ist. Die Meldung kann das nicht auffangen; nur die Prüfung vor dem Export tut
    es. Läuft dieser Test eines Tages rot (Jinja ändert das Verhalten), ist das
    eine Verbesserung, keine Regression.
    """
    pfad = _vorlage(tmp_path / "laenge.docx", ["Fragen: {{ fragen | length }}"])
    raus = engine.render_docx_from_path(pfad, {}, modul="cra")
    text = "\n".join(p.text for p in Document(io.BytesIO(raus)).paragraphs)
    assert text.strip() == "Fragen: 0"
    # Genau dieser Fall wird deshalb schon beim Standard-Setzen abgewiesen:
    assert not validate.pruefe_variablen("cra", pfad).ok


def test_fehlende_variable_ohne_filter_rendert_weiter_leer(tmp_path):
    """Die Robustheit aus #989 bleibt: nur Filter/Rechnung brechen ab.

    Sonst würde eine Vorlage mit einem optionalen Feld plötzlich scheitern.
    """
    pfad = _vorlage(tmp_path / "optional.docx", ["Notiz: {{ gibtsnicht.feld }}"])
    assert engine.render_docx_from_path(pfad, {}, modul="cra")[:2] == b"PK"


def test_fehler_ohne_modulangabe_nennt_wenigstens_die_zeile(tmp_path):
    """``render_docx_from_path`` ist auch ohne Modul aufrufbar (Tests, Gutachten-Pfad)."""
    pfad = _vorlage(tmp_path / "ohne.docx", ["{{ x.y | round }}"])
    with pytest.raises(engine.RenderFehler) as fehler:
        engine.render_docx_from_path(pfad, {}, modul="")
    assert fehler.value.zeile is not None
    assert fehler.value.unbekannt == []


def test_passende_vorlage_rendert_unveraendert(tmp_path):
    """Gegenprobe: Mit den Variablen des Moduls bleibt alles wie vorher."""
    pfad = _vorlage(tmp_path / "cra.docx", [
        "Projekt: {{ projekt.name }}",
        "Reifegrad: {{ meta.reifegrad_pct | round(1) }} %",
        "Anforderungen: {{ bewertungen | length }}",
    ])
    raus = engine.render_docx_from_path(
        pfad, {"projekt": {"name": "ACME"}, "meta": {"reifegrad_pct": 42.42},
               "bewertungen": [1, 2, 3]}, modul="cra")
    text = "\n".join(p.text for p in Document(io.BytesIO(raus)).paragraphs)
    assert "ACME" in text and "42.4" in text and "3" in text


def test_echter_programmfehler_bleibt_ein_programmfehler(tmp_path, monkeypatch):
    """Nur Jinja-Abbrüche werden übersetzt — sonst verschleierte man Defekte.

    Ein kaputter Konverter oder ein Fehler in docxtpl darf weiter als 500
    durchschlagen; 422 heißt „deine Vorlage", nicht „unser Server".
    """
    pfad = _vorlage(tmp_path / "ok.docx", ["{{ projekt.name }}"])
    import docxtpl

    def _explodiert(self, *a, **k):  # noqa: ANN001
        raise RuntimeError("docxtpl kaputt")

    monkeypatch.setattr(docxtpl.DocxTemplate, "render", _explodiert)
    with pytest.raises(RuntimeError, match="docxtpl kaputt"):
        engine.render_docx_from_path(pfad, {}, modul="cra")


# ──────────────────────────────────────────────────────────────────────────────
# 6 · REST: 422 im Berichts-Center
# ──────────────────────────────────────────────────────────────────────────────
def test_berichts_center_antwortet_422_nicht_500(rest, monkeypatch):
    """Der Anwender sah „Bericht konnte nicht erzeugt werden" und einen 500.

    Ein Vorlagenfehler ist kein Serverfehler: 422 mit der Meldung, die sagt,
    welche Variablen fehlen und welches Modul sie kennt.
    """
    client, auth_headers = rest
    from shared.reports import api as rapi

    fehler = engine.RenderFehler(
        "Die Vorlage passt nicht zum Modul CRA: wasserstand wird dort nicht "
        "geliefert (wasserstand gehört zu Compliance-Audit).",
        modul="cra", unbekannt=[{"name": "wasserstand", "wurzel": "wasserstand",
                                 "bekannt_in": ["compliance_audit"],
                                 "bekannt_in_labels": ["Compliance-Audit"]}],
        zeile=11, urspruenglich="TypeError: …")

    def _raise(*a, **k):
        raise fehler

    monkeypatch.setattr(rapi.vorlagen_bruecke, "render_vorlage", _raise)

    res = client.get("/api/cra/projekte/Egal/berichte/gesamt?template_id=1",
                     headers=auth_headers)
    assert res.status_code == 422, res.get_data(as_text=True)
    body = res.get_json()
    assert "wasserstand" in body["error"]
    assert "Compliance-Audit" in body["error"]
    assert body["zeile"] == 11
    assert body["unbekannt"][0]["name"] == "wasserstand"


def test_generate_antwortet_422_und_protokolliert_den_fehllauf(rest, monkeypatch):
    """Der zweite Weg: „Erzeugen + ablegen" statt Ad-hoc-Download.

    Hier liegt der Render hinter ``core.generate_and_store``, das jeden Fehler
    fängt und nur Text zurückgibt. Ohne den Wrapper wäre die Unterscheidung
    „Vorlage passt nicht" von „Server kaputt" an dieser Stelle verloren.
    """
    client, auth_headers = rest
    from shared.reports import api as rapi

    fehler = engine.RenderFehler("wasserstand gehört zu Compliance-Audit",
                                 modul="cra", zeile=11)

    def _callable(_modul, _tid):
        def _r(*a, **k):
            raise fehler
        return _r

    monkeypatch.setattr(rapi.vorlagen_bruecke, "vorlage_name", lambda *a: "Fremd v1")
    monkeypatch.setattr(rapi.vorlagen_bruecke, "render_callable", _callable)

    res = client.post("/api/cra/projekte/Egal/berichte/gesamt/generate",
                      json={"format": "docx", "template_id": 1}, headers=auth_headers)
    assert res.status_code == 422, res.get_data(as_text=True)
    assert "Compliance-Audit" in res.get_json()["error"]


def test_weg_ohne_template_id_bleibt_unveraendert(rest, monkeypatch):
    """Ohne Vorlage darf die neue Prüfung nichts anfassen (#1735-Invariante)."""
    client, auth_headers = rest
    from shared.reports import api as rapi

    def _nie(*a, **k):
        raise AssertionError("Vorlagen-Pfad ohne template_id betreten")

    monkeypatch.setattr(rapi.vorlagen_bruecke, "render_vorlage", _nie)
    res = client.get("/api/cra/projekte/GibtsNicht/berichte/gesamt", headers=auth_headers)
    # Kein Projekt → programmatischer Renderer scheitert mit 500; entscheidend ist,
    # dass der Vorlagen-Pfad gar nicht betreten wurde.
    assert res.status_code in (200, 500)


# ──────────────────────────────────────────────────────────────────────────────
# 7 · REST: Standard setzen
# ──────────────────────────────────────────────────────────────────────────────
@pytest.fixture
def hochgeladen(rest, tmp_path):
    """Lädt eine Vorlage für ``cra`` hoch und liefert (id, antwort-body)."""
    client, auth_headers = rest

    def _upload(zeilen: list[str]) -> tuple[int, dict]:
        pfad = _vorlage(tmp_path / f"{uuid.uuid4().hex}.docx", zeilen)
        res = client.post(
            "/api/templates",
            data={"file": (io.BytesIO(pfad.read_bytes()), "v.docx"),
                  "modul": "cra", "name": f"pytest-1741-{uuid.uuid4().hex[:8]}"},
            content_type="multipart/form-data", headers=auth_headers)
        assert res.status_code == 201, res.get_data(as_text=True)
        body = res.get_json()
        return int(body["id"]), body

    return _upload


def test_upload_antwort_traegt_das_pruefergebnis(hochgeladen):
    """Punkt 2 des Issues: Das Ergebnis wird zurückgegeben, nicht verworfen."""
    _, fremd = hochgeladen(["{{ wasserstand.quote_prozent }}"])
    assert fremd["pruefung"]["ok"] is False
    assert "wasserstand" in fremd["pruefung"]["meldung"]
    assert "Compliance-Audit" in fremd["pruefung"]["meldung"]

    _, passend = hochgeladen(["{{ projekt.name }}"])
    assert passend["pruefung"]["ok"] is True


def test_standard_setzen_wird_bei_fremden_variablen_abgelehnt(hochgeladen, rest):
    """Ein Standard wirkt auf **jeden** Export des Moduls.

    Hochladen bleibt erlaubt — eine Vorlage darf in Arbeit sein. Zum Standard wird
    sie nicht: Genau so war jeder CRA-Bericht lahmgelegt, bis jemand den Standard
    entfernte.
    """
    client, auth_headers = rest
    tid, _ = hochgeladen(["{{ wasserstand.quote_prozent | round }}"])
    res = client.put(f"/api/templates/{tid}/default", headers=auth_headers)
    assert res.status_code == 409, res.get_data(as_text=True)
    body = res.get_json()
    assert "wasserstand" in body["error"] and "Compliance-Audit" in body["error"]
    assert body["pruefung"]["ok"] is False


def test_standard_setzen_bleibt_fuer_passende_vorlagen_erlaubt(hochgeladen, rest):
    client, auth_headers = rest
    tid, _ = hochgeladen(["{{ projekt.name }}", "{{ meta.reifegrad_pct | round }}"])
    res = client.put(f"/api/templates/{tid}/default", headers=auth_headers)
    assert res.status_code == 200, res.get_data(as_text=True)
    assert res.get_json()["ist_default"] is True


# ──────────────────────────────────────────────────────────────────────────────
# 8 · „Funktioniert in allen Modulen" — der Nachweis
# ──────────────────────────────────────────────────────────────────────────────
#: Schema je Modul anlegen, bevor der Kontext-Builder liest. ``risikobewertung``
#: hat kein ``ensure_db``; dort legt ``_connect`` die Tabellen an.
_SCHEMA_ANLEGEN: dict[str, tuple[str, str]] = {
    "cra": ("cra.db", "ensure_db"),
    "wiba": ("wiba.db", "ensure_db"),
    "nis2": ("nis2.db", "ensure_db"),
    "aiact": ("ai_act.db", "ensure_db"),
    "dsgvo": ("dsgvo.db", "ensure_db"),
    "soc": ("soc.db", "ensure_db"),
    "compliance_audit": ("compliance_audit.db", "ensure_db"),
    "risikobewertung": ("risikobewertung.db", "_connect"),
}


def _schema_anlegen(modul: str) -> Path:
    import importlib

    db_path = schema.MODULE_DB[modul]
    pfad, attr = _SCHEMA_ANLEGEN[modul]
    fn = getattr(importlib.import_module(pfad), attr)
    raus = fn(db_path)
    if raus is not None and hasattr(raus, "close"):  # _connect liefert die Verbindung
        raus.close()
    return db_path


def _aufloesbar(ctx: object, key: str) -> bool:
    """Löst einen deklarierten Blattpfad im echten Kontext auf.

    Endet der Pfad in einer Liste (``meta.kapitel.titel`` gibt es nur je Eintrag),
    ist er ohne Index nicht prüfbar — das gilt als in Ordnung, nicht als Fehler.
    """
    cur = ctx
    for segment in key.split("."):
        if isinstance(cur, list):
            return True
        if not isinstance(cur, dict) or segment not in cur:
            return False
        cur = cur[segment]
    return True


@pytest.mark.parametrize("modul", list(schema.MODULE))
def test_jedes_modul_liefert_jede_deklarierte_variable(pg, modul):
    """Der Nachweis, den es bisher nicht gab (Punkt 5 des Issues).

    Deklaration und Kontext sind zwei Listen, die auseinanderlaufen können: eine
    dokumentierte Variable, die der Builder nicht liefert, ist derselbe Fehler wie
    in #1741 — nur mit dem Admin-UI als Quelle. Dieser Test wird rot, sobald das
    passiert, und zwar für **jedes** Modul.

    Zusätzlich wird eine Vorlage erzeugt, die jede Variable so benutzt, wie es ein
    Anwender täte: Zahlen mit ``| round``, Listen mit ``| length``. Genau diese
    Zugriffe sind es, die ``ChainableUndefined`` nicht abfängt.
    """
    db_path = _schema_anlegen(modul)
    builder = schema.get_context_builder(modul)
    assert builder is not None, f"{modul} ohne Kontext-Builder"

    projekt = f"pytest-1741-{uuid.uuid4().hex[:8]}"
    ctx = dict(builder(db_path, projekt) or {})

    vars_ = schema.get_variables(modul)
    fehlend = [v["key"] for v in vars_ if not _aufloesbar(ctx, v["key"])]
    assert not fehlend, f"{modul} deklariert, liefert aber nicht: {fehlend}"


@pytest.mark.parametrize("modul", list(schema.MODULE))
def test_vorlage_mit_allen_variablen_rendert(pg, modul, tmp_path):
    """Dieselbe Zusage, aber durch die Engine — inklusive Filter."""
    db_path = _schema_anlegen(modul)
    builder = schema.get_context_builder(modul)
    projekt = f"pytest-1741-{uuid.uuid4().hex[:8]}"
    ctx = dict(builder(db_path, projekt) or {})
    # Wie im Export: die Brücke legt diese beiden Variablen dazu.
    ctx.setdefault("bericht_typ", "gesamt")
    ctx.setdefault("zeitraum", {"von": "", "bis": ""})

    zeilen: list[str] = []
    for v in schema.get_variables(modul):
        key, typ = v["key"], str(v.get("typ") or "")
        if typ.startswith("zahl"):
            zeilen.append(f"{key}: {{{{ {key} | round(2) }}}}")
        elif typ.startswith("list"):  # „liste" und das ältere „list"
            zeilen.append(f"{key}: {{{{ {key} | length }}}}")
        else:
            zeilen.append(f"{key}: {{{{ {key} }}}}")

    pfad = _vorlage(tmp_path / f"{modul}.docx", zeilen)
    # Gegenprobe zur Prüfung: die generierte Vorlage muss als passend gelten.
    assert validate.pruefe_variablen(modul, pfad).ok

    raus = engine.render_docx_from_path(pfad, ctx, modul=modul)
    assert raus[:2] == b"PK"
