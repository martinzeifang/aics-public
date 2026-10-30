"""Widersprüche vor dem Ausstellen rechnen, nicht hinterher bemerken (#1885).

Fünf Bemerkungen kamen aus dem Betrieb zurück — als **KI-Bemerkung nach** dem
Dokument. Alle fünf sind rechenbar:

1. „Der Digest ist mit ``sha256:…`` unvollständig und wirkt wie ein Platzhalter"
2. „Widerspruch in der Versionsangabe — der Text nennt Release 8.0.11, Nr. 1
   und der Image-Tag nennen 8.2.3"
3. „Den Satz zur Versionsanzeige nur stehen lassen, wenn die Anzeige existiert"
4. „Zusatzangaben wurden als ``[FEHLT]`` übergeben"
5. „Das Ausstellungsdatum mit dem Freigabedatum des Release abgleichen"

Jeder Test unten trägt den Fall, aus dem er stammt. Zwei Richtungen sind dabei
genauso wichtig wie die Befunde selbst: dass ein **blockierender Fehlalarm**
nicht entsteht (er machte die Anwendung unbenutzbar) und dass der Prüfer nicht
auf seinen **eigenen** Hinweistext anschlägt.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from cra import db as cdb
from cra import doc_export as de
from cra import doc_pruefung as dp
from cra import konformitaet_db as kdb

CRA_DB = Path("data/db/cra.sqlite")
PROJ = "pytest-pruefung-1885"
VER = "8.2.3"
DIGEST = "sha256:" + "cd" * 32

VOLLSTAENDIG = {
    "produkt_name": "Testprodukt", "produkt_typ": "Web",
    "produkt_charge": f"Version {VER}",
    "hersteller": "pytest-firma-1885", "hersteller_adresse": "Weg 1, 00000 Ort",
    "gegenstand": "Testprodukt, Release 8.2.3 – Webanwendung.",
    "ort": "Ort", "datum": "2026-10-30",
    "unterzeichner_name": "A. Person", "unterzeichner_funktion": "Geschäftsführung",
    "keine_harmonisierten": True, "zusatzangaben_keine": True,
}


@pytest.fixture
def projekt():
    cdb.ensure_db(CRA_DB)
    cdb.delete_projekt(CRA_DB, PROJ)
    cdb.save_projekt(CRA_DB, PROJ, unternehmen="pytest-firma-1885",
                     produkt="Testprodukt")
    kdb.save_konformitaet(CRA_DB, PROJ, {"bewertungsweg": "A",
                                         "bewertung_abgeschlossen": True})
    yield PROJ
    cdb.delete_projekt(CRA_DB, PROJ)


# ── 1. Der unvollständige Digest ────────────────────────────────────────────

def test_unvollstaendiger_digest_blockiert():
    b = dp.pruefe({"gegenstand": f"P, Release {VER}. Digest sha256:…"},
                  version=VER, erlaubte_versionen={VER})
    treffer = [x for x in b if x["id"] == "digest_unvollstaendig"]
    assert treffer and treffer[0]["wirkung"] == dp.WIRKUNG_BLOCKIERT


def test_vollstaendiger_digest_ist_kein_befund():
    b = dp.pruefe({"gegenstand": f"P, Release {VER}. Digest {DIGEST}"},
                  version=VER, erlaubte_versionen={VER})
    assert not [x for x in b if x["id"] == "digest_unvollstaendig"]


@pytest.mark.parametrize("text", ["[zu ergänzen]", "[FEHLT]", "TBD", "TODO",
                                   "<…>", "XXX"])
def test_platzhalter_blockieren(text):
    b = dp.pruefe({"gegenstand": f"Produkt {text}"})
    assert [x for x in b if x["id"] == "platzhalter"
            and x["wirkung"] == dp.WIRKUNG_BLOCKIERT]


def test_ein_blockierender_befund_hat_keine_bestaetigung():
    """Es gibt keinen Fall, in dem „sha256:…" richtig ist. Eine Bestätigung
    anzubieten wäre ein Weg um eine nachweislich falsche Aussage herum."""
    b = dp.blockierend(dp.pruefe({"gegenstand": "Digest sha256:…"}))
    assert b and all("confirm" not in x for x in b)


# ── 2. Der Versions-Widerspruch ─────────────────────────────────────────────

def test_der_gemeldete_versions_widerspruch():
    """Der Fall Wort für Wort: Nr. 1 und der Image-Tag nennen 8.2.3, der Text
    von Nr. 4 noch 8.0.11."""
    b = dp.pruefe({"gegenstand": "Produkt, Release 8.0.11 – Webanwendung.",
                   "produkt_charge": "Version 8.2.3"},
                  version=VER, erlaubte_versionen={VER})
    treffer = [x for x in b if x["id"] == "version_widerspruch"]
    assert treffer and treffer[0]["wirkung"] == dp.WIRKUNG_BLOCKIERT
    assert "8.0.11" in treffer[0]["text"] and "8.2.3" in treffer[0]["text"]


def test_passende_version_ist_kein_befund():
    b = dp.pruefe({"gegenstand": f"Produkt, Release {VER}.",
                   "produkt_charge": f"Version {VER}"},
                  version=VER, erlaubte_versionen={VER})
    assert not [x for x in b if x["id"] == "version_widerspruch"]


@pytest.mark.parametrize("text", [
    "Verschlüsselung nach TLS 1.3",
    "Stückliste im Format SPDX 2.3",
    "Managementsystem nach ISO/IEC 27001:2022",
    "Prüfung gegen IEC 62443-4-1",
])
def test_zahlen_in_prosa_sind_keine_versionen(text):
    """**Der Fehlalarm, der die Anwendung unbenutzbar gemacht hätte.** Ein
    Regex auf jede Zahlenfolge hätte „TLS 1.3" als Produktversion gelesen und
    das Ausstellen blockiert. Erkannt wird nur, was der Text als Version
    ausgibt."""
    b = dp.pruefe({"gegenstand": f"Produkt, Release {VER}. {text}"},
                  version=VER, erlaubte_versionen={VER})
    assert not [x for x in b if x["id"] == "version_widerspruch"], text


def test_ohne_vergleichsmassstab_kein_abgleich():
    """Ist in der Akte keine Version bekannt, wäre jede genannte „falsch" —
    und ein Prüfer, der immer anschlägt, wird abgeschaltet."""
    b = dp.pruefe({"gegenstand": "Produkt, Release 9.9.9."})
    assert not [x for x in b if x["id"] == "version_widerspruch"]


def test_v_prefix_zaehlt_als_dieselbe_version():
    """``v8.2.3`` und ``8.2.3`` bezeichnen dasselbe Release; als Zeichenketten
    sind sie verschieden."""
    b = dp.pruefe({"gegenstand": "Produkt v8.2.3."},
                  version=VER, erlaubte_versionen={VER})
    assert not [x for x in b if x["id"] == "version_widerspruch"]


# ── 3. Die unbelegte Aussage ────────────────────────────────────────────────

def test_satz_zur_versionsanzeige_ohne_grundlage_ist_ein_hinweis():
    b = dp.pruefe({"gegenstand": "Die installierte Version wird unter „Über“ "
                                 "angezeigt."},
                  versionsanzeige_erfasst=False)
    treffer = [x for x in b if x["id"] == "versionsanzeige_unbelegt"]
    assert treffer and treffer[0]["wirkung"] == dp.WIRKUNG_HINWEIS


def test_mit_erfasster_anzeige_kein_hinweis():
    b = dp.pruefe({"gegenstand": "Die installierte Version wird angezeigt."},
                  versionsanzeige_erfasst=True)
    assert not [x for x in b if x["id"] == "versionsanzeige_unbelegt"]


# ── 4. Die optionale Angabe ohne Entscheidung ───────────────────────────────

def test_leere_zusatzangaben_ohne_schalter_sind_ein_hinweis():
    b = dp.pruefe({"zusatzangaben": ""})
    treffer = [x for x in b if x["id"] == "zusatzangaben_unbestaetigt"]
    assert treffer and treffer[0]["wirkung"] == dp.WIRKUNG_HINWEIS


def test_der_schalter_beendet_den_hinweis():
    b = dp.pruefe({"zusatzangaben": "", "zusatzangaben_keine": True})
    assert not [x for x in b if x["id"] == "zusatzangaben_unbestaetigt"]


def test_der_prompt_markiert_optionale_felder_nicht_als_fehlend():
    """**Diese Bemerkung hat die Anwendung selbst provoziert.**
    ``_doc_datenblock`` schrieb für *jedes* leere Feld ``[FEHLT]`` — ohne zu
    unterscheiden, ob die Angabe Pflicht ist. ``zusatzangaben`` trägt
    ``daten_pflicht=False``; leer ist dort der Normalfall. Die KI las
    ``[FEHLT]``, meldete es pflichtgemäß, und der Anwender bekam eine Bemerkung
    über ein Feld, das rechtlich nichts verlangt."""
    from cra.ai_wizards import _anhang_v_zeilen

    zeilen, _ = _anhang_v_zeilen(dict(VOLLSTAENDIG, zusatzangaben=""), False)
    nr8 = [z for z in zeilen.splitlines() if z.startswith("8. Zusatzangaben")]
    assert nr8, zeilen
    assert "[FEHLT]" not in nr8[0]
    assert "optional" in nr8[0]


def test_pflichtfelder_werden_weiter_als_fehlend_markiert():
    """Die Gegenrichtung — sonst hätte der Fix die Prüfung entwertet."""
    from cra.ai_wizards import _anhang_v_zeilen

    zeilen, fehlende = _anhang_v_zeilen(dict(VOLLSTAENDIG, ort=""), False)
    assert "[FEHLT]" in zeilen
    assert "ort" in fehlende


# ── 5. Das Datum vor der Freigabe ───────────────────────────────────────────

def test_datum_vor_der_freigabe_verlangt_bestaetigung():
    b = dp.pruefe({"datum": "2026-10-22"}, freigegeben_am="2026-10-25",
                  release_label=VER)
    treffer = [x for x in b if x["id"] == "datum_vor_freigabe"]
    assert treffer
    assert treffer[0]["wirkung"] == dp.WIRKUNG_BESTAETIGUNG
    assert treffer[0]["confirm"] == dp.CONFIRM_DATUM


def test_datum_am_tag_der_freigabe_ist_in_ordnung():
    """Die Erklärung kann Teil des Freigabepakets sein."""
    b = dp.pruefe({"datum": "2026-10-25"}, freigegeben_am="2026-10-25")
    assert not [x for x in b if x["id"] == "datum_vor_freigabe"]


def test_ohne_freigabedatum_kein_abgleich():
    b = dp.pruefe({"datum": "2020-01-01"})
    assert not [x for x in b if x["id"] == "datum_vor_freigabe"]


# ── Der Prüfer schlägt nicht auf sich selbst an ─────────────────────────────

def test_abgeleitete_felder_werden_nicht_geprueft():
    """**Der Fehler war meiner.** ``normen_satz`` trägt bei fehlender Angabe
    absichtlich ``[fehlt — Anhang V Nr. 6 …]``. Die erste Fassung ging über
    *alle* Schlüssel, fand diesen Text und blockierte **jede** Erklärung ohne
    Nr. 6 — der Prüfer schlug auf seinen eigenen Hinweis an (dieselbe Klasse
    wie #1834/#1867/#1878, nur blockierend)."""
    doc = dict(VOLLSTAENDIG, keine_harmonisierten=False, normen=[])
    doc["normen_satz"] = de.NORMEN_FEHLT
    doc["normen_text"] = ""
    b = dp.pruefe(doc, version=VER, erlaubte_versionen={VER})
    assert not dp.blockierend(b), [x["text"] for x in dp.blockierend(b)]


def test_geprueft_wird_eine_positivliste():
    """Eine Negativliste („alles außer normen_satz") bricht beim nächsten
    abgeleiteten Feld wieder."""
    pruefbar = set(dp._pruefbare_felder())
    assert pruefbar >= set(de.EINGABEFELDER)
    assert not (pruefbar & set(de.ABGELEITET))


# ── Nr. 6 und Nr. 8 im Dokument ─────────────────────────────────────────────

def test_nr6_ohne_aussage_ist_ein_hinweis():
    b = dp.pruefe({"normen": [], "keine_harmonisierten": False})
    treffer = [x for x in b if x["id"] == "normen_unbestimmt"]
    assert treffer and treffer[0]["wirkung"] == dp.WIRKUNG_HINWEIS


def test_das_dokument_nennt_nr6_und_nr8(projekt):
    """Rundlauf: Die Aussagen müssen im erzeugten DOCX stehen — behaupten lässt
    sich das nicht aus dem Quelltext ablesen (Lehre aus #1785)."""
    from docx import Document
    from io import BytesIO

    kdb.save_doc(CRA_DB, projekt, dict(VOLLSTAENDIG))
    roh = de.build_docx(projekt, db_path=CRA_DB)
    text = "\n".join(p.text for p in Document(BytesIO(roh)).paragraphs)
    assert "keine harmonisierten Normen" in text
    assert "Keine weiteren Angaben" in text
    assert "—" not in text.split("Angewandte Normen und Spezifikationen:")[1][:80]


def test_kein_gedankenstrich_mehr_bei_nr6(projekt):
    """Der gemeldete Befund: „‚--' wörtlich übernommen"."""
    from docx import Document
    from io import BytesIO

    kdb.save_doc(CRA_DB, projekt, dict(VOLLSTAENDIG, keine_harmonisierten=False,
                                       normen=[]))
    roh = de.build_docx(projekt, db_path=CRA_DB)
    text = "\n".join(p.text for p in Document(BytesIO(roh)).paragraphs)
    zeile = [z for z in text.splitlines()
             if z.startswith("Angewandte Normen und Spezifikationen:")]
    assert zeile and zeile[0].strip() != "Angewandte Normen und Spezifikationen: —"
    assert "Anhang V Nr. 6" in zeile[0]


# ── Das Ausstellen ──────────────────────────────────────────────────────────

def test_ausstellen_bricht_bei_widerspruch_ab(projekt):
    kdb.save_doc(CRA_DB, projekt, dict(
        VOLLSTAENDIG, gegenstand="Produkt, Release 8.0.11 – Digest sha256:…"))
    cdb.save_sbom(CRA_DB, projekt, {"release_version": VER})
    with pytest.raises(dp.ErklaerungUnstimmig):
        de.ausstellen(projekt, db_path=CRA_DB, pdf_bytes=b"%PDF-")


def test_ausstellen_laeuft_bei_widerspruchsfreier_erklaerung(projekt):
    kdb.save_doc(CRA_DB, projekt, dict(VOLLSTAENDIG))
    erg = de.ausstellen(projekt, db_path=CRA_DB, pdf_bytes=b"%PDF-")
    assert erg["sha256"]


def test_issue_doc_meldet_den_widerspruch(projekt):
    """Geprüft wird im **einen** Tor (``_pruefe_anhang_v``) — der REST-Endpunkt
    und jeder andere Aufrufer gehen durch dieselbe Funktion (#1814)."""
    with pytest.raises(dp.ErklaerungUnstimmig):
        kdb.issue_doc(CRA_DB, projekt,
                      dict(VOLLSTAENDIG, produkt_charge="Version 8.2.3",
                           gegenstand="Produkt, Release 1.0.0."),
                      "")


def test_bestaetigung_laesst_den_grenzfall_durch(projekt):
    """Ein echter Grenzfall braucht einen Menschen, keine Sperre."""
    kdb.save_konformitaet(CRA_DB, projekt, {"release_version": VER})
    con = __import__("shared.db", fromlist=["connect"]).connect(str(CRA_DB))
    try:
        con.execute("UPDATE cra_konformitaet SET freigegeben_am='2026-10-25' "
                    "WHERE projekt_name=? AND release_version=?", (projekt, VER))
        con.commit()
    finally:
        con.close()
    doc = dict(VOLLSTAENDIG, datum="2026-10-22")
    with pytest.raises(dp.BestaetigungErforderlich) as exc:
        kdb.issue_doc(CRA_DB, projekt, doc, VER)
    assert dp.CONFIRM_DATUM in exc.value.schluessel
    rec = kdb.issue_doc(CRA_DB, projekt, doc, VER, [dp.CONFIRM_DATUM])
    assert rec["doc_ausgestellt"]


def test_die_pruefung_ist_fail_closed(projekt, monkeypatch):
    """Eine Erklärung, deren Widerspruchsfreiheit niemand bestätigen konnte,
    ist das Gegenteil eines Nachweises (#1751)."""
    kdb.save_doc(CRA_DB, projekt, dict(VOLLSTAENDIG))

    def kaputt(*a, **k):
        raise RuntimeError("Quelle weg")

    monkeypatch.setattr("cra.doc_pruefung.befunde", kaputt)
    with pytest.raises(ValueError, match="Widersprüche"):
        kdb.issue_doc(CRA_DB, projekt, dict(VOLLSTAENDIG), "")


# ── Keine zweite Prüf-Liste ─────────────────────────────────────────────────

def test_die_feldbeschriftungen_kommen_aus_anhang_v():
    """Eine zweite Zuordnung stünde in einer Fehlermeldung — also genau da, wo
    ein Anwender sich auf sie verlässt (#1731/#1735)."""
    quelle = Path("cra/doc_pruefung.py").read_text(encoding="utf-8")
    assert "from cra.doc_export import ANHANG_V" in quelle


def _ohne_kommentare(text: str) -> str:
    """Kommentare und Docstring-Zeilen weg — das Programm prüfen, nicht die
    Erzählung (#1834/#1867/#1878).

    Dieser Wächter wurde beim ersten Anlauf rot, weil der **Kommentar**, der die
    Entfernung der alten Regel begründet, deren Wortlaut zitiert.
    """
    import io
    import tokenize

    raus = []
    for tok in tokenize.generate_tokens(io.StringIO(text).readline):
        if tok.type in (tokenize.COMMENT, tokenize.STRING):
            continue
        raus.append(tok.string)
    return " ".join(raus)


def test_der_wizard_hat_keine_eigenen_plausibilitaetsregeln_mehr():
    """Der alte Hinweis „Keine angewandten Normen angegeben (Annex V Pkt. 6)"
    war **falsch**: ``ANHANG_V`` führt Nr. 6 bewusst nicht als Pflichtangabe."""
    quelle = Path("cra/ai_wizards.py").read_text(encoding="utf-8")
    programm = _ohne_kommentare(quelle)
    assert "Keine angewandten Normen angegeben" not in programm
    assert "doc_pruefung" in programm


def test_parse_uebernimmt_nur_prosa_und_nennt_das_verworfene():
    from cra.ai_wizards import parse_konformitaetserklaerung_response as parse

    erg = parse('{"doc_text": "Entwurf", "felder": {"produkt_typ": "Web", '
                '"hersteller": "Fremd GmbH"}}')
    assert erg["felder"] == {"produkt_typ": "Web"}
    assert erg["verworfen"] == ["hersteller"]
    assert any("Nicht übernommen" in h for h in erg["plausibilitaet_hinweise"])
