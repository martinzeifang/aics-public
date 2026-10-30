"""Sicherheitsupdate und Funktionsupdate trennen (#1793).

Der ganze teure Teil eines Release-Vorgangs hängt an einer Einstufung: Die vier
Neubewertungs-Punkte erscheinen nur bei „wesentliche Änderung = ja". Erwägungs-
grund 39 der Verordnung (EU) 2024/2847 sagt, dass eine Sicherheitsaktualisierung,
die die Zweckbestimmung nicht verändert, keine wesentliche Änderung ist — und
dass die Bündelung mit einem Sicherheitsupdate ein Funktionsupdate **nicht** von
der Einstufung befreit.

Geprüft wird hier beides: dass die Erleichterung ankommt (ein Vorschlag für die
Kernfrage), und dass sie sich nicht missbrauchen lässt (die Erklärung wird gegen
die eigene Akte gerechnet).
"""
from __future__ import annotations

import pytest

from cra import release_art as ra


# ── Vorschlag ───────────────────────────────────────────────────────────────

def test_sicherheitsupdate_schlaegt_nein_vor_und_nennt_den_erwaegungsgrund():
    v = ra.vorschlag("sicherheitsupdate")
    assert v["wert"] == "nein"
    assert "Erwägungsgrund 39" in v["grund"]
    assert "Zweckbestimmung" in v["grund"]


def test_wesentliche_aenderung_schlaegt_ja_vor():
    assert ra.vorschlag("wesentliche_aenderung")["wert"] == "ja"


def test_funktionsupdate_bekommt_bewusst_keinen_vorschlag():
    """Erwägungsgrund 39 macht es von der Änderung abhängig: Eine geringfügige
    Verbesserung ist in der Regel nicht wesentlich, eine neue Funktion schon.
    Ein Vorschlag wäre hier geraten — und geraten sieht in der Akte aus wie
    beurteilt."""
    assert ra.vorschlag("funktionsupdate") == {"wert": "", "grund": ""}
    assert ra.vorschlag("") == {"wert": "", "grund": ""}


def test_entwurf_laesst_die_luecke_stehen():
    """Ein Entwurf, der schon vollständig klingt, wird nicht gelesen."""
    for art in ("sicherheitsupdate", "wesentliche_aenderung"):
        assert "[Hier ergänzen:" in ra.vorschlag(art)["grund"]


# ── Widersprüche ────────────────────────────────────────────────────────────

def test_ohne_erklaerte_art_kein_befund():
    """Bestandsvorgänge führen die Angabe nicht — das ist kein Mangel."""
    assert ra.widersprueche("", wesentliche_aenderung="ja",
                            neue_komponenten=7, behobene_schwachstellen=3) == []


def test_sauberes_sicherheitsupdate_erzeugt_keinen_befund():
    assert ra.widersprueche("sicherheitsupdate", wesentliche_aenderung="nein",
                            neue_komponenten=0, behobene_schwachstellen=4) == []


def test_sicherheitsupdate_mit_ja_ist_ein_widerspruch():
    ids = [w["id"] for w in ra.widersprueche("sicherheitsupdate",
                                             wesentliche_aenderung="ja")]
    assert ids == ["art_gegen_antwort"]


def test_wesentliche_aenderung_mit_nein_ist_ein_widerspruch():
    ids = [w["id"] for w in ra.widersprueche("wesentliche_aenderung",
                                             wesentliche_aenderung="nein")]
    assert ids == ["art_gegen_antwort"]


def test_neue_komponenten_sprechen_gegen_ein_reines_sicherheitsupdate():
    w = ra.widersprueche("sicherheitsupdate", wesentliche_aenderung="nein",
                         neue_komponenten=3)
    assert [x["id"] for x in w] == ["art_neue_komponenten"]
    assert "3" in w[0]["text"]


def test_angehobene_komponenten_allein_sind_kein_befund():
    """Eine angehobene Version **ist** der Normalfall eines Sicherheitsupdates.
    Nur neu aufgenommene Komponenten deuten auf Funktion hin — deshalb bekommt
    die Prüfung gar nicht erst die Zahl der angehobenen zu sehen."""
    assert ra.widersprueche("sicherheitsupdate", wesentliche_aenderung="nein",
                            neue_komponenten=0) == []


def test_funktionsupdate_mit_sicherheitsfix_verweist_auf_die_trennungspflicht():
    """Anhang I Teil II Nr. 2: soweit technisch machbar getrennt bereitstellen."""
    w = ra.widersprueche("funktionsupdate", behobene_schwachstellen=2)
    assert [x["id"] for x in w] == ["art_gebuendelt"]
    assert "Anhang I Teil II Nr. 2" in w[0]["rechtsbezug"]


def test_nicht_geprueft_erzeugt_keinen_befund():
    """``None`` heißt im ganzen Gate „nicht geprüft". Ein Projekt ohne SBOM darf
    keinen Widerspruch gemeldet bekommen, weil nichts zu vergleichen war."""
    assert ra.widersprueche("sicherheitsupdate", wesentliche_aenderung="nein",
                            neue_komponenten=None,
                            behobene_schwachstellen=None) == []
    assert ra.widersprueche("funktionsupdate", behobene_schwachstellen=None) == []


# ── Verdrahtung ─────────────────────────────────────────────────────────────

def test_jede_art_hat_label_und_erklaerung():
    for art in ra.ARTEN:
        assert ra.LABELS.get(art)
        assert ra.ERKLAERUNG.get(art)


def test_unbekannte_art_wird_abgewiesen():
    assert not ra.ist_gueltig("hauptsache_schnell")
    assert ra.ist_gueltig("")  # unbestimmt bleibt zulässig


def test_katalogpunkt_zeigt_auf_den_gerechneten_befund():
    """Die Falle aus #1762/#1774: Ein gerechneter Katalogpunkt braucht den
    flachen Schlüssel im Gate-Ergebnis, sonst meldet der Vorgangsbericht
    „nicht geprüft", obwohl gerechnet wurde."""
    from cra import release_bericht as rb
    from cra import release_checkliste as rc

    eintrag = rc.item("art_konsistent")
    assert eintrag is not None and eintrag.art == "gerechnet"
    assert eintrag.quelle == "release_art_widersprueche"
    assert eintrag.quelle in rb._GERECHNET_AUSWERTUNG


def test_bericht_wertet_den_befund_in_beide_richtungen_aus():
    from cra import release_bericht as rb

    assert rb.gerechnetes_ergebnis(
        "release_art_widersprueche", {"release_art_widersprueche": []})[0] == "erfuellt"
    assert rb.gerechnetes_ergebnis(
        "release_art_widersprueche", {"release_art_widersprueche": ["x"]})[0] == "offen"
    assert rb.gerechnetes_ergebnis(
        "release_art_widersprueche", {"release_art_widersprueche": None})[0] == "ungeprueft"


@pytest.mark.parametrize("art", list(ra.ARTEN))
def test_bericht_druckt_die_art(art):
    """Wer den Vorgang später liest, soll sehen, wovon die Antwort ausgeht."""
    from docx import Document
    from cra import release_bericht as rb

    doc = Document()
    rb._release_art(doc, {"release_art": art})
    text = "\n".join(p.text for p in doc.paragraphs)
    assert ra.LABELS[art] in text


def test_bericht_schweigt_nicht_ueber_eine_fehlende_art():
    from docx import Document
    from cra import release_bericht as rb

    doc = Document()
    rb._release_art(doc, {})
    assert "nicht erklärt" in "\n".join(p.text for p in doc.paragraphs)
