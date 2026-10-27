"""Ein kaufmännisches Und im Namen darf keinen Bericht verhindern (#1791).

Der Befund kam aus dem Betrieb: Für einen Mandanten, dessen Firmenname ein
**kaufmännisches Und** trägt, ließ sich kein einziger Vorlagen-Bericht
erzeugen —

    lxml.etree.XMLSyntaxError: xmlParseEntityRef: no name

Die übrigen Mandanten derselben Installation rendern fehlerfrei. Ursache:
``shared/templates/engine._sandbox_env()`` baute die Jinja-Umgebung mit
``autoescape=False``. docxtpl setzt den gerenderten Wert roh in die DOCX-XML,
und dort ist ein ``&`` der Anfang einer Entität.

Dass es lange niemand merkte, hat einen Grund: Der **programmatische** Export
schreibt den Text über python-docx, das selbst escapt. Erst seit #1735 nimmt
der Anwender im Berichts-Center tatsächlich den Vorlagen-Weg.

Die Tests prüfen beide Jinja-Renderer — den der Module und den des
Gutachten-Moduls, der denselben Defekt trug — und halten fest, dass die
Zusagen daneben bestehen bleiben: fehlende Variablen rendern leer, und
docxtpls Markup-Objekte werden nicht zerschrieben.
"""
from __future__ import annotations

import uuid
from pathlib import Path

import pytest
from docx import Document

from shared.templates.engine import render_docx_from_path

#: Dieselbe Form wie der Name, der den Befund ausgelöst hat: Wortmarke mit
#: Bindestrich, kaufmännisches Und, Rechtsform. Der echte Mandantenname steht
#: bewusst **nicht** im Code — der Baum wird ins öffentliche Repo gespiegelt,
#: und der Sync-Guard lässt Kundennamen dort zu Recht nicht durch.
FIRMA = "muster-it banking & network solutions GmbH"


def _vorlage(pfad: Path, zeilen: list[str]) -> Path:
    pfad.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    for z in zeilen:
        d.add_paragraph(z)
    d.save(str(pfad))
    return pfad


def _text(daten: bytes, tmp_path: Path) -> str:
    ziel = tmp_path / f"{uuid.uuid4().hex}.docx"
    ziel.write_bytes(daten)
    return "\n".join(p.text for p in Document(str(ziel)).paragraphs)


# ── Der Defekt ───────────────────────────────────────────────────────────────

def test_kaufmaennisches_und_im_firmennamen(tmp_path):
    """Vorher: XMLSyntaxError. Der Name muss **lesbar** im Dokument stehen —
    nicht als `&amp;`, denn Word löst die Entität beim Öffnen wieder auf."""
    v = _vorlage(tmp_path / "v.docx", ["Mandant: {{ firma }}"])
    text = _text(render_docx_from_path(v, {"firma": FIRMA}), tmp_path)
    assert text.strip() == f"Mandant: {FIRMA}"


@pytest.mark.parametrize("wert", [
    "A & B",
    "Müller < Meier",
    "x > y",
    'Er sagte "hallo"',
    "Fritz' Firma",
    "A & B < C > D",
    # Was aussieht wie eine Entität, aber keine ist — der Klassiker, an dem
    # halbherziges Escaping scheitert.
    "&amp; wörtlich",
    "&#60; wörtlich",
])
def test_sonderzeichen_kommen_unveraendert_an(tmp_path, wert):
    v = _vorlage(tmp_path / "v.docx", ["{{ w }}"])
    assert _text(render_docx_from_path(v, {"w": wert}), tmp_path).strip() == wert


def test_sonderzeichen_in_einer_schleife(tmp_path):
    """Listen sind der Normalfall in diesen Vorlagen (Anforderungen, Risiken)."""
    v = _vorlage(tmp_path / "v.docx", ["{% for z in zeilen %}{{ z }}|{% endfor %}"])
    text = _text(render_docx_from_path(v, {"zeilen": ["A & B", "C < D"]}), tmp_path)
    assert text.strip() == "A & B|C < D|"


# ── Die Zusagen daneben bleiben ──────────────────────────────────────────────

def test_fehlende_variablen_rendern_weiterhin_leer(tmp_path):
    """`ChainableUndefined` (#1741) — auch verschachtelt."""
    v = _vorlage(tmp_path / "v.docx", ["[{{ gibtsnicht }}]", "[{{ a.b.c }}]"])
    assert _text(render_docx_from_path(v, {}), tmp_path).strip() == "[]\n[]"


def test_docxtpl_markup_wird_nicht_escaped(tmp_path):
    """`RichText` & Co. tragen `__html__`; Jinja übernimmt sie unverändert.

    Ohne diese Eigenschaft hätte `autoescape=True` jede formatierte Einfügung
    in sichtbaren XML-Text verwandelt — der Grund, warum die Umstellung vorher
    zu prüfen war und nicht einfach gesetzt werden durfte.
    """
    from docxtpl import RichText

    rt = RichText("fett & formatiert", bold=True)
    assert hasattr(rt, "__html__")
    # `{{r … }}` ist docxtpls eigene Syntax für einen RichText-Platzhalter —
    # nur so ersetzt es den umgebenden `w:r` statt den Text zu verschachteln.
    v = _vorlage(tmp_path / "v.docx", ["{{r inhalt }}"])
    text = _text(render_docx_from_path(v, {"inhalt": rt}), tmp_path)
    assert text.strip() == "fett & formatiert"
    assert "w:r" not in text and "&amp;" not in text


def test_zahlen_und_none_bleiben_unauffaellig(tmp_path):
    v = _vorlage(tmp_path / "v.docx", ["{{ n }}|{{ f }}|{{ b }}"])
    text = _text(render_docx_from_path(v, {"n": 42, "f": 3.5, "b": True}), tmp_path)
    assert text.strip() == "42|3.5|True"


# ── Derselbe Defekt im Gutachten-Modul ───────────────────────────────────────

def test_gutachten_vorlage_vertraegt_das_und(tmp_path):
    """`gutachten/template_render.render_with_template` baute die Sandbox ohne
    `autoescape` — Jinjas Vorgabe ist aus. Ein Auftraggeber mit `&` im Namen
    hätte dort denselben Abbruch erzeugt."""
    from gutachten.template_render import render_with_template

    v = _vorlage(tmp_path / "g.docx", ["Auftraggeber: {{ projekt.auftraggeber }}"])
    ziel = tmp_path / "out.docx"
    render_with_template(v, {"projekt": {"auftraggeber": FIRMA}}, ziel)
    text = "\n".join(p.text for p in Document(str(ziel)).paragraphs)
    assert text.strip() == f"Auftraggeber: {FIRMA}"


# ── Wächter: keine weitere Render-Umgebung ohne autoescape ───────────────────

def test_keine_jinja_umgebung_ohne_autoescape():
    """Beide Fundstellen hatten dieselbe Ursache an verschiedenen Orten. Dieser
    Test wird rot, wenn eine dritte dazukommt — eine DOCX-Vorlage ohne
    Escaping ist kein Stilfehler, sondern ein Export, der für bestimmte
    Mandanten gar nicht läuft."""
    import re

    verdaechtig: list[str] = []
    for pfad in (Path("shared/templates/engine.py"),
                 Path("gutachten/template_render.py")):
        quelle = pfad.read_text(encoding="utf-8")
        for treffer in re.finditer(r"SandboxedEnvironment\(([^)]*)\)", quelle):
            args = treffer.group(1)
            if "autoescape=True" not in args:
                verdaechtig.append(f"{pfad}: SandboxedEnvironment({args.strip()})")
    assert not verdaechtig, "\n".join(verdaechtig)
