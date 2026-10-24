"""EU-Konformitätserklärung als Dokument mit eigener Word-Vorlage (#1753).

Geprüft wird nicht das Layout, sondern die Aussagen, auf die sich später jemand
beruft:

* Der Typ ist am geteilten Berichts-Center registriert — sonst gibt es DOCX,
  PDF, Historie und Word-Vorlage nicht, und es wäre ein Sonderweg daneben
  entstanden (#1735/#1748).
* Eine Vorlage, der Pflichtangaben aus Anhang V fehlen, lässt sich **nicht als
  Standard setzen**, und die Meldung nennt die Nummer. Bei einem Bericht ist
  eine freie Vorlage Geschmacksfrage, bei der Konformitätserklärung ein
  Rechtsrisiko.
* Unvollständige Pflichtfelder ergeben einen **Entwurf** mit benannten Lücken.
  Eine Erklärung, in der „Ort und Datum“ fehlt, sieht aus wie eine Erklärung
  und ist keine — stillschweigend leer drucken ist der schlimmere Fehler
  (#1741/#1751).
* Die ausgestellte Fassung liegt mit SHA-256 unveränderlich im Archiv und
  ändert sich danach nicht mehr.
* Jede deklarierte ``doc.*``-Variable wird vom Kontext geliefert, auch ohne
  DoC-Datensatz: leere Zeichenketten, nie ``None`` — Jinja druckt ``None``
  wörtlich.
"""
from __future__ import annotations

import io
import uuid
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from cra import berichte_provider as bp
from cra import doc_export as de
from shared.templates import validate

PROJ = 'pytest-doc-vorlage-1753'
DB = Path('data/db/cra.sqlite')   # Schema 'cra' — dieselbe Quelle wie das Modul

#: Vollständige Angaben nach Anhang V (ohne die abgeleiteten Gesetzesformeln).
VOLL = {
    'produkt_name': 'AI Compliance Suite',
    'produkt_typ': 'AICS-Server',
    'produkt_charge': 'Version 7.4.8',
    'hersteller': 'Cyberwoks GmbH',
    'hersteller_adresse': 'Musterweg 1, 12345 Musterstadt, Deutschland',
    'gegenstand': 'Webanwendung zur Verwaltung von Compliance-Nachweisen.',
    'normen': ['EN 18031-1', 'ETSI EN 303 645'],
    'ort': 'Musterstadt',
    'datum': '2026-10-22',
    'unterzeichner_name': 'Martin Zeifang',
    'unterzeichner_funktion': 'Geschäftsführer',
}


def _text(docx_bytes: bytes) -> str:
    """Fließtext **und** Tabellenzellen — die Kopfangaben stehen in einer Tabelle."""
    doc = Document(BytesIO(docx_bytes))
    teile = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            teile.extend(c.text for c in row.cells)
    return '\n'.join(teile)


def _pdf_moeglich() -> bool:
    """Lokal ist in der Regel kein Konverter da — dann wird übersprungen."""
    try:
        from shared.templates import pdf_converter as pc
        return bool(pc.is_soffice_available() or pc.is_gotenberg_available())
    except Exception:  # noqa: BLE001
        return False


def _vorlage(pfad: Path, zeilen: list[str]) -> Path:
    pfad.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    for z in zeilen:
        d.add_paragraph(z)
    d.save(str(pfad))
    return pfad


#: Zeilen einer vollständigen Erklärungs-Vorlage.
VORLAGE_VOLL = [
    '{{ doc.produkt_name }} / {{ doc.produkt_typ }} / {{ doc.produkt_charge }}',
    '{{ doc.hersteller }}, {{ doc.hersteller_adresse }}',
    '{{ doc.verantwortungssatz }}',
    '{{ doc.gegenstand }}',
    '{{ doc.konformitaetsaussage }}',
    'Normen: {{ doc.normen_text }}',
    '{{ doc.ort }}, {{ doc.datum }}',
    '{{ doc.unterzeichner_name }} — {{ doc.unterzeichner_funktion }}',
    'Typ: {{ bericht_typ }}',
]


@pytest.fixture
def entwurf():
    """Projekt mit angefangener Erklärung: Ort, Datum und Unterzeichner fehlen."""
    from cra import konformitaet_db as kdb
    from cra.db import save_projekt

    save_projekt(DB, PROJ, unternehmen='Cyberwoks GmbH',
                 produkt='AI Compliance Suite')
    kdb.save_konformitaet(DB, PROJ, {'bewertungsweg': 'A',
                                     'bewertung_abgeschlossen': True})
    luecke = {k: v for k, v in VOLL.items()
              if k not in ('ort', 'datum', 'unterzeichner_name',
                           'unterzeichner_funktion')}
    kdb.issue_doc(DB, PROJ, luecke)
    return DB


@pytest.fixture
def vollstaendig():
    """Dasselbe Projekt mit allen Anhang-V-Angaben, Erklärung ausgestellt."""
    from cra import konformitaet_db as kdb
    from cra.db import save_projekt

    save_projekt(DB, PROJ, unternehmen='Cyberwoks GmbH',
                 produkt='AI Compliance Suite')
    kdb.save_konformitaet(DB, PROJ, {'bewertungsweg': 'A',
                                     'bewertung_abgeschlossen': True})
    kdb.issue_doc(DB, PROJ, dict(VOLL))
    return DB


# ── 1 · Registrierung am Berichts-Center ────────────────────────────────────
def test_typ_steht_im_katalog():
    """Ohne Katalog-Eintrag weist die Export-Route den Typ mit 404 ab."""
    spec = next((s for s in bp.catalog() if s.key == 'konformitaetserklaerung'), None)
    assert spec is not None, 'Berichtstyp fehlt'
    assert spec.titel and spec.norm, 'Titel und Rechtsbezug gehören in die Auswahl'
    assert 'Anhang V' in spec.norm


def test_typ_erscheint_im_berichts_center(client, auth_headers, entwurf):
    """Der Weg, den der Anwender nimmt: die Typ-Liste des Moduls."""
    res = client.get(f'/api/cra/projekte/{PROJ}/berichte', headers=auth_headers)
    assert res.status_code == 200, res.json
    assert 'konformitaetserklaerung' in [r['key'] for r in res.json['reports']]


def test_docx_entsteht_und_enthaelt_die_pflichtangaben(vollstaendig):
    daten = bp.render('konformitaetserklaerung', 'docx', projekt=PROJ)
    assert daten[:2] == b'PK' and len(daten) > 5000
    text = _text(daten)
    assert 'EU-Konformitätserklärung' in text
    # Jede erfasste Pflichtangabe muss im Dokument stehen — sonst ist es keine.
    for wert in ('AICS-Server', 'Version 7.4.8', 'Cyberwoks GmbH',
                 'Musterweg 1', 'Musterstadt', '2026-10-22',
                 'Martin Zeifang', 'Geschäftsführer'):
        assert wert in text, f'{wert} fehlt im Dokument'
    # Nr. 3 ist Festtext des Renderers, nicht Formulareingabe.
    assert de.VERANTWORTUNGSSATZ in text
    # Nr. 5 nennt die Verordnung.
    assert '2024/2847' in text


def test_ohne_jede_erfassung_entsteht_ein_entwurf():
    """Ein Projekt ohne Erklärung darf keinen Abbruch erzeugen: Den Entwurf
    braucht man, um die Lücken überhaupt zu sehen."""
    text = _text(de.build_docx(PROJ + '-leer', db_path=DB))
    assert de.ENTWURF_MARKER in text
    assert 'Nr. 8: Ort und Datum' in text


@pytest.mark.skipif(not _pdf_moeglich(), reason='kein PDF-Konverter verfügbar')
def test_pdf_laeuft_ueber_denselben_docx_weg(vollstaendig):
    daten = bp.render('konformitaetserklaerung', 'pdf', projekt=PROJ)
    assert daten[:4] == b'%PDF'


# ── 2 · Word-Vorlage ────────────────────────────────────────────────────────
def test_vorlagen_weg_laeuft_und_kennt_den_berichtstyp(vollstaendig, tmp_path,
                                                       monkeypatch):
    """Mit Vorlage rendert die Engine, und ``bericht_typ`` steht im Kontext.

    Nur dadurch kann eine Vorlage auf das Dokument verzweigen, statt dass die
    Engine es ihr vorgibt (#1735).
    """
    from shared.reports import vorlagen as vb

    pfad = _vorlage(tmp_path / 'doc.docx', VORLAGE_VOLL)

    class _FakeDB:
        DEFAULT_DB_PATH = Path('/dev/null')

        def get_template(self, _db, tid):  # noqa: ANN001
            return {'id': int(tid), 'modul': 'cra', 'name': 'DoC', 'version': 1,
                    'aktiv': 1, 'datei_pfad': str(pfad)}

    gesehen: dict = {}

    # ``engine.render_docx`` liest die Registry selbst; die Vorlage liegt hier
    # im tmp_path. Gerendert wird deshalb über den Pfad — die echte Engine mit
    # dem echten Kontext, nur ohne Registry-Eintrag.
    def _render(_db, _tid, ctx):  # noqa: ANN001
        gesehen.update(ctx)
        from shared.templates.engine import render_docx_from_path
        return render_docx_from_path(pfad, ctx, modul='cra')

    monkeypatch.setattr(vb, 'tdb', _FakeDB())
    monkeypatch.setattr(vb.engine, 'render_docx', _render)
    raus = vb.render_vorlage('cra', 1, projekt=PROJ, typ='konformitaetserklaerung',
                             fmt='docx')
    assert raus[:2] == b'PK'
    assert gesehen['bericht_typ'] == 'konformitaetserklaerung'
    text = _text(raus)
    assert 'Geschäftsführer' in text and 'konformitaetserklaerung' in text


def test_vollstaendige_vorlage_gilt_als_passend(tmp_path):
    pfad = _vorlage(tmp_path / 'voll.docx', VORLAGE_VOLL)
    erg = validate.pruefe_variablen('cra', pfad)
    assert erg.ok, erg.meldung
    assert erg.fehlende_pflicht == []


def test_vorlage_ohne_pflichtvariable_nennt_die_anhang_v_nummer(tmp_path):
    """Die Meldung muss die Nummer nennen — sie sagt, was fehlt und warum es
    nicht verhandelbar ist."""
    ohne_ort = [z for z in VORLAGE_VOLL if 'doc.ort' not in z]
    erg = validate.pruefe_variablen('cra', _vorlage(tmp_path / 'ohne.docx', ohne_ort))

    assert not erg.ok
    assert erg.unbekannt == [], 'die Variablen passen zum Modul — es fehlt eine Angabe'
    assert [(f['nr'], f['bezeichnung']) for f in erg.fehlende_pflicht] == \
        [(8, 'Ort und Datum')]
    assert 'Nr. 8: Ort und Datum fehlt' in erg.meldung
    assert 'Anhang V' in erg.meldung


def test_gesetzesformel_darf_woertlich_in_der_vorlage_stehen(tmp_path):
    """Nr. 3 und Nr. 5 sind Gesetzeswortlaut. Eine Vorlage, die sie ausschreibt,
    statt die Variable zu benutzen, ist vollständig — eine reine
    Variablenprüfung würde sie zu Unrecht abweisen."""
    zeilen = [z for z in VORLAGE_VOLL
              if 'verantwortungssatz' not in z and 'konformitaetsaussage' not in z]
    zeilen += ['Diese Erklärung wird unter der alleinigen Verantwortung des '
               'Herstellers ausgestellt.',
               'Der Gegenstand erfüllt die Verordnung (EU) 2024/2847.']
    erg = validate.pruefe_variablen('cra', _vorlage(tmp_path / 'text.docx', zeilen))
    assert erg.ok, erg.meldung


def test_bericht_ohne_doc_variablen_bleibt_unberuehrt(tmp_path):
    """Die Pflichtprüfung gilt nur für die Erklärung. Ein Gesamtbericht benutzt
    ``doc.*`` nicht und darf von ihr nichts merken."""
    pfad = _vorlage(tmp_path / 'bericht.docx', [
        '{{ projekt.name }}', '{{ meta.reifegrad_pct | round(1) }}',
        '{{ bewertungen | length }}'])
    erg = validate.pruefe_variablen('cra', pfad)
    assert erg.ok and erg.fehlende_pflicht == []


def test_andere_module_haben_keine_pflichtvorlagen():
    """Nur CRA kennt dieses Dokument — kein anderes Modul darf plötzlich
    Pflichtvariablen bekommen."""
    from shared.templates import schema

    for modul in schema.MODULE:
        erwartet = modul == 'cra'
        assert bool(validate.pflicht_vorlagen(modul)) is erwartet, modul


# ── 3 · Standard setzen (der Weg des Anwenders) ─────────────────────────────
@pytest.fixture
def hochgeladen(pg, request, tmp_path):
    """Lädt eine Vorlage für ``cra`` hoch und liefert (id, Antwort-Body).

    ``client``/``auth_headers`` hängen am session-weiten ``app``-Fixture und
    werden deshalb erst nach dem Postgres-Check geholt (Muster aus #1741).
    """
    client = request.getfixturevalue('client')
    auth_headers = request.getfixturevalue('auth_headers')

    def _upload(zeilen: list[str]) -> tuple[int, dict, object, dict]:
        pfad = _vorlage(tmp_path / f'{uuid.uuid4().hex}.docx', zeilen)
        res = client.post(
            '/api/templates',
            data={'file': (io.BytesIO(pfad.read_bytes()), 'v.docx'),
                  'modul': 'cra', 'name': f'pytest-1753-{uuid.uuid4().hex[:8]}'},
            content_type='multipart/form-data', headers=auth_headers)
        assert res.status_code == 201, res.get_data(as_text=True)
        body = res.get_json()
        return int(body['id']), body, client, auth_headers

    return _upload


def test_standard_setzen_wird_bei_fehlender_pflichtangabe_abgelehnt(hochgeladen):
    """Ein Standard wirkt auf **jeden** Export des Moduls. Eine Erklärungs-
    Vorlage ohne „Ort und Datum“ erzeugt ein Dokument, das aussieht wie eine
    Erklärung und keine ist — Hochladen bleibt erlaubt, Standard nicht."""
    ohne_ort = [z for z in VORLAGE_VOLL if 'doc.ort' not in z]
    tid, body, client, auth_headers = hochgeladen(ohne_ort)
    assert body['pruefung']['ok'] is False
    assert 'Nr. 8: Ort und Datum' in body['pruefung']['meldung']

    res = client.put(f'/api/templates/{tid}/default', headers=auth_headers)
    assert res.status_code == 409, res.get_data(as_text=True)
    assert 'Nr. 8: Ort und Datum fehlt' in res.get_json()['error']


def test_standard_setzen_bleibt_fuer_vollstaendige_vorlage_erlaubt(hochgeladen):
    tid, body, client, auth_headers = hochgeladen(VORLAGE_VOLL)
    assert body['pruefung']['ok'] is True, body['pruefung']['meldung']
    res = client.put(f'/api/templates/{tid}/default', headers=auth_headers)
    assert res.status_code == 200, res.get_data(as_text=True)


# ── 4 · Leere Pflichtfelder brechen die ausgestellte Fassung ────────────────
def test_unvollstaendig_gibt_entwurf_mit_benannten_luecken(entwurf):
    text = _text(de.build_docx(PROJ, db_path=DB))
    assert de.ENTWURF_MARKER in text
    assert 'ENTWURF' in text
    assert 'Nr. 8: Ort und Datum' in text
    assert 'Nr. 8: Unterzeichner mit Name und Funktion' in text
    # Was erfasst ist, steht trotzdem drin — der Entwurf dient dem Nacharbeiten.
    assert 'AICS-Server' in text


def test_ausgestellter_datensatz_mit_luecken_bleibt_entwurf(entwurf):
    """``doc_ausgestellt`` allein macht keine Erklärung.

    Der Datensatz kann vor dieser Erweiterung (oder über die API) als
    ausgestellt markiert worden sein, obwohl Anhang-V-Angaben fehlen. Das
    Dokument darf dann nicht als ausgestellte Fassung erscheinen.
    """
    d = de.doc_daten(DB, PROJ)
    assert d['doc_ausgestellt'] is True, 'Vorbedingung: Datensatz ist ausgestellt'
    assert d['ist_entwurf'] is True
    assert [f['bezeichnung'] for f in d['fehlend']] == \
        ['Ort und Datum', 'Unterzeichner mit Name und Funktion']


def test_ausgestellte_fassung_verlangt_vollstaendigkeit(entwurf):
    with pytest.raises(de.PflichtfelderFehlen) as fehler:
        de.build_docx(PROJ, db_path=DB, nur_ausgestellt=True)
    assert 'Nr. 8: Ort und Datum' in str(fehler.value)

    with pytest.raises(de.PflichtfelderFehlen) as zweiter:
        de.ausstellen(PROJ, db_path=DB, pdf_bytes=b'%PDF-1.7 egal')
    assert [f['nr'] for f in zweiter.value.fehlend] == [8, 8]


def test_vollstaendige_angaben_tragen_kein_entwurfskennzeichen(vollstaendig):
    d = de.doc_daten(DB, PROJ)
    assert d['fehlend'] == [] and d['ist_entwurf'] is False
    text = _text(de.build_docx(PROJ, db_path=DB, nur_ausgestellt=True))
    assert de.ENTWURF_MARKER not in text
    assert 'Ausgestellte Fassung' in text


def test_notifizierte_stelle_nur_wenn_der_weg_sie_einschliesst(vollstaendig):
    """Anhang V Nr. 7 gilt „soweit zutreffend“.

    Bei Modul A ist niemand beteiligt — die Angabe zu verlangen würde jede
    Selbstbewertung blockieren. Bei B+C ist sie Pflicht.
    """
    from cra import konformitaet_db as kdb

    assert de.doc_daten(DB, PROJ)['fehlend'] == []

    # #1759: „abgeschlossen" verlangt jetzt die notifizierte Stelle. Für die
    # Aussage dieses Tests (Nr. 7 fehlt → Entwurf) ist der Haken nicht nötig —
    # er wird deshalb **ausdrücklich** gelöst.
    #
    # #1805: Vorher stand hier nur ``{'bewertungsweg': 'B+C'}``, und der Haken
    # fiel von selbst weg, weil jedes Speichern alle nicht genannten Felder
    # leerte. Seit das ein Teil-Update ist, muss gesagt werden, was gelten soll.
    # Dass dieser Test dadurch rot wurde, ist die Probe aufs Exempel: Genau so
    # verlor das Konformitäts-Formular bei jedem Speichern die `notizen`.
    kdb.save_konformitaet(DB, PROJ, {'bewertungsweg': 'B+C',
                                     'bewertung_abgeschlossen': False})
    d = de.doc_daten(DB, PROJ)
    assert [f['nr'] for f in d['fehlend']] == [7]
    assert d['ist_entwurf'] is True


# ── 5 · Ausgestellte Fassung unveränderlich ─────────────────────────────────
def test_ausgestellte_fassung_liegt_mit_sha256_im_archiv(vollstaendig):
    import hashlib

    pdf = b'%PDF-1.7 ausgestellte Fassung'
    rec = de.ausstellen(PROJ, von='gf@example.com', db_path=DB, pdf_bytes=pdf)

    assert rec['sha256'] == hashlib.sha256(pdf).hexdigest()
    pfad = Path(rec['pdf_path'])
    assert pfad.is_file() and pfad.read_bytes() == pdf
    assert rec['sha256'][:12] in pfad.name, 'der Hash gehört in den Dateinamen'

    eintraege = de.list_archiv(DB, PROJ)
    assert eintraege and eintraege[0]['sha256'] == rec['sha256']
    assert eintraege[0]['ausgestellt_von'] == 'gf@example.com'
    assert de.get_archiv(DB, PROJ, eintraege[0]['id'])['pdf_path'] == str(pfad)


def test_archiv_wird_nie_ueberschrieben(vollstaendig):
    """Append-only: Eine zweite Ausstellung legt eine zweite Datei an und lässt
    die erste unberührt. Die Erklärung ist das Stück, auf das sich Marktaufsicht
    und Kunde berufen."""
    erst = de.ausstellen(PROJ, von='gf@example.com', db_path=DB,
                         pdf_bytes=b'%PDF-1.7 Fassung 1')
    zweit = de.ausstellen(PROJ, von='gf@example.com', db_path=DB,
                          pdf_bytes=b'%PDF-1.7 Fassung 2')

    assert erst['pdf_path'] != zweit['pdf_path']
    assert Path(erst['pdf_path']).read_bytes() == b'%PDF-1.7 Fassung 1'
    assert len(de.list_archiv(DB, PROJ)) >= 2


def test_archiv_ist_eine_eigene_tabelle(vollstaendig):
    """Nicht ``cra_konformitaet_final``: Dessen Zeilen erscheinen als
    „Archivierte Nachweise“ im Release-Vorgangsbericht und unter
    ``/nachweis/archiv``. Eine Erklärung dort hineinzulegen hieße, sie als
    Nachweis auszuweisen und unter falschem Namen zum Download anzubieten."""
    from cra import konformitaet_final_db as fdb

    de.ausstellen(PROJ, von='gf@example.com', db_path=DB,
                  pdf_bytes=b'%PDF-1.7 Erklaerung')
    assert fdb.list_final(DB, PROJ) == [], \
        'die Erklärung darf nicht im Nachweis-Archiv auftauchen'


def test_archiv_tabelle_ist_beim_loeschen_zugeordnet():
    """Sonst bliebe nach dem Löschen eines Mandanten eine PDF ohne Zeile auf
    der Platte — für die Anwendung unsichtbar (#1752)."""
    from cra.db import PROJEKT_TABELLEN

    assert 'cra_doc_final' in {t for t, _ in PROJEKT_TABELLEN}


# ── 6 · Kontext: jede deklarierte Variable kommt an ─────────────────────────
def _aufloesbar(ctx: object, key: str) -> bool:
    cur = ctx
    for segment in key.split('.'):
        if isinstance(cur, list):
            return True
        if not isinstance(cur, dict) or segment not in cur:
            return False
        cur = cur[segment]
    return True


def test_kontext_liefert_jede_deklarierte_doc_variable():
    """Auch für ein Projekt **ohne** Erklärung: leere Zeichenketten, keine
    ``None`` — Jinja druckt ``None`` wörtlich (#1741)."""
    from cra.template_context import CRA_VARIABLES, build_cra_context

    ctx = build_cra_context(DB, f'pytest-1753-{uuid.uuid4().hex[:8]}')
    keys = [v['key'] for v in CRA_VARIABLES if v['key'].startswith('doc.')]
    assert len(keys) >= 20, 'die Anhang-V-Variablen fehlen in CRA_VARIABLES'

    fehlend = [k for k in keys if not _aufloesbar(ctx, k)]
    assert not fehlend, f'deklariert, aber nicht geliefert: {fehlend}'
    assert all(v is not None for v in ctx['doc'].values())


def test_jede_pflichtvariable_ist_auch_deklariert():
    """Eine Pflichtvariable, die das Modul nicht deklariert, wäre eine Vorlage,
    die niemand schreiben kann: Das Admin-UI zeigt die Tokens aus
    ``CRA_VARIABLES``."""
    from cra.template_context import CRA_VARIABLES

    deklariert = {v['key'] for v in CRA_VARIABLES}
    for spec in validate.pflicht_vorlagen('cra'):
        for angabe in spec['angaben']:
            for token in angabe['tokens']:
                assert token in deklariert, token
