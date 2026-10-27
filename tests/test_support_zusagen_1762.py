"""Unterstützungszeitraum: Zusagen, Mindestdauer, Kaufhinweis (#1762, Punkte 1-4).

Vier Lücken, die alle am selben Datensatz hängen (``cra_support_period``):

1. Die Zusagen „Sicherheitsupdates kostenlos" und „getrennt von
   Funktionsupdates" standen nur als Hinweistext im Anforderungskatalog;
   ``update_kanal`` ist **ein** Freitextfeld und kann zwei Kanäle nicht
   abbilden.
2. Ein Support-Zeitraum unter fünf Jahren wurde stillschweigend akzeptiert,
   obwohl die Ausnahme eine Begründung braucht.
3. Der Kaufhinweis (Support-Ende vor dem Kauf sichtbar) war ein ungeprüftes
   Freitextfeld ohne Prüfpunkt.
4. Die Aufbewahrungsfristen wurden gerechnet, aber nicht überwacht.

Die Prüfungen liegen **serverseitig** (``cra.db.pruefe_support_period``), nicht
im Formular: Das Feld wird auch vom Autodetect und von Wizards beschrieben, und
eine Prüfung, die nur im Browser steht, ist bei einem Audit kein Argument.
"""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

import pytest

from cra import db as cdb

CRA_DB = Path("data/db/cra.sqlite")
PROJ = "pytest-zusagen-1762"
VERSION = "1.2.3"


@pytest.fixture(autouse=True)
def _full_license():
    from server import license_state
    cur = license_state._current
    prev = (cur.state, list(cur.modules))
    cur.state, cur.modules = "ok", ["*"]
    yield
    cur.state, cur.modules = prev[0], prev[1]


@pytest.fixture
def projekt():
    """Projekt mit leerem Support-Datensatz.

    Unter Postgres ist ein SQLite-Pfad nur ein Schema-Selektor (#1332): Zeilen
    überleben den Lauf. Ohne explizites Leeren sieht ein Test die Daten des
    vorigen und das Ergebnis hängt an der Reihenfolge (#1706).
    """
    cdb.ensure_db(CRA_DB)
    con = cdb._connect(CRA_DB)
    try:
        con.try_execute("DELETE FROM cra_support_period WHERE projekt_name=?", (PROJ,))
        con.commit()
    finally:
        con.close()
    cdb.delete_projekt(CRA_DB, PROJ)
    cdb.save_projekt(CRA_DB, PROJ, unternehmen="pytest-firma-1762", produkt="Testprodukt")
    yield PROJ
    con = cdb._connect(CRA_DB)
    try:
        con.try_execute("DELETE FROM cra_support_period WHERE projekt_name=?", (PROJ,))
        con.commit()
    finally:
        con.close()
    cdb.delete_projekt(CRA_DB, PROJ)


# ── Punkt 2: Mindestdauer braucht eine Begründung ────────────────────────────

def test_unter_fuenf_jahren_ohne_begruendung_ist_nicht_speicherbar(projekt):
    with pytest.raises(ValueError) as e:
        cdb.save_support_period(CRA_DB, projekt, {
            "markteintritt_datum": "2026-01-01", "support_jahre": 3})
    text = str(e.value)
    # Die Meldung muss sagen, was fehlt UND warum — sonst rät der Anwender.
    assert "Begründung" in text
    assert "5" in text
    assert "Nutzungsdauer" in text
    assert cdb.load_support_period(CRA_DB, projekt) is None


def test_unter_fuenf_jahren_mit_begruendung_geht(projekt):
    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 3,
        "rationale": "Erwartete Nutzungsdauer der Messsonde liegt bei 3 Jahren; "
                     "danach Austausch durch Nachfolgemodell."})
    sp = cdb.load_support_period(CRA_DB, projekt)
    assert sp["support_jahre"] == 3


def test_floskel_zaehlt_nicht_als_begruendung(projekt):
    """„n/a" ist keine Begründung. Die Mindestlänge fängt genau das ab."""
    with pytest.raises(ValueError):
        cdb.save_support_period(CRA_DB, projekt, {
            "markteintritt_datum": "2026-01-01", "support_jahre": 2,
            "rationale": "n/a"})


def test_fuenf_jahre_brauchen_keine_begruendung(projekt):
    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 5})
    assert cdb.load_support_period(CRA_DB, projekt)["support_jahre"] == 5


def test_null_jahre_sind_auch_mit_begruendung_nicht_zulaessig(projekt):
    """Eine kurze Nutzungsdauer kann 0 Jahre nicht tragen: Ein Produkt ohne
    jede Update-Zusage darf nicht in Verkehr gebracht werden."""
    with pytest.raises(ValueError) as e:
        cdb.save_support_period(CRA_DB, projekt, {
            "support_jahre": 0,
            "rationale": "Wird nach dem Verkauf nicht weiter betreut, Einwegware."})
    assert "mindestens 1 Jahr" in str(e.value)


def test_unsinnige_jahresangabe_wird_benannt(projekt):
    with pytest.raises(ValueError) as e:
        cdb.save_support_period(CRA_DB, projekt, {"support_jahre": "fünf"})
    assert "keine Zahl" in str(e.value)


# ── Punkt 3: Kaufhinweis wird geprüft ────────────────────────────────────────

def test_kaufhinweis_url_muss_eine_adresse_sein(projekt):
    with pytest.raises(ValueError) as e:
        cdb.save_support_period(CRA_DB, projekt, {
            "markteintritt_datum": "2026-01-01", "support_jahre": 5,
            "support_end_kaufhinweis_url": "siehe Webshop"})
    assert "keine Adresse" in str(e.value)
    # Der Hinweis nennt den Ausweg: Freitext gehört ins Nachweis-Feld.
    assert "Nachweis-Feld" in str(e.value)


def test_freitext_im_nachweisfeld_bleibt_erlaubt(projekt):
    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 5,
        "support_end_kaufhinweis_nachweis": "Aufdruck auf der Verpackung, Muster in Ablage X"})
    sp = cdb.load_support_period(CRA_DB, projekt)
    assert "Verpackung" in sp["support_end_kaufhinweis_nachweis"]


# ── Punkt 1: zwei Kanäle, zwei Zusagen ──────────────────────────────────────

def test_zwei_kanaele_und_zwei_zusagen_werden_gespeichert(projekt):
    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 5,
        "sicherheitsupdate_kanal": "security.example.test/apt",
        "funktionsupdate_kanal": "releases.example.test",
        "sicherheitsupdates_kostenlos": True,
        "kanaele_getrennt": True,
        "update_zusagen_nachweis": "https://example.test/support-policy"})
    sp = cdb.load_support_period(CRA_DB, projekt)
    assert sp["sicherheitsupdate_kanal"] == "security.example.test/apt"
    assert sp["funktionsupdate_kanal"] == "releases.example.test"
    # Nach außen Wahrheitswerte, nicht 0/1 — sonst castet jeder Aufrufer selbst.
    assert sp["sicherheitsupdates_kostenlos"] is True
    assert sp["kanaele_getrennt"] is True


def test_bestandszeile_behauptet_keine_zusage(projekt):
    """Additive Spalten mit Default 0: Ein Altprojekt bekommt **keine** Zusage
    angedichtet, die niemand gegeben hat."""
    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 5,
        "update_kanal": "Auto-Update"})
    sp = cdb.load_support_period(CRA_DB, projekt)
    assert sp["sicherheitsupdates_kostenlos"] is False
    assert sp["kanaele_getrennt"] is False
    assert sp["update_kanal"] == "Auto-Update"   # Alt-Feld bleibt unberührt


# ── Gate-Prüfpunkte ────────────────────────────────────────────────────────

def test_gate_meldet_fehlende_zusagen_und_kaufhinweis(projekt):
    from cra.konformitaet_gate import _check_update_zusagen

    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 5})
    befunde, punkte = {}, []
    _check_update_zusagen(projekt, VERSION, befunde, punkte)
    assert befunde["update_zusagen_fehlen"] is True
    assert befunde["kaufhinweis_fehlt"] is True
    assert len(punkte) == 3      # kostenlos, Trennung, Kaufhinweis
    # Wirkung bedingt, nicht rot: Was fehlt, ist die dokumentierte Zusage —
    # ein Mangel der Nachweiskette, keine Eigenschaft des Produkts.
    assert {p["ampel_wirkung"] for p in punkte} == {"bedingt"}
    # Keine erfundene Absatznummer: Die Stelle ist im Issue als unsicher markiert.
    assert all("nicht geprüft" in p["ref"] for p in punkte)


def test_gate_schweigt_bei_vollstaendigen_zusagen(projekt):
    from cra.konformitaet_gate import _check_update_zusagen

    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 5,
        "sicherheitsupdate_kanal": "security.example.test",
        "sicherheitsupdates_kostenlos": True, "kanaele_getrennt": True,
        "support_end_kaufhinweis_url": "https://example.test/eol"})
    befunde, punkte = {}, []
    _check_update_zusagen(projekt, VERSION, befunde, punkte)
    assert punkte == []
    assert befunde["update_zusagen_fehlen"] is False
    assert befunde["kaufhinweis_fehlt"] is False


def test_getrennt_ohne_kanal_zaehlt_nicht(projekt):
    """Ein Haken ohne benannten Kanal ist keine Trennung: Dann steht nirgends,
    wo der Kunde die Sicherheitsupdates allein beziehen kann."""
    from cra.konformitaet_gate import _check_update_zusagen

    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 5,
        "sicherheitsupdates_kostenlos": True, "kanaele_getrennt": True,
        "support_end_kaufhinweis_url": "https://example.test/eol"})
    befunde, punkte = {}, []
    _check_update_zusagen(projekt, VERSION, befunde, punkte)
    assert len(punkte) == 1
    assert "getrennt" in punkte[0]["titel"]


def test_gate_meldet_unbestimmbare_fristen(projekt):
    from cra.konformitaet_gate import _check_aufbewahrungsfristen

    cdb.save_support_period(CRA_DB, projekt, {"support_jahre": 5})   # kein Markteintritt
    befunde, punkte = {}, []
    _check_aufbewahrungsfristen(projekt, VERSION, befunde, punkte)
    assert befunde["aufbewahrungsfristen_unbestimmbar"] is True
    assert len(punkte) == 1
    assert punkte[0]["ampel_wirkung"] == "bedingt"
    assert "Markteintritt" in punkte[0]["hinweis"]


def test_gate_meldet_verstrichene_fristen_bedingt(projekt):
    """Verstrichene Aufbewahrungsfrist → bedingt, **nie rot**.

    Diese Fristen sind endende Pflichten: Läuft die Aufbewahrungsfrist ab, darf
    die Akte weg. Das ist kein Versäumnis — nur ein Hinweis, dass der erfasste
    Markteintritt nicht zur aktuellen Fassung passt.
    """
    from cra.konformitaet_gate import _check_aufbewahrungsfristen

    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2005-01-01", "support_jahre": 5})
    befunde, punkte = {}, []
    _check_aufbewahrungsfristen(projekt, VERSION, befunde, punkte)
    assert punkte, "verstrichene Fristen müssen auffallen"
    assert {p["ampel_wirkung"] for p in punkte} == {"bedingt"}
    assert befunde["aufbewahrung"]["schlechteste_ampel"] == "verstrichen"


def test_laufende_fristen_erzeugen_keinen_punkt(projekt):
    from cra.konformitaet_gate import _check_aufbewahrungsfristen

    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": date.today().isoformat(), "support_jahre": 5})
    befunde, punkte = {}, []
    _check_aufbewahrungsfristen(projekt, VERSION, befunde, punkte)
    assert punkte == []
    assert befunde["aufbewahrung"]["schlechteste_ampel"] == "gruen"


def test_neue_pruefungen_sind_verdrahtet():
    """Eine Prüfung, die nicht in ``_RELEASE_CHECKS`` steht, läuft nie."""
    from cra import konformitaet_gate as gate

    namen = {c.__name__ for c in gate._RELEASE_CHECKS}
    for erwartet in ("_check_update_zusagen", "_check_aufbewahrungsfristen",
                     "_check_marktaufsicht"):
        assert erwartet in namen, erwartet


# ── Endpunkte ────────────────────────────────────────────────────────────────

def test_api_antwortet_mit_400_und_begruendung(client, auth_headers, projekt):
    r = client.post(f"/api/cra/projekte/{projekt}/support-period",
                    headers=auth_headers,
                    json={"markteintritt_datum": "2026-01-01", "support_jahre": 2})
    assert r.status_code == 400, r.json
    assert "Begründung" in r.json["error"]


def test_api_fristen_endpunkt(client, auth_headers, projekt):
    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": date.today().isoformat(), "support_jahre": 5})
    r = client.get(f"/api/cra/projekte/{projekt}/support-period/fristen",
                   headers=auth_headers)
    assert r.status_code == 200, r.json
    assert r.json["berechenbar"] is True
    assert len(r.json["fristen"]) == 4
    assert r.json["naechste"]["key"] == "support_ende"


def test_api_fristen_unbekanntes_projekt_404(client, auth_headers):
    r = client.get("/api/cra/projekte/gibt-es-nicht-1762/support-period/fristen",
                   headers=auth_headers)
    assert r.status_code == 404


# ── Kundenfassung ────────────────────────────────────────────────────────────

def test_kundenfassung_druckt_die_zusagen(projekt):
    """Die Zusagen gehören in die Kundenfassung — dafür sind sie gemacht."""
    import io

    from docx import Document

    from cra import kundenfassung

    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 5,
        "sicherheitsupdate_kanal": "security.example.test",
        "funktionsupdate_kanal": "releases.example.test",
        "sicherheitsupdates_kostenlos": True, "kanaele_getrennt": True,
        "support_end_kaufhinweis_url": "https://example.test/eol"})
    d = Document(io.BytesIO(kundenfassung.build_docx(projekt, db_path=CRA_DB)))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "kostenlos bereitgestellt" in text
    assert "getrennt von Funktionsupdates" in text
    assert "security.example.test" in text
    assert "https://example.test/eol" in text


def test_kundenfassung_behauptet_keine_fehlende_zusage(projekt):
    """Ohne dokumentierte Zusage darf im Kundendokument keine stehen.

    Die Lücke wird benannt statt verschwiegen: Eine Kundenfassung, die eine
    nicht dokumentierte Zusage behauptet, ist schlimmer als eine, die ihre
    Lücke zeigt.
    """
    import io

    from docx import Document

    from cra import kundenfassung

    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 5})
    d = Document(io.BytesIO(kundenfassung.build_docx(projekt, db_path=CRA_DB)))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "kostenlos bereitgestellt" not in text
    assert "keine dokumentierte" in text


def test_kundenfassung_bleibt_ohne_internas(projekt):
    """Die Sperrliste aus #1758 gilt für die neuen Absätze genauso."""
    import io

    from docx import Document

    from cra import kundenfassung

    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 5,
        "sicherheitsupdates_kostenlos": True, "kanaele_getrennt": True,
        "sicherheitsupdate_kanal": "security.example.test"})
    d = Document(io.BytesIO(kundenfassung.build_docx(projekt, db_path=CRA_DB)))
    text = "\n".join(p.text for p in d.paragraphs)
    for wort in kundenfassung.VERBOTENE_BEGRIFFE:
        assert wort.lower() not in text.lower(), wort


def test_eol_und_fristen_werden_weiter_berechnet(projekt):
    """Regression: Die Rechnung aus #1209 bleibt unverändert."""
    cdb.save_support_period(CRA_DB, projekt, {
        "markteintritt_datum": "2026-01-01", "support_jahre": 7})
    sp = cdb.load_support_period(CRA_DB, projekt)
    mt = date(2026, 1, 1)
    assert sp["eol_datum"] == (mt + timedelta(days=7 * 365)).isoformat()
    assert sp["update_availability_until"] == (mt + timedelta(days=10 * 365)).isoformat()
    assert sp["doku_retention_until"] == sp["doc_retention_until"]
