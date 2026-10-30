"""Vorgangsbericht je Release im Berichts-Center (#1748).

Geprüft wird nicht das Layout, sondern die Aussagen, auf die sich später jemand
beruft:

* Der Typ ist am geteilten Berichts-Center registriert — sonst gibt es DOCX, PDF,
  Historie und Word-Vorlage nicht, und es wäre ein Sonderweg daneben entstanden.
* Ein **Entwurf** ist berichtbar und als solcher gekennzeichnet. Das ist der
  Hauptfall im Tagesgeschäft: Man braucht den Bericht, um ihn prüfen zu lassen.
* Beim freigegebenen Vorgang stehen Prüfer, Freigeber **und** der SHA-256 des
  archivierten Nachweises im Dokument — der Bericht zeigt auf unveränderliche
  Stücke, statt selbst die Beweislast zu tragen.
* Eine Entscheidung erscheint mit Begründung **und Person**. Fehlt die
  Begründung, darf der Punkt nicht als erledigt dargestellt werden; ein Haken
  ohne Aussage ist als Nachweis wertlos (#1744).
* Ein nicht ermittelbarer Rechenwert heißt „nicht geprüft", nie „erfüllt"
  (#1741/#1751).
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest

from cra import berichte_provider as bp
from cra import release_bericht as rb

PROJ = 'pytest-release-bericht-1748'
VERSION = '7.4.8'
DB = Path('data/db/cra.sqlite')   # Schema 'cra' — dieselbe Quelle wie das Gate


def _text(docx_bytes: bytes) -> str:
    """Fließtext **und** Tabellenzellen — die Hälfte des Berichts sind Tabellen."""
    from docx import Document

    doc = Document(BytesIO(docx_bytes))
    teile = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            teile.extend(c.text for c in row.cells)
    return '\n'.join(teile)


def _pdf_moeglich() -> bool:
    """Lokal ist in der Regel kein Konverter da — dann wird übersprungen."""
    try:
        from shared.templates import pdf_converter as pc
        return bool(pc.is_soffice_available() or pc.is_gotenberg_available())
    except Exception:  # noqa: BLE001
        return False


@pytest.fixture
def entwurf():
    """Vorgang im Arbeitsstand: Checkliste teilweise gefüllt, keine Freigabe."""
    from cra import konformitaet_db as kdb
    from cra.db import save_projekt

    save_projekt(DB, PROJ, unternehmen='Cyberwoks GmbH', produkt='AI Compliance Suite')
    kdb.save_konformitaet(DB, PROJ, {'release_version': VERSION, 'bewertungsweg': 'A',
                                     'bewertung_abgeschlossen': True})
    kdb.set_wesentliche_aenderung(
        DB, PROJ, VERSION, 'nein',
        'Nur Fehlerbehebungen; Funktion und Sicherheitseigenschaften unverändert.')
    kdb.set_release_checkliste(
        DB, PROJ, VERSION, 'release_notes_veroeffentlicht', 'erledigt',
        'Release Notes im Kundenportal veröffentlicht', 'dev@example.com')
    kdb.set_release_checkliste(
        DB, PROJ, VERSION, 'psirt_kontakt', 'erledigt', '', 'dev@example.com')
    return DB


@pytest.fixture
def freigegeben(entwurf):
    """Derselbe Vorgang, durch die Statusmaschine geschickt + Nachweis archiviert."""
    from cra import konformitaet_db as kdb
    from cra import konformitaet_final_db as fdb

    kdb.einreichen(DB, PROJ, VERSION, von='dev@example.com')
    kdb.freigeben(DB, PROJ, VERSION, von='compliance@example.com')
    archiv = fdb.add_final(DB, PROJ, version=VERSION, pdf_bytes=b'%PDF-1.7 nachweis',
                           ampel='bedingt', freigegeben_von='compliance@example.com')
    return archiv


# ── Registrierung am Berichts-Center ────────────────────────────────────────
def test_typ_steht_im_katalog():
    """Ohne Katalog-Eintrag weist die Export-Route den Typ mit 404 ab."""
    keys = [s.key for s in bp.catalog()]
    assert 'release' in keys
    spec = next(s for s in bp.catalog() if s.key == 'release')
    assert spec.titel and spec.norm, 'Titel und Rechtsbezug gehören in die Auswahl'


def test_typ_erscheint_im_berichts_center(client, auth_headers, entwurf):
    """Der Weg, den der Anwender nimmt: die Typ-Liste des Moduls."""
    res = client.get(f'/api/cra/projekte/{PROJ}/berichte', headers=auth_headers)
    assert res.status_code == 200, res.json
    keys = [r['key'] for r in res.json['reports']]
    assert 'release' in keys
    assert res.json['ki_summary'] is True


# ── DOCX ────────────────────────────────────────────────────────────────────
def test_docx_entsteht_und_ist_nicht_leer(entwurf):
    daten = bp.render('release', 'docx', projekt=PROJ)
    assert daten[:2] == b'PK' and len(daten) > 5000
    assert 'CRA-Vorgangsbericht' in _text(daten)


def test_entwurf_ist_berichtbar_und_gekennzeichnet(entwurf):
    """Hauptfall im Tagesgeschäft — und der, der leicht verweigert würde."""
    text = _text(rb.build_docx(PROJ, db_path=DB))
    assert rb.ENTWURF_MARKER in text
    assert 'Entwurf' in text
    assert 'kein freigegebener Nachweis' in text


def test_ohne_vorgang_entsteht_trotzdem_ein_bericht():
    """Ein Projekt ohne Release-Vorgang darf keinen Abbruch erzeugen; der Bericht
    sagt, dass nichts erfasst ist, statt Zahlen zu erfinden."""
    from cra.db import save_projekt

    save_projekt(DB, PROJ + '-leer', unternehmen='Cyberwoks GmbH', produkt='X')
    text = _text(rb.build_docx(PROJ + '-leer', db_path=DB))
    assert 'kein Release-Vorgang erfasst' in text
    assert rb.ENTWURF_MARKER in text


def test_freigegebener_vorgang_nennt_pruefer_freigeber_und_hash(freigegeben):
    text = _text(rb.build_docx(PROJ, db_path=DB))
    assert rb.ENTWURF_MARKER not in text, 'ein freigegebener Vorgang ist kein Entwurf'
    assert 'dev@example.com' in text, 'Einreicher fehlt'
    assert 'compliance@example.com' in text, 'Prüfer/Freigeber fehlt'
    assert freigegeben['sha256'] in text, (
        'ohne den Hash zeigt der Bericht nicht auf ein unveränderliches Stück und '
        'müsste die Beweislast selbst tragen')
    assert 'freigegeben' in text.lower()


def test_entscheidung_erscheint_mit_begruendung_und_person(entwurf):
    text = _text(rb.build_docx(PROJ, db_path=DB))
    assert 'Release Notes im Kundenportal veröffentlicht' in text
    assert 'dev@example.com' in text
    # Art. 3 Nr. 30 — Antwort *und* Begründung, auch bei „nein".
    assert 'Funktion und Sicherheitseigenschaften unverändert' in text
    # Die Kernfrage steckt in eigenen Spalten; sie darf in der Checkliste nicht
    # als unbeantwortet erscheinen, wenn sie beantwortet ist.
    assert 'beantwortet: nein' in text


def test_eingefrorener_rechenstand_hat_vorrang(entwurf):
    """Geprüft wurde gegen den Stand bei der Einreichung — nicht gegen heute.

    Sonst behauptet der Bericht, es sei etwas geprüft worden, was inzwischen
    anders aussieht.
    """
    from cra import konformitaet_db as kdb

    kdb.einreichen(DB, PROJ, VERSION, von='dev@example.com',
                   gate={'ampel': 'gruen', 'offene_punkte': [],
                         'sbom_version_fehlt': False})
    text = _text(rb.build_docx(PROJ, db_path=DB))
    assert 'eingefroren bei Einreichung' in text
    assert 'KONFORM' in text
    # Der heutige Stand ist rot (nichts bewertet) — die Abweichung muss stehen.
    assert 'heutige Rechenstand' in text


def test_fehlender_eingefrorener_stand_wird_benannt(freigegeben):
    """``einreichen`` ohne Gate kommt vor (Altbestand, direkter DB-Zugriff).

    Dann ist nicht belegbar, gegen welche Zahlen geprüft wurde — und genau das
    muss im Bericht stehen, statt den heutigen Stand als den geprüften auszugeben.
    """
    text = _text(rb.build_docx(PROJ, db_path=DB))
    assert 'kein bei der Einreichung' in text


def test_wesentliche_aenderung_ohne_antwort_ist_nicht_nein(entwurf):
    """„offen" darf nicht als „nein" gelesen werden — „nein" hat Folgen."""
    from cra import konformitaet_db as kdb

    kdb.save_konformitaet(DB, PROJ, {'release_version': '9.0.0', 'bewertungsweg': 'A',
                                     'bewertung_abgeschlossen': True})
    text = _text(rb.build_docx(PROJ, version='9.0.0', db_path=DB))
    assert '9.0.0' in text, 'der Bericht muss den angefragten Vorgang zeigen'
    assert 'Nicht beantwortet' in text


def test_drei_arten_sind_sichtbar_getrennt(entwurf):
    text = _text(rb.build_docx(PROJ, db_path=DB))
    for titel in rb.ART_LABEL.values():
        assert titel in text, titel


def test_phase1_ohne_feststellung(entwurf):
    """Ohne ausdrückliche Feststellung steht genau das im Bericht (#1746)."""
    text = _text(rb.build_docx(PROJ, db_path=DB))
    assert VERSION in text, 'ohne den Vorgang im Dokument prüft der Test nichts'
    assert 'Nicht festgestellt' in text


def test_phase1_feststellung_erscheint_mit_person(entwurf):
    from cra import release_db as rdb

    rdb.cra_ready_feststellen(DB, PROJ, von='compliance@example.com',
                              gate={'ampel': 'gruen'})
    text = _text(rb.build_docx(PROJ, db_path=DB))
    assert 'Festgestellt' in text and 'compliance@example.com' in text


def test_nachweis_wird_zitiert_nicht_neu_erzeugt(entwurf, monkeypatch):
    """Zwei Renderer für dasselbe Dokument wären zwei Wahrheiten (#1727).

    Der Vorgangsbericht darf den juristischen Nachweis also nicht aufrufen —
    sonst wäre er ein zweiter Erzeugungsweg mit möglicherweise anderen Zahlen.
    """
    from cra import konformitaet_report as kr

    def _verboten(*_a, **_k):
        raise AssertionError('der Vorgangsbericht erzeugt den Nachweis neu')

    monkeypatch.setattr(kr, 'build_docx', _verboten)
    monkeypatch.setattr(kr, 'build_pdf', _verboten)
    assert rb.build_docx(PROJ, db_path=DB)[:2] == b'PK'


@pytest.mark.skipif(not _pdf_moeglich(),
                    reason='kein DOCX→PDF-Konverter (Gotenberg/soffice) verfügbar')
def test_pdf_laeuft_ueber_denselben_docx_weg(entwurf):
    daten = bp.render('release', 'pdf', projekt=PROJ)
    assert daten[:4] == b'%PDF'


# ── Word-Vorlage (#1735) ────────────────────────────────────────────────────
def test_vorlagen_weg_traegt_bericht_typ_release(monkeypatch):
    """Mit ``template_id`` läuft der Vorlagen-Weg; die Vorlage muss erkennen
    können, dass sie einen Release-Vorgangsbericht füllt."""
    from shared.reports import vorlagen as vb
    from shared.templates import schema

    rec = {'id': 11, 'modul': 'cra', 'name': 'Partner-Vorlage', 'version': 1,
           'ist_default': 1, 'aktiv': 1}

    class _FakeDB:
        DEFAULT_DB_PATH = Path('/dev/null')

        def get_template(self, _db, tid):  # noqa: ANN001
            return rec if int(tid) == rec['id'] else None

        def list_templates(self, _db, modul=None):  # noqa: ANN001
            return [rec] if modul in (None, 'cra') else []

    monkeypatch.setattr(vb, 'tdb', _FakeDB())
    monkeypatch.setattr(schema, 'get_context_builder',
                        lambda m: (lambda db, p: {'projekt': {'name': p}}))
    gesehen: dict = {}

    def _render(_db, tid, context):  # noqa: ANN001
        gesehen.update({'tid': tid, 'context': context})
        return b'PK-DOCX'

    monkeypatch.setattr(vb.engine, 'render_docx', _render)

    out = vb.render_vorlage('cra', 11, projekt=PROJ, typ='release', fmt='docx')
    assert out == b'PK-DOCX'
    assert gesehen['context']['bericht_typ'] == 'release'
    # Die Vorlage gehört dem Modul — sonst käme man über CRA an fremde Berichte.
    assert rec['modul'] == vb.engine_modul('cra')


# ── Reine Funktionen: die Regeln ohne Datenbank ─────────────────────────────
def test_fehlende_begruendung_ist_nicht_erledigt():
    """Die wichtigste Regel des Berichts. ``set_release_checkliste`` verhindert
    diesen Zustand heute, Altbestand kann ihn aber tragen — und dann darf er
    nicht als Nachweis durchgehen."""
    from cra import release_checkliste as rc

    item = rc.item('release_notes_veroeffentlicht')
    erg, label = rb.entscheidungs_ergebnis(item, {'status': 'erledigt',
                                                  'begruendung': '   '})
    assert erg == 'maengel'
    assert 'nicht nachgewiesen' in label
    assert label != 'erledigt'

    erg, label = rb.entscheidungs_ergebnis(
        item, {'status': 'erledigt', 'begruendung': 'Advisory veröffentlicht'})
    assert (erg, label) == ('erfuellt', 'erledigt')

    erg, label = rb.entscheidungs_ergebnis(item, {})
    assert (erg, label) == ('offen', 'offen')


def test_nicht_geprueft_wird_nie_erfuellt():
    """Ein Rechenwert, den die Anwendung nicht ermitteln konnte, ist ``None``.

    Würde er als 0/„ok" durchgehen, sähe ein ungeprüfter Punkt wie ein
    erfüllter aus — in einem Konformitätsnachweis der gefährlichste Fehler.
    """
    assert rb.gerechnetes_ergebnis('sbom_version_fehlt', {}) [0] == 'ungeprueft'
    assert rb.gerechnetes_ergebnis(
        'sbom_version_fehlt', {'sbom_version_fehlt': None})[0] == 'ungeprueft'
    assert rb.gerechnetes_ergebnis(
        'behobene_ohne_advisory', {'behobene_ohne_advisory': None})[0] == 'ungeprueft'
    # Unbekannte Quelle = Katalogfehler, nicht erfüllt.
    assert rb.gerechnetes_ergebnis('gibt_es_nicht', {'gibt_es_nicht': 0})[0] == 'ungeprueft'
    # Und die positiven Fälle, damit der Test nicht nur „alles grau" prüft.
    assert rb.gerechnetes_ergebnis(
        'sbom_version_fehlt', {'sbom_version_fehlt': False})[0] == 'erfuellt'
    assert rb.gerechnetes_ergebnis(
        'offene_kritische_schwachstellen',
        {'offene_kritische_schwachstellen': 2})[0] == 'offen'
    assert rb.gerechnetes_ergebnis(
        'annex_vollstaendig', {'annex_vollstaendig': True})[0] == 'erfuellt'


def test_jede_gerechnete_quelle_hat_eine_auswertung():
    """Ein neuer gerechneter Katalogpunkt ohne Auswertung erschiene stumm als
    „nicht geprüft" — der Test macht das sichtbar."""
    from cra import release_checkliste as rc

    fehlend = [i.id for i in rc.KATALOG
               if i.art == 'gerechnet' and i.quelle not in rb._GERECHNET_AUSWERTUNG]
    assert not fehlend, f'ohne Auswertung im Vorgangsbericht: {fehlend}'


def test_vorgangswahl_nimmt_den_juengsten_und_nie_den_produktdatensatz():
    """Das Berichts-Center reicht keine Version durch — die Wahl muss
    vorhersagbar sein. Der Datensatz ohne Version ist Phase 1, kein Release."""
    vorgaenge = [
        {'release_version': '', 'updated_at': '2030-01-01 00:00:00'},
        {'release_version': '1.0.0', 'updated_at': '2026-01-01 00:00:00'},
        {'release_version': '2.0.0', 'updated_at': '2026-05-01 00:00:00'},
    ]
    assert rb.waehle_vorgang(vorgaenge)['release_version'] == '2.0.0'
    assert rb.waehle_vorgang(vorgaenge, '1.0.0')['release_version'] == '1.0.0'
    assert rb.waehle_vorgang(vorgaenge, '9.9.9') is None
    assert rb.waehle_vorgang([{'release_version': ''}]) is None


# ── KI-Zusammenfassung (#1725) ──────────────────────────────────────────────
def test_summary_context_hat_die_pflichtschluessel(entwurf):
    ctx = bp.summary_context(PROJ)
    assert 'reifegrad' in ctx and 'offene' in ctx, (
        'nur über diese zwei Schlüssel kommen die Zahlen in den Prompt (#1725)')
    assert isinstance(ctx['offene'], list)


def test_summary_context_traegt_den_vorgang_als_zusatz(freigegeben):
    """Zusatzschlüssel gehen als JSON-Block mit — die Zusammenfassung soll
    wissen, ob das Release geprüft und freigegeben ist."""
    ctx = bp.summary_context(PROJ)
    assert 'reifegrad' in ctx and 'offene' in ctx
    vorgang = ctx['release_vorgang']
    assert vorgang['release_version'] == VERSION
    assert vorgang['ist_entwurf'] is False
    assert vorgang['freigegeben_von'] == 'compliance@example.com'
    assert vorgang['archivierte_nachweise'] == 1
