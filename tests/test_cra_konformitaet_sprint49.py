"""Sprint #49: Konformitätsnachweis je Release — Akzeptanz (#1608/#1613),
Gate (#1612), Nachweis-Report (#1610, A4/Arial), Sign-off (#1611)."""
import shutil
from io import BytesIO
from pathlib import Path

import pytest

CRA = "/api/cra"
KONF = "/api/cra-konformitaet"
FIRMA = "pytest-firma-1607"
PROJ = "pytest-cra-1607"
CRA_DB = Path("data/db/cra.sqlite")
RB_DB = Path("data/db/risikobewertung.sqlite")


@pytest.fixture(autouse=True)
def _full_license():
    from server import license_state
    cur = license_state._current
    prev = (cur.state, list(cur.modules))
    cur.state, cur.modules = "ok", ["*"]
    yield
    cur.state, cur.modules = prev[0], prev[1]


@pytest.fixture
def projekt(client, auth_headers):
    client.delete(f"{CRA}/projekte/{PROJ}", headers=auth_headers)
    client.post(f"{CRA}/projekte", headers=auth_headers, json={"name": PROJ, "unternehmen": FIRMA})
    yield PROJ
    client.delete(f"{CRA}/projekte/{PROJ}", headers=auth_headers)


def test_behandlung_akzeptanz_braucht_kategorie():
    """#1613: Akzeptanz ohne Begründungskategorie ist unzulässig."""
    from risikobewertung.db import set_risiko_behandlung, save_risiko, save_projekt
    save_projekt(RB_DB, "pytest-rb-1607", "STRIDE", "")
    rid = save_risiko(RB_DB, {"projekt_name": "pytest-rb-1607", "nr": 1, "risk_name": "R",
                              "beschreibung": "x", "framework": "STRIDE", "felder": {}})
    with pytest.raises(ValueError):
        set_risiko_behandlung(RB_DB, rid, behandlung="akzeptiert")   # keine Kategorie
    res = set_risiko_behandlung(RB_DB, rid, behandlung="akzeptiert", kategorie="restrisiko",
                                grund="vertretbar", von="GF")
    assert res["behandlung"] == "akzeptiert" and res["is_resolved"] == 1


def test_gate_begruendete_abweichung_zaehlt(client, auth_headers, projekt):
    """#1612/#1613: eine akzeptierte Abweichung zählt nicht als offene Nicht-Konformität."""
    from cra.konformitaet_gate import evaluate
    from shared import req_risks
    from risikobewertung.db import set_risiko_behandlung
    # Gap-Risiko für eine unerfüllte Anforderung + akzeptieren
    req_risks.create_rb("cra", projekt, "STRIDE")
    gap = req_risks.create_gap_risk("cra", projekt, "requirement", "AI1-01", titel="Risikobasierte Cybersicherheit")
    offen_vorher = evaluate(projekt)["offen"]
    set_risiko_behandlung(RB_DB, int(gap["risk_id"]), behandlung="akzeptiert",
                          kategorie="restrisiko", grund="Kompensiert durch Betriebsprozess",
                          von="CRA-Verantwortliche", rechtsgrundlage="Annex I Part I, Abs.1 (where applicable)")
    g = evaluate(projekt)
    # AI1-01 ist jetzt begründet, nicht mehr offen
    assert g["begruendet_akzeptiert"] >= 1
    assert g["offen"] == offen_vorher - 1
    abw = {a["requirement_id"]: a for a in g["abweichungen"]}
    assert abw["AI1-01"]["art"] == "begruendet" and abw["AI1-01"]["kategorie"] == "restrisiko"


def test_nachweis_docx_enthaelt_rechtsgrundlage(client, auth_headers, projekt):
    """#1610: Nachweis-DOCX enthält Rechtsgrundlage (Regulation + Artikelrefs)."""
    from cra.konformitaet_report import build_docx
    data = build_docx(projekt, version="v1.0")
    from docx import Document
    d = Document(BytesIO(data))
    txt = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            txt += " " + " ".join(c.text for c in row.cells)
    assert "2024/2847" in txt                     # Rechtsgrundlage sichtbar
    assert "Annex I" in txt or "Annex" in txt
    # A4-Section
    s = d.sections[0]
    assert round(s.page_width / 36000) == 210


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice nicht verfügbar")
def test_nachweis_pdf_a4_vektor(client, auth_headers, projekt, monkeypatch):
    monkeypatch.setenv("GOTENBERG_URL", "http://127.0.0.1:1")
    from cra.konformitaet_report import build_pdf
    pdf = build_pdf(projekt, version="v1.0")
    import pdfplumber
    with pdfplumber.open(BytesIO(pdf)) as p:
        pg = p.pages[0]
        assert 590 < pg.width < 600 and 835 < pg.height < 848
        assert "Konformit" in (pg.extract_text() or "")


def test_release_check_endpoint(client, auth_headers, projekt):
    r = client.get(f"{KONF}/projekte/{projekt}/release-check", headers=auth_headers)
    assert r.status_code == 200
    d = r.get_json()
    assert "ampel" in d and "offene_punkte" in d and "abweichungen" in d


def test_release_versions_endpoint_ohne_repo(client, auth_headers, projekt):
    """#1662: Versions-Endpunkt liefert leere Liste, wenn kein Repo verknüpft ist (kein externer Call)."""
    r = client.get(f"{KONF}/projekte/{projekt}/versions", headers=auth_headers)
    assert r.status_code == 200
    d = r.get_json()
    assert d["versions"] == [] and d["repo"] in ("", None)


def test_gate_deeplinks_sind_echte_tab_ids(client, auth_headers, projekt):
    """#1661: offene Punkte verweisen auf echte CRA-Tab-IDs (nicht 'Übersicht')."""
    from cra.konformitaet_gate import evaluate
    g = evaluate(projekt)
    gueltig = {"requirements", "pflichtdoku", "traceability", "konformitaet"}
    for o in g["offene_punkte"]:
        if o.get("link"):
            assert o["link"] in gueltig, o


def test_gate_erkennt_freigegebene_doc_als_managed_doc(client, auth_headers, projekt):
    """#1665: eine im Dokumente-Tab freigegebene EU-Konformitätserklärung zählt im Gate als DoC."""
    from cra.konformitaet_gate import evaluate
    from shared.documents.db import create_document, set_status
    assert evaluate(projekt)["doc_ausgestellt"] is False
    did = create_document(CRA_DB, "cra", projekt=projekt, doc_type="konformitaetserklaerung",
                          titel="EU-DoC", content_html="<p>Konformität erklärt.</p>")
    set_status(CRA_DB, "cra", did, "freigegeben")
    assert evaluate(projekt)["doc_ausgestellt"] is True


def test_ce_kennzeichnung_persistenz_und_report(client, auth_headers, projekt):
    """#1671: CE-Kennzeichnungsnachweis wird gespeichert und erscheint im Nachweis."""
    from cra import konformitaet_db as kdb
    from cra.konformitaet_report import build_docx
    kdb.save_konformitaet(CRA_DB, projekt, {
        "bewertungsweg": "A", "ce_status": "ce_angebracht",
        "ce_angebracht_am": "2026-10-19", "ce_ort": "Info-Dialog + Handbuch S.1",
        "ce_nachweis": "Screenshot build-42"})
    rec = kdb.get_konformitaet(CRA_DB, projekt, "")
    assert rec["ce_ort"] == "Info-Dialog + Handbuch S.1"
    assert str(rec["ce_angebracht_am"])[:10] == "2026-10-19"
    from docx import Document
    from io import BytesIO
    d = Document(BytesIO(build_docx(projekt)))
    txt = " ".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "CE-Kennzeichnung" in txt and "Info-Dialog" in txt


def test_akzeptiertes_restrisiko_im_nachweis(client, auth_headers, projekt):
    """#1674: ein direkt in der RB akzeptiertes Risiko (nicht aus Anforderung) erscheint im Nachweis."""
    from shared import req_risks
    from risikobewertung.db import save_risiko, set_risiko_behandlung
    from cra.konformitaet_report import build_docx
    rb = req_risks.create_rb("cra", projekt, "STRIDE")
    rid = save_risiko(RB_DB, {"projekt_name": rb, "nr": 1, "risk_name": "Eigenständiges Risiko",
                              "beschreibung": "x", "framework": "STRIDE", "felder": {}})
    set_risiko_behandlung(RB_DB, rid, behandlung="akzeptiert", kategorie="restrisiko",
                          grund="bewusst getragen", von="GF", rechtsgrundlage="CRA Annex I")
    from docx import Document
    from io import BytesIO
    txt = "\n".join(p.text for p in Document(BytesIO(build_docx(projekt))).paragraphs)
    assert "Akzeptierte Restrisiken" in txt and "Eigenständiges Risiko" in txt and "bewusst getragen" in txt
