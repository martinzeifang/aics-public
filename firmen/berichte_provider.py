"""Firmen-Adapter für das geteilte Berichts-Center (v7.0.0).

Berichte auf Firmenebene — modulübergreifend aus :func:`firmen.uebersicht.build_uebersicht`:

* ``gesamt``  — Firmen-Gesamtbericht (Management): Umsetzungsgrad je Modul/Projekt,
  Risiko-Lage, DSGVO-TOM/WiBA, offene Aufgaben, Produkte/Nachweise.
* ``risiken`` — Risikobericht: alle offenen Risiken/Schwachstellen des Risiko-Cockpits.

DOCX via python-docx; PDF per Gotenberg/LibreOffice (gleicher Weg wie die Modul-Berichte).
"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

from shared.reports.core import ReportSpec

DB_PATH = Path('data/db/firmen.sqlite')

CATALOG: list[ReportSpec] = [
    ReportSpec('gesamt', 'Firmen-Gesamtbericht', 'CRA · NIS2 · AI Act · DSGVO · WiBA',
               'Management-Überblick: Umsetzungsgrad aller Module und Projekte, Risiko-Lage, '
               'TOM/WiBA, offene Aufgaben.'),
    ReportSpec('risiken', 'Risikobericht (Risiko-Cockpit)', 'firmenweit',
               'Alle offenen Risiken und Schwachstellen der Firma nach Schwere, mit Herkunft.'),
]

SEV_LABEL = {'critical': 'kritisch', 'high': 'hoch', 'medium': 'mittel', 'low': 'gering',
             'unknown': 'unbewertet'}
SEV_ORDER = ['critical', 'high', 'medium', 'low', 'unknown']


def catalog() -> list[ReportSpec]:
    return CATALOG


def summary_context(firma: str) -> dict:
    """Kennzahlen für die KI-Management-Zusammenfassung (#1393) — firmenweit."""
    from firmen.uebersicht import build_uebersicht
    u = build_uebersicht(firma)
    return {
        'gesamt_umsetzungsgrad_pct': u['gesamt_pct'],
        'module': [{'modul': m['label'], 'umsetzungsgrad_pct': m['pct'],
                    'projekte': [{k: p.get(k) for k in ('name', 'pct', 'risiken', 'offen')}
                                 for p in m['projekte']]}
                   for m in u['module'] if m['projekte']],
        'risiken': {'offen_gesamt': u['risiken']['total'], 'nach_schwere': u['risiken']['by_severity'],
                    'top': u['risiken']['top']},
        'dsgvo_tom': u['tom'],
        'offene_aufgaben': u['aufgaben'],
    }


def _table(doc, header: list[str], rows: list[list[Any]]):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(header))
    tbl.style = 'Light Grid Accent 1'
    for ci, h in enumerate(header):
        tbl.cell(0, ci).paragraphs[0].add_run(h).bold = True
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            tbl.cell(ri, ci).text = '' if val is None else str(val)
    return tbl


def _pct(v: Any) -> str:
    return '—' if v is None else f"{float(v):.0f} %"


def _kopf(doc, titel: str, u: dict[str, Any]) -> None:
    from docx.shared import Inches
    for s in doc.sections:
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)
    doc.add_heading(f"{titel} — {u['firma']}", level=0)
    p = doc.add_paragraph()
    p.add_run('Erstellt: ').bold = True
    p.add_run(date.today().strftime('%d.%m.%Y'))
    if u.get('unternehmen'):
        p.add_run('   ·   Unternehmen: ').bold = True
        p.add_run(u['unternehmen'])
    if u.get('berater'):
        p.add_run('   ·   Berater: ').bold = True
        p.add_run(u['berater'])


def _gesamt_docx(u: dict[str, Any]) -> bytes:
    from docx import Document
    doc = Document()
    _kopf(doc, 'Firmen-Gesamtbericht', u)

    doc.add_heading('1. Zusammenfassung', level=1)
    r = u['risiken']
    sev = r.get('by_severity') or {}
    doc.add_paragraph(
        f"Gesamt-Umsetzungsgrad (Ø der Module CRA, NIS2, AI Act, DSGVO, WiBA): {_pct(u['gesamt_pct'])}. "
        f"Offene Risiken/Schwachstellen: {r.get('total', 0)} "
        f"(kritisch {sev.get('critical', 0)}, hoch {sev.get('high', 0)}, mittel {sev.get('medium', 0)}, "
        f"gering {sev.get('low', 0)}). Produkte: {len(u['produkte'])}, Nachweise: {u['nachweise']}.")

    doc.add_heading('2. Umsetzungsgrad je Modul', level=1)
    _table(doc, ['Modul', 'Umsetzungsgrad', 'Projekte'],
           [[m['label'], _pct(m['pct']) if m['key'] != 'risikobewertung' else '—', len(m['projekte'])]
            for m in u['module'] if m['projekte']])
    for m in u['module']:
        if not m['projekte']:
            continue
        doc.add_heading(m['label'], level=2)
        if m['key'] == 'risikobewertung':
            _table(doc, ['Bewertung', 'Framework', 'Risiken', 'offen', 'unbewertet'],
                   [[p['name'], p.get('framework', ''), p.get('risiken', 0), p.get('offen', 0),
                     p.get('unbewertet', 0)] for p in m['projekte']])
        else:
            _table(doc, ['Projekt', 'Produkt', 'Umsetzungsgrad'],
                   [[p['name'], p.get('produkt') or '—', _pct(p.get('pct'))] for p in m['projekte']])

    doc.add_heading('3. Risiko-Lage (Risiko-Cockpit)', level=1)
    if r.get('unassigned'):
        doc.add_paragraph('Die Firma ist keiner Firmen-ID zugeordnet — Risiko-Cockpit nicht verfügbar.')
    elif not r.get('top'):
        doc.add_paragraph('Keine offenen Risiken oder Schwachstellen.')
    else:
        _table(doc, ['Schwere', 'Quelle', 'Projekt', 'Risiko'],
               [[SEV_LABEL.get(t.get('severity'), t.get('severity')), (t.get('source') or '').upper(),
                 t.get('projekt'), t.get('titel')] for t in r['top']])
        doc.add_paragraph('Vollständige Liste: Risikobericht (Risiko-Cockpit).')

    if u['tom']:
        doc.add_heading('4. DSGVO — TOM (Art. 32) und BSI-WiBA', level=1)
        _table(doc, ['DSGVO-Projekt', 'TOM', 'Erfüllung', 'Status aus WiBA', 'mit Umsetzungsbeschreibung'],
               [[t['projekt'], t['anzahl'], _pct(t['pct']), t['aus_wiba'], t['mit_umsetzung']]
                for t in u['tom']])

    if u['aufgaben']:
        doc.add_heading('5. Offene Aufgaben und Fristen (DSGVO-DSMS)', level=1)
        _table(doc, ['Bereich', 'Aufgabe', 'Frist'],
               [[a.get('area_label') or a.get('area') or '', a.get('text') or '',
                 f"{a.get('due') or '—'}{' (überfällig)' if a.get('overdue') else ''}"]
                for a in u['aufgaben']])

    doc.add_paragraph()
    doc.add_paragraph('Hinweis: Werte aus der AI Compliance Suite, Stand Berichtserstellung. '
                      'Umsetzungsgrade sind Selbstbewertungen bzw. aus BSI-WiBA abgeleitet.')
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _risiken_docx(u: dict[str, Any], items: list[dict[str, Any]]) -> bytes:
    from docx import Document
    doc = Document()
    _kopf(doc, 'Risikobericht', u)
    if u['risiken'].get('unassigned'):
        doc.add_paragraph('Die Firma ist keiner Firmen-ID zugeordnet — Risiko-Cockpit nicht verfügbar.')
    for sev in SEV_ORDER:
        rows = [it for it in items if it.get('severity') == sev]
        if not rows:
            continue
        doc.add_heading(f"{SEV_LABEL[sev].capitalize()} ({len(rows)})", level=1)
        _table(doc, ['Quelle', 'Projekt', 'Risiko / Schwachstelle', 'Beschreibung'],
               [[(it.get('source') or '').upper(), it.get('projekt'), it.get('titel'),
                 it.get('beschreibung') or ''] for it in rows])
    if not items and not u['risiken'].get('unassigned'):
        doc.add_paragraph('Keine offenen Risiken oder Schwachstellen.')
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render(typ: str, fmt: str, *, projekt: str, von: str | None = None,
           bis: str | None = None, **_ctx: Any) -> bytes:
    """``projekt`` = Firmenname."""
    from firmen.uebersicht import build_uebersicht
    u = build_uebersicht(projekt)
    if typ == 'risiken':
        items: list[dict[str, Any]] = []
        if u['firmen_id'] is not None:
            from shared.risk_cockpit import build_cockpit
            items = build_cockpit(u['firmen_id']).get('items') or []
        docx_bytes = _risiken_docx(u, items)
    else:
        docx_bytes = _gesamt_docx(u)
    if fmt == 'pdf':
        from shared.templates.pdf_converter import convert_docx_to_pdf
        return convert_docx_to_pdf(docx_bytes)
    return docx_bytes
