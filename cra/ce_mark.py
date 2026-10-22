"""CE-Kennzeichnung für Dokumente — eine Quelle, gesperrt gegen Missbrauch.

Die CE-Kennzeichnung ist kein Logo, sondern ein **Rechtsakt**: Sie erklärt,
dass das Produkt die Konformitätsbewertung durchlaufen hat und eine
EU-Konformitätserklärung vorliegt (Art. 29/30 i.V.m. Art. 28). Sie auf ein
Dokument zu drucken, ohne dass beides zutrifft, ist unzulässige Verwendung —
kein Darstellungsfehler.

Deshalb entscheidet nicht der Aufrufer, ob das Zeichen erscheint, sondern
:func:`darf_zeichen`. Wer die Prüfung umgehen will, muss das ausdrücklich tun.

Bei **Software** ist das Dokument der vorgesehene Ort: Da sich das Zeichen
nicht am Produkt anbringen lässt, erscheint es auf der Konformitätserklärung
bzw. der begleitenden Website.

Die Grafik liegt als SVG (Anzeige, Download) und PNG (DOCX — python-docx
kennt kein SVG) unter ``cra/assets/``. Proportionen nach Verordnung (EG)
Nr. 765/2008, Anhang II; Mindesthöhe 5 mm.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

ASSETS = Path(__file__).resolve().parent / "assets"
SVG = ASSETS / "ce_kennzeichen.svg"
PNG = ASSETS / "ce_kennzeichen.png"

#: Mindesthöhe nach Anhang II. Kleiner darf das Zeichen nicht gedruckt werden.
MIN_HOEHE_MM = 5.0

#: Höhe im Dokument. Deutlich über dem Mindestmaß, damit die Verkleinerung
#: beim Druck die 5 mm nicht unterschreitet.
HOEHE_MM = 12.0


def darf_zeichen(konformitaet: dict[str, Any] | None,
                 doc_ausgestellt: bool) -> tuple[bool, str]:
    """Darf die Kennzeichnung gedruckt werden? ``(ja, grund)``.

    Beide Bedingungen müssen erfüllt sein — die Kennzeichnung bezieht sich auf
    eine ausgestellte Erklärung, und eine Erklärung setzt die abgeschlossene
    Bewertung voraus. Der Grund wird mitgeliefert, damit das Dokument sagen
    kann, **warum** das Zeichen fehlt, statt es wortlos wegzulassen.
    """
    k = konformitaet or {}
    if not doc_ausgestellt:
        return False, ("CE-Kennzeichnung noch nicht ausweisbar — "
                       "EU-Konformitätserklärung nicht ausgestellt")
    angebracht = str(k.get("ce_angebracht_am") or "").strip()
    if not angebracht and str(k.get("ce_status") or "") != "ce_angebracht":
        return False, "CE-Kennzeichnung noch nicht angebracht"
    return True, ""


def fundstelle(konformitaet: dict[str, Any] | None) -> str:
    """Einzeiler für den Bericht: wann und wo angebracht."""
    k = konformitaet or {}
    am = str(k.get("ce_angebracht_am") or "")[:10]
    ort = str(k.get("ce_ort") or "").strip()
    nachweis = str(k.get("ce_nachweis") or "").strip()
    teile = []
    if am:
        teile.append(f"angebracht am {am}")
    if ort:
        teile.append(f"Fundstelle: {ort}")
    if nachweis:
        teile.append(nachweis)
    return " · ".join(teile) or "angebracht"


def einfuegen(doc: Any, konformitaet: dict[str, Any] | None,
              doc_ausgestellt: bool, *, hoehe_mm: float = HOEHE_MM) -> bool:
    """Zeichen in ein ``python-docx``-Dokument setzen, wenn es zulässig ist.

    Gibt zurück, ob gedruckt wurde. Ist es unzulässig, wird der Grund als Text
    gesetzt — ein Dokument, das die Kennzeichnung stillschweigend weglässt,
    lässt den Leser im Unklaren, ob sie fehlt oder nur nicht gedruckt wurde.
    """
    from docx.shared import Mm

    ja, grund = darf_zeichen(konformitaet, doc_ausgestellt)
    if not ja:
        doc.add_paragraph(grund)
        return False
    if not PNG.exists():       # Asset fehlt im Image → Text statt stillem Nichts
        doc.add_paragraph(f"CE-Kennzeichnung {fundstelle(konformitaet)}")
        return False
    doc.add_picture(str(PNG), height=Mm(max(hoehe_mm, MIN_HOEHE_MM)))
    doc.add_paragraph(f"CE-Kennzeichnung {fundstelle(konformitaet)}")
    return True
