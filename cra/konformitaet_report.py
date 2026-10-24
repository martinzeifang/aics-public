"""Kundenfähiger CRA-Konformitätsnachweis je Release (#1610/#1613).

Erzeugt einen versionsgebundenen Nachweis (DOCX/PDF, A4, vektoriell) mit sichtbarer
Rechtsgrundlage: Erfüllungsgrad je Anforderung (Artikel-/Annex-Referenz), begründete
Abweichungen inkl. akzeptiertem Risiko, Annex-VII-Vollständigkeit, Schwachstellen-Status,
EU-Konformitätserklärung/CE und das Konformitäts-Gate-Fazit.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor

CRA_DB = Path("data/db/cra.sqlite")

_EU_BLUE = RGBColor(0x00, 0x33, 0x99)
_REGULATION = "Verordnung (EU) 2024/2847 (Cyber Resilience Act)"
_VERFAHREN = {"A": "Modul A — interne Kontrolle (Selbstbewertung), Annex VIII",
              "B+C": "Modul B+C — EU-Baumusterprüfung, Annex VIII",
              "H": "Modul H — umfassende Qualitätssicherung, Annex VIII",
              "EUCC": "EUCC-Zertifizierung"}
_AMPEL = {"gruen": "KONFORM", "bedingt": "BEDINGT KONFORM", "rot": "NICHT KONFORM"}
_STATUS_LABEL = {"erfuellt": "erfüllt", "begruendet": "Abweichung begründet akzeptiert",
                 "offen": "offen (nicht konform)"}


def _ce_text(konf: dict | None) -> str:
    """CE-Kennzeichnungs-Status für den Nachweis (#1671)."""
    k = konf or {}
    if k.get("ce_status") == "ce_angebracht" or k.get("ce_angebracht_am"):
        am = str(k.get("ce_angebracht_am") or "")[:10]
        ort = k.get("ce_ort") or ""
        teile = ["angebracht"]
        if am:
            teile.append(f"am {am}")
        if ort:
            teile.append(f"({ort})")
        return " ".join(teile)
    if k.get("doc_ausgestellt"):
        return "DoC ausgestellt — CE anzubringen"
    return "noch nicht angebracht"


def _h(doc, text, size=13, color=None, space_before=10):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    return p


def _p(doc, text, size=10):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size)
    return p


#: Was eine Word-Vorlage des Nachweises mindestens benutzen muss (#1777).
#:
#: Gelesen von ``shared.templates.validate`` (guarded import, Muster
#: ``cra/doc_export.PFLICHT_VORLAGEN``). Die Liste ist **kurz und bewusst
#: nicht nummeriert**: Der Nachweis ist kein Formular aus einem Anhang, sein
#: Aufbau gehört uns. Was nicht uns gehört, ist sein Zweck — und daraus folgen
#: genau drei Angaben: Ohne Produkt und Release belegt er nichts Bestimmtes,
#: ohne Rechtsgrundlage ist er ein Produktblatt, und ohne Gesamturteil trifft
#: er keine Aussage. Eine Vorlage, die eine davon weglässt, erzeugt ein
#: Dokument, das wie ein Nachweis aussieht und keiner ist. Eine längere Liste
#: wäre eine erfundene Rechtspflicht — in einer Nachweisakte schlechter als
#: keine.
PFLICHT_VORLAGEN: list[dict[str, Any]] = [{
    "dokument": "Konformitätsnachweis",
    "rechtsgrundlage": "Verordnung (EU) 2024/2847, Art. 13 i.V.m. Anhang I und VII",
    "erkennung": "nachweis",
    "angaben": [
        {"bezeichnung": "Produkt und Release, auf das sich der Nachweis bezieht",
         "tokens": ["nachweis.produkt", "nachweis.version"]},
        {"bezeichnung": "Rechtsgrundlage",
         "tokens": ["nachweis.rechtsgrundlage"],
         "textbeleg": "2024/2847"},
        {"bezeichnung": "Gesamturteil (Konformitäts-Gate)",
         "tokens": ["nachweis.ampel_label"],
         "alternativen": ["nachweis.ampel", "nachweis.konform"]},
    ],
}]


HINWEIS_TEXT = (
    "Dieser Nachweis dokumentiert den Stand zum angegebenen Datum/Release. "
    "Begründete Abweichungen beruhen auf dem risikobasierten Ansatz des CRA; "
    "offene Nicht-Konformitäten sind vor der CE-Kennzeichnung zu schließen.")


def daten(projekt_name: str, *, version: str | None = None,
          db_path: Path = CRA_DB) -> dict[str, Any]:
    """Alle Angaben des Nachweises — **eine** Erhebung für beide Wege (#1777).

    Der Nachweis entsteht seit #1610 programmatisch; seit #1777 kann ihn auch
    eine hochgeladene Word-Vorlage erzeugen. Beide müssen dieselben Zahlen
    nennen, sonst gibt es zwei Konformitätsaussagen zum selben Release — und die
    Abweichung fällt erst auf, wenn jemand die Dokumente nebeneinander legt.
    Deshalb erhebt diese Funktion, und :func:`build_docx` setzt nur noch.

    ``evaluate`` rechnet bewusst gegen die Produktiv-Datenbank des Moduls (sie
    liest ein Dutzend Register); ``db_path`` gilt für die Stammdaten.
    """
    from cra.db import load_bewertungen, load_projekt
    from cra.requirements import CRA_ANFORDERUNGEN, KAPITEL, PRODUKTKLASSEN
    from cra.konformitaet_gate import evaluate
    from cra import konformitaet_db as kdb

    projekt = load_projekt(db_path, projekt_name) or {}
    gate = evaluate(projekt_name, version=version)
    bew_raw = load_bewertungen(db_path, projekt_name)
    bew = {rid: d.get("bewertung", 0) for rid, d in bew_raw.items()}

    # Konformitätsbewertungsverfahren + Produktklasse
    konf = None
    try:
        for k in kdb.list_konformitaet(db_path, projekt_name):
            if not version or str(k.get("release_version") or "") == str(version):
                konf = k
                break
    except Exception:  # noqa: BLE001
        konf = None
    pklasse = (konf or {}).get("produktklasse") or projekt.get("produktklasse") or "default"
    pk_info = PRODUKTKLASSEN.get(pklasse, PRODUKTKLASSEN.get("default", {}))
    weg = (konf or {}).get("bewertungsweg") or "A"

    abw_by_id = {a["requirement_id"]: a for a in gate.get("abweichungen", [])}
    anforderungen: list[dict[str, Any]] = []
    for req in CRA_ANFORDERUNGEN:
        rid = req["id"]
        b = int(bew.get(rid, 0) or 0)
        if b >= 4:
            status, hinweis = "erfuellt", ""
        elif rid in abw_by_id and abw_by_id[rid]["art"] == "begruendet":
            status = "begruendet"
            a = abw_by_id[rid]
            hinweis = f"{a.get('kategorie')}: {a.get('grund') or ''}".strip(": ")
        else:
            status, hinweis = "offen", "nicht erfüllt / nicht begründet"
        kref = KAPITEL.get(req.get("kapitel"), {}).get("referenz", "")
        anforderungen.append({
            "id": rid,
            "titel": str(req.get("titel") or ""),
            "rechtsgrundlage": f"{req.get('ref', '')} · {kref}".strip(" ·"),
            "status": status,
            "status_label": _STATUS_LABEL[status],
            "hinweis": hinweis,
            "bewertung": b,
        })

    annex_vii: list[dict[str, Any]] = []
    annex_verfuegbar = True
    try:
        from cra.traceability_db import annex_vii_status
        for b in annex_vii_status(db_path, projekt_name).get("bausteine", []):
            annex_vii.append({"label": str(b.get("label") or ""),
                              "belegt": b.get("ampel") == "belegt",
                              "nachweis_count": int(b.get("nachweis_count", 0) or 0)})
    except Exception:  # noqa: BLE001 — Teilangabe darf den Nachweis nicht verhindern
        annex_verfuegbar = False

    return {
        "projekt_name": projekt_name,
        "version": str(version or (konf or {}).get("release_version") or ""),
        "hersteller": str(projekt.get("unternehmen") or ""),
        "produkt": str(projekt.get("produkt") or projekt_name),
        "produktklasse": str(pklasse),
        "produktklasse_label": str(pk_info.get("label") or pklasse),
        "produktklasse_referenz": str(pk_info.get("referenz") or ""),
        "bewertungsweg": str(weg),
        "bewertungsweg_label": _VERFAHREN.get(weg, weg),
        "rechtsgrundlage": _REGULATION,
        "ce_kennzeichnung": _ce_text(konf),
        "datum": date.today().isoformat(),
        "ampel": str(gate.get("ampel") or ""),
        "ampel_label": _AMPEL.get(gate.get("ampel"), str(gate.get("ampel") or "")),
        "konform": bool(gate.get("konform")),
        "reifegrad_pct": gate.get("reifegrad_pct"),
        "erfuellt": int(gate.get("erfuellt") or 0),
        "begruendet_akzeptiert": int(gate.get("begruendet_akzeptiert") or 0),
        "offen": int(gate.get("offen") or 0),
        "gesamt_anforderungen": int(gate.get("gesamt_anforderungen") or 0),
        "annex_vollstaendig": bool(gate.get("annex_vollstaendig")),
        "doc_ausgestellt": bool(gate.get("doc_ausgestellt")),
        "offene_kritische_schwachstellen": int(
            gate.get("offene_kritische_schwachstellen") or 0),
        "anforderungen": anforderungen,
        "abweichungen_begruendet": [a for a in gate.get("abweichungen", [])
                                    if a.get("art") == "begruendet"],
        "nicht_konformitaeten": [a for a in gate.get("abweichungen", [])
                                 if a.get("art") == "offen"],
        "annex_vii": annex_vii,
        "annex_vii_verfuegbar": annex_verfuegbar,
        "akzeptierte_restrisiken": _akzeptierte_restrisiken(projekt),
        "hinweis": HINWEIS_TEXT,
    }


def _akzeptierte_restrisiken(projekt: dict[str, Any]) -> list[dict[str, Any]]:
    """Formal akzeptierte Restrisiken der verknüpften Risikobewertung (#1674).

    Auch nicht-anforderungsbezogene Risiken sind hier nachgewiesen — eine
    Akzeptanz ohne Bezug zu einer Annex-I-Anforderung bleibt eine Entscheidung
    über das Produkt.
    """
    meta = projekt.get("meta") if isinstance(projekt.get("meta"), dict) else {}
    rb_name = str((meta or {}).get("linked_risk_projekt") or "").strip()
    if not rb_name:
        return []
    try:
        from risikobewertung.db import load_risiken
        rows = load_risiken(Path("data/db/risikobewertung.sqlite"), rb_name)
    except Exception:  # noqa: BLE001
        return []
    return [{"nr": r.get("nr"), "name": str(r.get("risk_name") or ""),
             "kategorie": str(r.get("akzeptanz_kategorie") or ""),
             "von": str(r.get("akzeptanz_von") or ""),
             "rechtsgrundlage": str(r.get("rechtsgrundlage") or ""),
             "grund": str(r.get("akzeptanz_grund") or "")}
            for r in rows if (r.get("behandlung") or "") == "akzeptiert"]


def kontext(db_path: Path | str, projekt_name: str,
            version: str | None = None) -> dict[str, Any]:
    """``nachweis``-Teilbaum für den CRA-Vorlagen-Kontext (#1777).

    Wie bei der Erklärung (``doc_export.kontext``) gilt: kein ``None``, kein
    fehlender Schlüssel. Jede in ``CRA_VARIABLES`` deklarierte Variable muss
    auch für ein Projekt ohne jede Erfassung ankommen — Jinja druckt ``None``
    wörtlich, und eine Vorlage darf nie an einer Teilangabe scheitern.
    """
    try:
        d = daten(projekt_name, version=version, db_path=Path(db_path))
    except Exception:  # noqa: BLE001
        d = {}
    leer = {
        "projekt_name": projekt_name, "version": "", "hersteller": "",
        "produkt": projekt_name, "produktklasse": "", "produktklasse_label": "",
        "produktklasse_referenz": "", "bewertungsweg": "", "bewertungsweg_label": "",
        "rechtsgrundlage": _REGULATION, "ce_kennzeichnung": "",
        "datum": date.today().isoformat(), "ampel": "", "ampel_label": "",
        "konform": False, "reifegrad_pct": 0, "erfuellt": 0,
        "begruendet_akzeptiert": 0, "offen": 0, "gesamt_anforderungen": 0,
        "annex_vollstaendig": False, "doc_ausgestellt": False,
        "offene_kritische_schwachstellen": 0, "anforderungen": [],
        "abweichungen_begruendet": [], "nicht_konformitaeten": [],
        "annex_vii": [], "annex_vii_verfuegbar": False,
        "akzeptierte_restrisiken": [], "hinweis": HINWEIS_TEXT,
    }
    out = {**leer, **{k: v for k, v in d.items() if v is not None}}
    if out.get("reifegrad_pct") is None:
        out["reifegrad_pct"] = 0
    return out


def build_docx(projekt_name: str, *, version: str | None = None,
               db_path: Path = CRA_DB) -> bytes:
    """Der programmatische Nachweis. Setzt nur, erhebt nicht (siehe :func:`daten`)."""
    from shared.reports.docx_pdf import set_a4

    d = daten(projekt_name, version=version, db_path=db_path)

    doc = Document(); set_a4(doc)

    _h(doc, "EU-Konformitätsnachweis (Cyber Resilience Act)", size=18, color=_EU_BLUE,
       space_before=0)
    _p(doc, _REGULATION)

    kopf = [
        ("Hersteller", d["hersteller"] or "–"),
        ("Produkt", d["produkt"]),
        ("Version / Release", d["version"] or "–"),
        ("Produktklasse",
         f"{d['produktklasse_label']} — {d['produktklasse_referenz']}".strip(" —")),
        ("Konformitätsbewertungsverfahren", d["bewertungsweg_label"]),
        ("Rechtsgrundlage", d["rechtsgrundlage"]),
        ("CE-Kennzeichnung", d["ce_kennzeichnung"]),
        ("Datum", d["datum"]),
    ]
    t = doc.add_table(rows=len(kopf), cols=2); t.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate(kopf):
        t.rows[i].cells[0].text = k; t.rows[i].cells[1].text = str(v)

    # Gesamturteil
    _h(doc, "Gesamturteil (Konformitäts-Gate)", color=_EU_BLUE)
    _p(doc, f"Ergebnis: {d['ampel_label']}  ·  Reifegrad {d['reifegrad_pct']} %  ·  "
            f"erfüllt {d['erfuellt']}, begründet akzeptiert {d['begruendet_akzeptiert']}, "
            f"offen {d['offen']} von {d['gesamt_anforderungen']} Anforderungen.")
    _p(doc, f"Annex VII vollständig: {'ja' if d['annex_vollstaendig'] else 'nein'}  ·  "
            f"EU-Konformitätserklärung ausgestellt: "
            f"{'ja' if d['doc_ausgestellt'] else 'nein'}  ·  "
            f"offene kritische/hohe Schwachstellen: "
            f"{d['offene_kritische_schwachstellen']}.")

    # Anforderungsstatus je Kapitel (mit Rechtsgrundlage)
    _h(doc, "Erfüllung der essenziellen Anforderungen (Annex I) — mit Rechtsgrundlage",
       color=_EU_BLUE)
    tbl = doc.add_table(rows=1, cols=4); tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    hdr[0].text = "Anforderung"; hdr[1].text = "Rechtsgrundlage"
    hdr[2].text = "Status"; hdr[3].text = "Hinweis"
    for a in d["anforderungen"]:
        row = tbl.add_row().cells
        row[0].text = f"{a['id']} — {a['titel']}"
        row[1].text = a["rechtsgrundlage"]
        row[2].text = a["status_label"]
        row[3].text = a["hinweis"]

    # Begründete Abweichungen (akzeptierte Restrisiken)
    _h(doc, "Begründete Abweichungen (als Risiko akzeptiert)", color=_EU_BLUE)
    if not d["abweichungen_begruendet"]:
        _p(doc, "Keine — alle anwendbaren Anforderungen sind erfüllt.")
    else:
        for a in d["abweichungen_begruendet"]:
            _p(doc, f"• {a['requirement_id']} — {a.get('titel', '')}", size=10)
            _p(doc, f"   Kategorie: {a.get('kategorie')} · akzeptiert von: "
                    f"{a.get('von') or '—'} · Rechtsgrundlage: "
                    f"{a.get('rechtsgrundlage') or '—'}", size=9)
            _p(doc, f"   Begründung: {a.get('grund') or '—'}", size=9)

    # Offene Nicht-Konformitäten
    if d["nicht_konformitaeten"]:
        _h(doc, "Offene Nicht-Konformitäten (Korrekturmaßnahmen erforderlich)",
           color=RGBColor(0xC6, 0x28, 0x28))
        for a in d["nicht_konformitaeten"]:
            _p(doc, f"• {a['requirement_id']} — {a.get('titel', '')} "
                    f"({a.get('ref', '')})", size=10)

    # Annex VII
    _h(doc, "Technische Dokumentation (Annex VII)", color=_EU_BLUE)
    if not d["annex_vii_verfuegbar"]:
        _p(doc, "Annex-VII-Status nicht verfügbar.", size=9)
    for b in d["annex_vii"]:
        mark = "✓" if b["belegt"] else "✗"
        _p(doc, f"{mark} {b['label']} ({b['nachweis_count']} Nachweis(e))", size=9)

    _h(doc, "Akzeptierte Restrisiken (Risikobewertung)", color=_EU_BLUE)
    if not d["akzeptierte_restrisiken"]:
        _p(doc, "Keine formal akzeptierten Restrisiken.")
    else:
        for r in d["akzeptierte_restrisiken"]:
            _p(doc, f"• #{r['nr']} {r['name']}", size=10)
            _p(doc, f"   Kategorie: {r['kategorie'] or '—'} · akzeptiert von: "
                    f"{r['von'] or '—'} · Rechtsgrundlage: "
                    f"{r['rechtsgrundlage'] or '—'}", size=9)
            if r["grund"]:
                _p(doc, f"   Begründung: {r['grund']}", size=9)

    _h(doc, "Hinweis", color=_EU_BLUE)
    _p(doc, d["hinweis"], size=9)

    bio = BytesIO(); doc.save(bio); return bio.getvalue()


def build_pdf(projekt_name: str, *, version: str | None = None,
              db_path: Path = CRA_DB) -> bytes:
    from shared.reports.docx_pdf import docx_bytes_to_pdf
    return docx_bytes_to_pdf(build_docx(projekt_name, version=version,
                                        db_path=db_path))


def build(projekt_name: str, *, fmt: str = "docx", version: str | None = None,
          db_path: Path = CRA_DB) -> bytes:
    """Format-Weiche für das Berichts-Center (#1776) — Muster ``doc_export.build``."""
    if (fmt or "docx").lower() == "pdf":
        return build_pdf(projekt_name, version=version, db_path=db_path)
    return build_docx(projekt_name, version=version, db_path=db_path)
