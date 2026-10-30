"""Vorgangsbericht je Release — für Management und Ablage (#1748).

**Abgrenzung, die diese Datei trägt.** Der juristische Konformitätsnachweis ist
:mod:`cra.konformitaet_report` (Anforderungen mit Rechtsgrundlage, Verfahren nach
Anhang VIII, begründete Abweichungen, Anhang VII). Dieser Bericht hier ist das
andere Stück: Er beschreibt den **Vorgang** — wer hat was wann festgestellt,
geprüft, freigegeben, und worauf stützt sich das. Er **zitiert** den Nachweis
(Archivverweis mit SHA-256) und erzeugt ihn nicht neu. Zwei Renderer für dasselbe
Dokument wären zwei Wahrheiten; genau das war die Lehre aus #1727.

Drei Entscheidungen, die den Inhalt prägen:

* **Gerechnet und erklärt bleiben getrennt.** Die gerechneten Prüfpunkte kommen
  aus :func:`cra.konformitaet_gate.evaluate`, die Entscheidungen aus der
  Release-Checkliste des Vorgangs. Der Bericht rechnet nichts selbst nach.
* **Der eingefrorene Stand hat Vorrang.** Ist der Vorgang eingereicht, wurde
  gegen den damals eingefrorenen Rechenstand (``gate_json``) geprüft. Genau der
  steht im Bericht; der heutige Stand wird daneben genannt, wenn er abweicht.
  Andernfalls behauptete der Bericht, es sei etwas geprüft worden, was heute
  anders aussieht.
* **Nicht geprüft ist nicht erfüllt.** Ein Befund, den die Anwendung nicht
  ermitteln konnte (``None``), erscheint als „nicht geprüft" und nie als Haken.
  Ein stiller Rückfall auf 0/„ok" wäre in einem Nachweis der gefährlichste
  Fehler, weil er wie ein gutes Ergebnis aussieht (#1741/#1751).

Ein Vorgang **ohne** Freigabe ist der Normalfall im Tagesgeschäft — man braucht
den Bericht, um ihn prüfen zu lassen. Er wird deshalb erzeugt und sichtbar als
**Entwurf** gekennzeichnet, nicht verweigert.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any, Callable

DB_PATH = Path("data/db/cra.sqlite")
RB_DB = Path("data/db/risikobewertung.sqlite")

_REGULATION = "Verordnung (EU) 2024/2847 (Cyber Resilience Act)"

AMPEL_LABEL = {"gruen": "KONFORM", "bedingt": "BEDINGT KONFORM", "rot": "NICHT KONFORM"}
STATUS_LABEL = {"entwurf": "Entwurf", "in_pruefung": "in Prüfung",
                "freigegeben": "freigegeben", "zurueckgewiesen": "zurückgewiesen",
                "veroeffentlicht": "veröffentlicht"}
ART_LABEL = {"gerechnet": "Gerechnete Prüfpunkte",
             "entscheidung": "Entscheidungen (mit Begründung und Person)",
             "gleichbleibend": "Gleichbleibende Angaben (je Release bestätigt)"}

#: Zustände, in denen der Vorgang kein Entwurf mehr ist.
FREIGEGEBEN = ("freigegeben", "veroeffentlicht")

#: Kennzeichnung eines nicht freigegebenen Vorgangs — sie steht im Titel, im Band
#: und in der Kopftabelle. Ein Entwurf, der wie ein Nachweis aussieht, ist der
#: Grund, warum später niemand mehr weiß, welche Fassung gilt.
ENTWURF_MARKER = "ENTWURF — nicht freigegeben"

# ── Farben (#1727: Farbe folgt dem Erfüllungsgrad, nicht der Dekoration) ────
FARBEN = {
    "kopf": "1565C0",
    "gruen": "2E7D32",
    "bedingt": "F9A825",
    "rot": "C62828",
    "unbekannt": "90A4AE",
}

#: Zellenhintergrund je Ergebnis — hell, damit der Text lesbar bleibt.
ERG_FARBE = {
    "erfuellt": "E8F5E9",
    "bestaetigt": "E8F5E9",
    "offen": "FFEBEE",
    "maengel": "FFF3E0",
    "ungeprueft": "ECEFF1",
    "nicht_anwendbar": "F5F5F5",
}


# ── DOCX-Kleinteile ─────────────────────────────────────────────────────────
# Dieselben Helfer stehen in ``compliance_audit/berichte_provider.py``. Sie
# gehören fachlich nach ``shared/reports/docx_pdf.py``; das ist in diesem Vorgang
# fremder Arbeitsbereich, deshalb hier lokal (siehe Bericht zu #1748).

def _schattieren(zelle, hex_farbe: str) -> None:
    """Zellenhintergrund setzen — python-docx kann das nur über w:shd."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_farbe)
    zelle._tc.get_or_add_tcPr().append(shd)


def _weiss_fett(zelle, text: str) -> None:
    from docx.shared import RGBColor

    run = zelle.paragraphs[0].add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _tabelle(doc, kopf: list[str], zeilen: int):
    tbl = doc.add_table(rows=zeilen + 1, cols=len(kopf))
    tbl.style = "Table Grid"
    for ci, titel in enumerate(kopf):
        z = tbl.cell(0, ci)
        _schattieren(z, FARBEN["kopf"])
        _weiss_fett(z, titel)
    return tbl


def _ueberschrift(doc, text: str, level: int = 1):
    from docx.shared import RGBColor

    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)
    return h


def _band(doc, farbe: str, zeile1: str, zeile2: str = "") -> None:
    """Farbiges Band — das Erste, was man sehen soll: Ampel bzw. Entwurfs-Hinweis."""
    from docx.shared import Pt, RGBColor

    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    z = tbl.cell(0, 0)
    _schattieren(z, farbe)
    r = z.paragraphs[0].add_run(zeile1)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if zeile2:
        r2 = z.add_paragraph().add_run(zeile2)
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    doc.add_paragraph()


def _datum(wert: Any) -> str:
    """Zeitstempel kurz und lesbar; leer bleibt leer (nie „None")."""
    s = str(wert or "").strip()
    return s[:19].replace("T", " ") if s else ""


def _person_am(von: Any, am: Any) -> str:
    """„von X am Y" — der Satz, den ein Nachweis braucht."""
    p = str(von or "").strip()
    d = _datum(am)
    if p and d:
        return f"{p} am {d}"
    return p or d or "—"


# ── Daten sammeln ───────────────────────────────────────────────────────────

def _aktivitaet(v: dict[str, Any]) -> str:
    """Zeitpunkt der letzten Bewegung eines Vorgangs (für „der jüngste")."""
    return max(str(v.get(k) or "") for k in (
        "veroeffentlicht_am", "freigegeben_am", "geprueft_am", "eingereicht_am",
        "updated_at", "created_at"))


def waehle_vorgang(vorgaenge: list[dict[str, Any]],
                   version: str | None = None) -> dict[str, Any] | None:
    """Welcher Release-Vorgang wird berichtet?

    Das Berichts-Center kennt nur Modul, Typ und Projekt — eine Release-Version
    reicht es nicht durch. Die Auswahl muss also hier fallen, und sie muss
    vorhersagbar sein: der **jüngste** versionierte Vorgang, also der, an dem
    gerade gearbeitet wird.

    Der Datensatz mit **leerer** Version ist bewusst außen vor: Das ist der
    Produkt-Datensatz (Bewertungsweg, Notified Body, CE) aus Phase 1 und kein
    Release-Vorgang.
    """
    versioniert = [v for v in vorgaenge if str(v.get("release_version") or "").strip()]
    if version:
        ziel = str(version).strip()
        return next((v for v in versioniert
                     if str(v.get("release_version") or "").strip() == ziel), None)
    if not versioniert:
        return None
    return max(versioniert, key=_aktivitaet)


def sammeln(projekt: str, version: str | None = None,
            db_path: Path = DB_PATH) -> dict[str, Any]:
    """Alle Angaben des Vorgangsberichts aus den vorhandenen Quellen holen.

    Jede Quelle ist einzeln gekapselt: Fällt eine aus, fehlt ihr Abschnitt mit
    einem lesbaren Hinweis, statt den ganzen Bericht zu verhindern. Erfunden
    wird nichts — fehlt ein Wert, steht das im Bericht.
    """
    from cra import konformitaet_db as kdb
    from cra import konformitaet_final_db as fdb
    from cra import release_db as rdb
    from cra.konformitaet_gate import evaluate

    daten: dict[str, Any] = {"projekt_name": projekt, "hinweise": []}

    def _quelle(name: str, fn: Callable[[], Any], default: Any = None) -> Any:
        try:
            return fn()
        except Exception as exc:  # noqa: BLE001 — Nebenquelle darf den Bericht nicht kippen
            daten["hinweise"].append(f"{name} konnte nicht geladen werden "
                                     f"({type(exc).__name__}: {exc})")
            return default

    from cra.db import load_projekt
    proj = _quelle("Projekt-Stammdaten", lambda: load_projekt(db_path, projekt)) or {}
    daten["projekt"] = proj
    meta = proj.get("meta") if isinstance(proj.get("meta"), dict) else {}

    vorgaenge = _quelle("Release-Vorgänge",
                        lambda: kdb.list_konformitaet(db_path, projekt), []) or []
    vorgang = waehle_vorgang(vorgaenge, version)
    daten["vorgaenge"] = vorgaenge
    daten["vorgang"] = vorgang
    rv = str((vorgang or {}).get("release_version") or version or "")
    daten["release_version"] = rv
    status = str((vorgang or {}).get("freigabe_status") or "entwurf")
    daten["status"] = status
    # Ohne Freigabe ist es ein Entwurf — auch wenn gar kein Vorgang erfasst ist.
    daten["ist_entwurf"] = status not in FREIGEGEBEN

    # Heutiger Rechenstand: Release-Sicht (mit Version) und Produkt-Sicht.
    # Die Produkt-Sicht wird getrennt gebraucht, weil ``bereit_ueberholt`` die
    # Phase-1-Feststellung bewertet — ein offener Release-Punkt würde sie sonst
    # fälschlich als überholt ausweisen.
    daten["gate_heute"] = _quelle(
        "Konformitäts-Gate", lambda: evaluate(projekt, version=rv or None), {}) or {}
    daten["gate_produkt"] = _quelle(
        "Konformitäts-Gate (Produktsicht)", lambda: evaluate(projekt), {}) or {}

    # Eingefrorener Stand bei Einreichung — er ist die Grundlage der Prüfung.
    eingefroren = (vorgang or {}).get("gate") or {}
    daten["gate_eingefroren"] = eingefroren if isinstance(eingefroren, dict) else {}
    nutze_eingefroren = bool(daten["gate_eingefroren"]) and status != "entwurf"
    daten["gate_basis"] = (daten["gate_eingefroren"] if nutze_eingefroren
                           else daten["gate_heute"])
    daten["gate_basis_quelle"] = ("eingefroren bei Einreichung" if nutze_eingefroren
                                  else "heutiger Rechenstand")

    # Phase 1 — CRA-Ready als ausdrückliche Feststellung (#1746).
    phase1 = _quelle("Phase-1-Feststellung", lambda: rdb.get_release(db_path, projekt))
    daten["phase1"] = phase1 or {}
    daten["phase1_ueberholt"] = bool(_quelle(
        "Phase-1-Feststellung", lambda: rdb.bereit_ueberholt(phase1, daten["gate_produkt"]),
        False))

    # Archivierte Nachweise (SHA-256) — der Bericht zeigt darauf, er ersetzt sie nicht.
    archiv = _quelle("Nachweis-Archiv", lambda: fdb.list_final(db_path, projekt), []) or []
    daten["archiv_alle"] = archiv
    daten["archiv"] = [a for a in archiv
                       if not rv or str(a.get("version") or "") == rv] if archiv else []

    daten["restrisiken"] = _restrisiken(daten["gate_basis"], meta, daten["hinweise"])
    return daten


def _restrisiken(gate: dict[str, Any], meta: dict[str, Any],
                 hinweise: list[str]) -> dict[str, Any]:
    """Akzeptierte Restrisiken — zwei Herkünfte, bewusst getrennt gehalten.

    Maßgeblich für die Konformität sind die **Abweichungen** aus dem Gate: eine
    nicht erfüllte Anforderung gilt nur dann als konform, wenn das verknüpfte
    Nichterfüllungs-Risiko begründet akzeptiert ist (#1613). Daneben kann die
    verknüpfte Risikobewertung weitere akzeptierte Risiken führen, die an keiner
    Anforderung hängen; die stehen als eigener Block, damit nicht der Eindruck
    entsteht, sie hätten dieselbe Wirkung.
    """
    abw = [a for a in (gate.get("abweichungen") or [])
           if str(a.get("art")) == "begruendet"]
    offen = [a for a in (gate.get("abweichungen") or []) if str(a.get("art")) == "offen"]
    weitere: list[dict[str, Any]] = []
    rb_name = str((meta or {}).get("linked_risk_projekt") or "").strip()
    if rb_name:
        try:
            from risikobewertung.db import load_risiken
            belegt = {f"cra:requirement:{a.get('requirement_id')}" for a in abw}
            for r in load_risiken(RB_DB, rb_name):
                if str(r.get("behandlung") or "") != "akzeptiert":
                    continue
                felder = r.get("felder") or {}
                if str(felder.get("_source_id") or "") in belegt:
                    continue   # steht schon als Abweichung — keine Dublette
                weitere.append({
                    "nr": r.get("nr"), "titel": r.get("risk_name"),
                    "kategorie": r.get("akzeptanz_kategorie"),
                    "grund": r.get("akzeptanz_grund"), "von": r.get("akzeptanz_von"),
                    "am": r.get("akzeptanz_am"), "review": r.get("akzeptanz_review"),
                    "rechtsgrundlage": r.get("rechtsgrundlage"),
                    "level": r.get("risk_level")})
        except Exception as exc:  # noqa: BLE001
            hinweise.append(f"Verknüpfte Risikobewertung '{rb_name}' nicht lesbar "
                            f"({type(exc).__name__}: {exc})")
    return {"abweichungen": abw, "offene_abweichungen": offen, "weitere": weitere,
            "rb_projekt": rb_name}


# ── Gerechnete Prüfpunkte auswerten ─────────────────────────────────────────
# Die Zuordnung Katalog-``quelle`` → Gate-Schlüssel steht im Katalog (#1744);
# hier wird der **Wert** nur in eine Aussage übersetzt. ``None`` heißt überall
# „nicht geprüft" und wird nie zu „erfüllt".

def _aus_zahl(wert: Any, *, ok_text: str, fehl: str) -> tuple[str, str]:
    if wert is None:
        return "ungeprueft", "nicht geprüft"
    try:
        n = int(wert)
    except (TypeError, ValueError):
        return "ungeprueft", f"nicht auswertbar ({wert!r})"
    return ("erfuellt", ok_text) if n == 0 else ("offen", fehl.format(n=n))


def _aus_liste(wert: Any, *, ok_text: str, fehl: str) -> tuple[str, str]:
    if wert is None:
        return "ungeprueft", "nicht geprüft"
    if isinstance(wert, (list, tuple)):
        return ("erfuellt", ok_text) if not wert else ("offen", fehl.format(n=len(wert)))
    return _aus_zahl(wert, ok_text=ok_text, fehl=fehl)


def _aus_mangel_flag(wert: Any, *, ok_text: str, fehl: str) -> tuple[str, str]:
    """Flag, bei dem ``True`` den **Mangel** meldet (``sbom_version_fehlt``)."""
    if wert is None:
        return "ungeprueft", "nicht geprüft"
    return ("offen", fehl) if bool(wert) else ("erfuellt", ok_text)


def _aus_ok_flag(wert: Any, *, ok_text: str, fehl: str) -> tuple[str, str]:
    """Flag, bei dem ``True`` die **Erfüllung** meldet (``annex_vollstaendig``)."""
    if wert is None:
        return "ungeprueft", "nicht geprüft"
    return ("erfuellt", ok_text) if bool(wert) else ("offen", fehl)


_GERECHNET_AUSWERTUNG: dict[str, Callable[[Any], tuple[str, str]]] = {
    "offene_kritische_schwachstellen": lambda w: _aus_zahl(
        w, ok_text="keine offenen kritischen oder hohen Befunde",
        fehl="{n} offene kritische/hohe Schwachstelle(n)"),
    "behobene_ohne_advisory": lambda w: _aus_liste(
        w, ok_text="alle in dieser Version behobenen Schwachstellen sind veröffentlicht",
        fehl="{n} behobene Schwachstelle(n) ohne Advisory"),
    "sbom_version_fehlt": lambda w: _aus_mangel_flag(
        w, ok_text="SBOM zur Release-Version vorhanden",
        fehl="keine SBOM zu dieser Release-Version"),
    "support_ueberschritten": lambda w: _aus_mangel_flag(
        w, ok_text="Release liegt im zugesagten Support-Zeitraum",
        fehl="Release liegt nach dem Support-Ende"),
    "ueberfaellige_meldungen": lambda w: _aus_liste(
        w, ok_text="keine überfällige Meldefrist",
        fehl="{n} überfällige Meldefrist(en)"),
    "annex_vollstaendig": lambda w: _aus_ok_flag(
        w, ok_text="technische Dokumentation vollständig belegt",
        fehl="Nachweise der technischen Dokumentation fehlen"),
    "doc_ausgestellt": lambda w: _aus_ok_flag(
        w, ok_text="EU-Konformitätserklärung ausgestellt bzw. freigegeben",
        fehl="EU-Konformitätserklärung nicht ausgestellt"),
    # #1759: Der Bewertungsweg muss zur Produktklasse passen. Ein unzulässiger
    # Weg ist keine Lücke in der Akte, sondern eine Nicht-Konformität — im
    # Bericht darf er deshalb nicht als bloß „offener Punkt" erscheinen.
    "bewertungsweg_unzulaessig": lambda w: _aus_mangel_flag(
        w, ok_text="Konformitätsbewertungsweg für die Produktklasse zulässig",
        fehl="Konformitätsbewertungsweg ist für diese Produktklasse NICHT zulässig"),
    # #1761: Sorgfaltspflicht bei Komponenten Dritter und die Meldung an deren
    # Hersteller. Beide Punkte waren vorher keine Rechenwerte — der eine ein
    # Sammelhaken, der andere gar nicht vorhanden. ``None`` heißt hier wie
    # überall „nicht geprüft": keine SBOM, keine Komponentenliste, oder die
    # Quelle fiel aus. Ein ``0`` an dieser Stelle wäre die Aussage, es sei
    # geprüft und in Ordnung.
    "fremdkomponenten_unbewertet": lambda w: _aus_liste(
        w, ok_text="neue und angehobene Fremdkomponenten sind bewertet",
        fehl="{n} neu aufgenommene oder angehobene Fremdkomponente(n) ohne "
             "dokumentierte Bewertung"),
    # #1774: Der Widerspruch zwischen Bewertung und eigenem Befund. ``None``
    # heißt auch hier „nicht geprüft" — ohne Komponentenliste ist nichts zu
    # vergleichen, und ein „0" wäre die Aussage, es sei geprüft und sauber.
    "komponenten_befund_widerspruch": lambda w: _aus_liste(
        w, ok_text="keine Bewertung widerspricht einem offenen Befund",
        fehl="{n} Fremdkomponente(n) als „nicht relevant“ bewertet trotz "
             "offenem Befund"),
    # #1793: Der Abgleich der erklärten Release-Art mit der Akte. ``None`` =
    # nicht geprüft (Phase 1 oder Quellenausfall), leere Liste = geprüft und
    # widerspruchsfrei. Eine nicht erklärte Art erzeugt keinen Widerspruch —
    # dass sie fehlt, steht im Abschnitt zur wesentlichen Änderung.
    "release_art_widersprueche": lambda w: _aus_liste(
        w, ok_text="die erklärte Art des Releases passt zur Akte",
        fehl="{n} Widerspruch/Widersprüche zwischen erklärter Release-Art und "
             "dem, was in der Akte steht"),
    "upstream_meldung_offen": lambda w: _aus_liste(
        w, ok_text="eigene Funde in Fremdkomponenten sind dem jeweiligen "
                   "Hersteller gemeldet",
        fehl="{n} Schwachstelle(n) in Fremdkomponenten ohne dokumentierte "
             "Meldung an deren Hersteller"),
    # #1762: Vier gerechnete Befunde, die bis v8.0.1 hier fehlten. Ohne Eintrag
    # liefert ``gerechnetes_ergebnis`` „nicht geprüft" — ein Punkt, der gerechnet
    # wurde, sähe im Bericht aus wie einer, den niemand angesehen hat.
    "update_zusagen_fehlen": lambda w: _aus_mangel_flag(
        w, ok_text="Sicherheitsupdates kostenlos und über einen getrennten Kanal "
                   "zugesagt",
        fehl="Update-Zusagen nicht vollständig dokumentiert"),
    "kaufhinweis_fehlt": lambda w: _aus_mangel_flag(
        w, ok_text="Support-Ende ist vor dem Kauf zugänglich",
        fehl="kein Hinweis auf das Support-Ende beim Kauf"),
    "aufbewahrungsfristen_unbestimmbar": lambda w: _aus_mangel_flag(
        w, ok_text="Aufbewahrungs- und Update-Fristen sind bestimmbar",
        fehl="Aufbewahrungs-/Update-Fristen nicht bestimmbar (Bezugsdaten fehlen)"),
    "marktaufsicht_versaeumt": lambda w: _aus_zahl(
        w, ok_text="keine versäumte Frist gegenüber der Marktüberwachung",
        fehl="{n} versäumte Frist(en) gegenüber der Marktüberwachung"),
    # #1824: Anhang I Teil II Nr. 3. ``None`` = nicht geprüft, leere Liste =
    # ein Testnachweis liegt vor und belegt den ausgelieferten Stand. Der
    # erklärte Verzicht erscheint hier als **offener** Punkt, und das ist der
    # Zweck: Er ist eine Entscheidung des Herstellers, die der Nachweis nennen
    # muss, und keine erfüllte Pflicht.
    "sicherheitstests_mangel": lambda w: _aus_liste(
        w, ok_text="Sicherheitstests belegen den ausgelieferten Stand",
        fehl="{n} Feststellung(en) zum regelmäßigen und wirksamen Testen"),
    # #1884: Anhang V Nr. 6. ``None`` = nicht geprüft, leere Liste = die
    # Erklärung trifft eine Aussage und sie passt zum Bewertungsweg.
    "normen_widerspruch": lambda w: _aus_liste(
        w, ok_text="Anhang V Nr. 6 ist erklärt und passt zum Bewertungsweg",
        fehl="{n} Feststellung(en) zu den angewandten Normen (Anhang V Nr. 6)"),
}


def gerechnetes_ergebnis(quelle: str, gate: dict[str, Any]) -> tuple[str, str]:
    """(Ergebnis-Schlüssel, Klartext) eines gerechneten Punktes.

    Eine ``quelle``, die das Gate nicht kennt, ist ein Fehler im Katalog und
    wird als solcher gezeigt — nicht als erfüllter Punkt.
    """
    auswerten = _GERECHNET_AUSWERTUNG.get(quelle)
    if auswerten is None:
        return "ungeprueft", f"kein Rechenergebnis zur Quelle '{quelle}'"
    if quelle not in gate:
        return "ungeprueft", "nicht geprüft (kein Rechenstand vorhanden)"
    return auswerten(gate.get(quelle))


def wesentliche_zeile(vorgang: dict[str, Any], eintrag: dict[str, Any],
                      item) -> tuple[str, str, str, str]:
    """(Ergebnis, Label, Begründung, Person) der Kernfrage nach Art. 3 Nr. 30.

    Dieser Punkt steht im Katalog wie jede andere Entscheidung, wird aber über
    **eigene Spalten** gepflegt (``set_wesentliche_aenderung``, #1744). Wer nur
    den Checklisten-Eintrag liest, zeigt die Frage als „offen“, obwohl sie
    beantwortet und begründet ist — genau die Verwechslung, die ein Prüfer dem
    Bericht anlasten würde.

    Die Spalten halten **keine Person und keinen Zeitpunkt** fest. Das wird hier
    benannt und nicht ersetzt: eine erfundene Person wäre schlimmer als eine
    fehlende.
    """
    wert = str(vorgang.get("wesentliche_aenderung") or "").strip()
    grund = (str(vorgang.get("wesentliche_aenderung_grund") or "").strip()
             or str((eintrag or {}).get("begruendung") or "").strip())
    person = _person_am((eintrag or {}).get("von"), (eintrag or {}).get("am"))
    if not wert:
        erg, label = entscheidungs_ergebnis(item, eintrag)
        return erg, label, grund or "—", person
    if not grund:
        return "maengel", "ohne Begründung — nicht nachgewiesen", "—", person
    if person == "—":
        person = "— (nicht erfasst)"
    return "erfuellt", f"beantwortet: {wert}", grund, person


def entscheidungs_ergebnis(item, stand: dict[str, Any]) -> tuple[str, str]:
    """(Ergebnis-Schlüssel, Klartext) einer Entscheidung.

    Der wichtigste Fall ist der dritte: Ein „erledigt" **ohne** Begründung darf
    im Bericht nicht als erledigt erscheinen. Sonst trägt der Nachweis einen
    Haken, hinter dem keine Aussage steht — und genau das soll die Trennung in
    drei Arten (#1744) verhindern.
    """
    from cra import release_checkliste as rc

    status = str((stand or {}).get("status") or "offen")
    begruendung = str((stand or {}).get("begruendung") or "").strip()
    if status == "offen":
        return "offen", "offen"
    braucht_grund = item.id in rc.BEGRUENDUNG_PFLICHT
    if braucht_grund and not begruendung:
        return "maengel", "ohne Begründung — nicht nachgewiesen"
    label = "erledigt" if status == "erledigt" else "nicht anwendbar"
    return ("erfuellt" if status == "erledigt" else "nicht_anwendbar"), label


# ── Gleichbleibende Angaben: Stand aus dem Gate ─────────────────────────────

def _stand_support(gate: dict[str, Any]) -> str:
    sup = (gate.get("release_befunde") or {}).get("support")
    if not isinstance(sup, dict):
        return "Stand nicht ausgewiesen"
    eol = sup.get("eol") or "nicht festgelegt"
    jahre = sup.get("support_jahre")
    return (f"Support-Ende {eol}"
            + (f", zugesagt {jahre} Jahr(e)" if jahre else "")
            + (f", Markteintritt {sup.get('markteintritt')}"
               if sup.get("markteintritt") else ""))


def _stand_kanal(gate: dict[str, Any]) -> str:
    sup = (gate.get("release_befunde") or {}).get("support")
    if not isinstance(sup, dict):
        return "Stand nicht ausgewiesen"
    return sup.get("update_kanal") or "nicht dokumentiert"


def _stand_psirt(gate: dict[str, Any]) -> str:
    vh = (gate.get("release_befunde") or {}).get("vuln_handling")
    if not isinstance(vh, dict):
        return "Stand nicht ausgewiesen"
    return (f"{vh.get('psirt_kontakte', 0)} Kontakt(e), Offenlegungs-Richtlinie "
            f"{'hinterlegt' if vh.get('policy_url') else 'fehlt'}")


_GLEICHBLEIBEND_STAND: dict[str, Callable[[dict[str, Any]], str]] = {
    "support_zeitraum": _stand_support,
    "update_kanal": _stand_kanal,
    "psirt_kontakt": _stand_psirt,
    "disclosure_policy": _stand_psirt,
}


# ── Dokument ────────────────────────────────────────────────────────────────

def build_docx(projekt: str, *, version: str | None = None,
               db_path: Path = DB_PATH) -> bytes:
    """Vorgangsbericht als DOCX-Bytes. Erzeugt auch für einen Entwurf."""
    from docx import Document

    from shared.reports.docx_pdf import apply_base_font, fit_tables, set_a4

    from cra import release_checkliste as rc

    d = sammeln(projekt, version, db_path=db_path)
    gate = d["gate_basis"]
    vorgang = d["vorgang"] or {}
    proj = d["projekt"]
    rv = d["release_version"]

    doc = Document()
    set_a4(doc)
    apply_base_font(doc)

    titel = "CRA-Vorgangsbericht je Release"
    if d["ist_entwurf"]:
        titel += " (Entwurf)"
    _ueberschrift(doc, titel, level=0)
    doc.add_paragraph(f"{_REGULATION} · Berichtsdatum {date.today().strftime('%d.%m.%Y')}")

    # Entwurf zuerst sagen, dann die Ampel: Wer das Band liest, muss wissen,
    # dass noch niemand freigegeben hat, bevor er die Farbe deutet.
    if d["ist_entwurf"]:
        _band(doc, FARBEN["unbekannt"], ENTWURF_MARKER,
              f"Vorgangsstatus: {STATUS_LABEL.get(d['status'], d['status'])} — dieser "
              "Bericht dient der Prüfung und ist kein freigegebener Nachweis.")
    ampel = str(gate.get("ampel") or "")
    _band(doc, FARBEN.get(ampel, FARBEN["unbekannt"]),
          f"Ampel: {AMPEL_LABEL.get(ampel, 'NICHT BEWERTET')}",
          f"Grundlage: {d['gate_basis_quelle']}"
          + (f" · Reifegrad {gate.get('reifegrad_pct')} %"
             if gate.get("reifegrad_pct") is not None else ""))

    _kopf_tabelle(doc, d, proj, vorgang, rv)
    _rechenstand_hinweis(doc, d)
    _phase1(doc, d)
    _wesentliche_aenderung(doc, d, vorgang)
    _checkliste(doc, d, gate, vorgang, rc)
    _offene_punkte(doc, gate)
    _restrisiken_abschnitt(doc, d)
    _vier_augen(doc, d, vorgang)
    _archiv(doc, d)
    _weitere_vorgaenge(doc, d)

    if d["hinweise"]:
        _ueberschrift(doc, "Hinweise zur Datenlage", level=1)
        doc.add_paragraph(
            "Die folgenden Quellen waren beim Erzeugen nicht lesbar. Was sie "
            "belegen sollten, ist damit nicht nachgewiesen — nicht erfüllt.")
        for h in d["hinweise"]:
            doc.add_paragraph(str(h), style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph(
        "Abgrenzung: Dieser Bericht beschreibt den Release-Vorgang. Der juristische "
        "Konformitätsnachweis (Anforderungen mit Rechtsgrundlage, Verfahren nach "
        "Anhang VIII, begründete Abweichungen, Anhang VII) ist ein eigenes Dokument; "
        "er wird oben über das Archiv mit seinem SHA-256-Wert zitiert und hier nicht "
        "erneut erzeugt."
    ).italic = True

    fit_tables(doc)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _kopf_tabelle(doc, d: dict[str, Any], proj: dict[str, Any],
                  vorgang: dict[str, Any], rv: str) -> None:
    zeilen = [
        ("Hersteller", proj.get("unternehmen") or "—"),
        ("Produkt", proj.get("produkt") or d["projekt_name"]),
        ("CRA-Projekt", d["projekt_name"]),
        ("Release-Version", rv or "— (kein Release-Vorgang erfasst)"),
        ("Vorgangsstatus", STATUS_LABEL.get(d["status"], d["status"])
         + (f" — {ENTWURF_MARKER}" if d["ist_entwurf"] else "")),
        ("Bewertungsweg (Anhang VIII)", vorgang.get("bewertungsweg") or "—"),
        ("Berichtsdatum", date.today().isoformat()),
    ]
    tbl = _tabelle(doc, ["Angabe", "Wert"], len(zeilen))
    for ri, (k, v) in enumerate(zeilen, start=1):
        tbl.cell(ri, 0).paragraphs[0].add_run(k).bold = True
        tbl.cell(ri, 1).text = str(v)
    doc.add_paragraph()
    if not vorgang:
        doc.add_paragraph(
            "Für dieses Projekt ist noch kein Release-Vorgang erfasst. Der Bericht "
            "zeigt deshalb nur den gerechneten Stand; Checkliste, Prüfung und "
            "Freigabe sind leer.")


def _rechenstand_hinweis(doc, d: dict[str, Any]) -> None:
    """Auf welchem Rechenstand der Bericht steht — und wo er vom heutigen abweicht."""
    if d["gate_basis"] is not d["gate_eingefroren"]:
        if d["status"] != "entwurf" and not d["gate_eingefroren"]:
            # Ein eingereichter Vorgang *soll* seinen Stand mitgebracht haben
            # (``einreichen(..., gate=…)``). Fehlt er, ist nicht belegbar, gegen
            # welche Zahlen geprüft wurde — das gehört in den Bericht.
            doc.add_paragraph(
                "Hinweis: Für diesen Vorgang ist kein bei der Einreichung "
                "eingefrorener Rechenstand hinterlegt. Der Bericht zeigt daher den "
                "heutigen Stand; worauf die Prüfung sich stützte, ist aus den Daten "
                "nicht belegbar.")
        return
    heute = str((d["gate_heute"] or {}).get("ampel") or "")
    eingefroren = str((d["gate_eingefroren"] or {}).get("ampel") or "")
    if not heute or heute == eingefroren:
        return
    doc.add_paragraph(
        f"Hinweis: Geprüft wurde gegen den bei der Einreichung eingefrorenen Stand "
        f"({AMPEL_LABEL.get(eingefroren, eingefroren)}). Der heutige Rechenstand "
        f"lautet {AMPEL_LABEL.get(heute, heute)} — die Abweichung ist zu bewerten, "
        f"bevor das Release weiterverwendet wird.")


def _phase1(doc, d: dict[str, Any]) -> None:
    """Phase 1 — CRA-Ready: ausdrückliche Feststellung oder nicht festgestellt."""
    _ueberschrift(doc, "Phase 1 — CRA-Ready (Produktzustand)", level=1)
    p1 = d["phase1"] or {}
    status = str(p1.get("bereit_status") or "offen")
    if status == "bereit":
        doc.add_paragraph(
            f"Festgestellt: {_person_am(p1.get('bereit_von'), p1.get('bereit_am'))}.")
        if d["phase1_ueberholt"]:
            doc.add_paragraph(
                "Die Feststellung ist durch den heutigen Produktstand überholt — die "
                "Phase-1-Ampel ist nicht mehr grün. Sie wurde nicht gelöscht (eine "
                "menschliche Aussage entfernt kein Rechenlauf), gilt aber nicht als "
                "aktueller Nachweis.")
    elif status == "entzogen":
        doc.add_paragraph(
            f"Feststellung entzogen ({_person_am(p1.get('bereit_von'), p1.get('bereit_am'))}). "
            f"Begründung: {p1.get('bereit_entzug_grund') or '— (nicht dokumentiert)'}")
    else:
        doc.add_paragraph(
            "Nicht festgestellt. Für dieses Produkt liegt keine ausdrückliche "
            "CRA-Ready-Feststellung vor.")
    gp = d["gate_produkt"] or {}
    if gp:
        doc.add_paragraph(
            f"Heutiger Produktstand: {AMPEL_LABEL.get(str(gp.get('ampel')), '—')} · "
            f"erfüllt {gp.get('erfuellt')}, begründet akzeptiert "
            f"{gp.get('begruendet_akzeptiert')}, offen {gp.get('offen')} von "
            f"{gp.get('gesamt_anforderungen')} Anforderungen.")


def _release_art(doc, vorgang: dict[str, Any]) -> None:
    """#1793: Welcher Art war dieses Release?

    Steht **vor** der Antwort auf die Kernfrage, weil die Einstufung sie
    erklärt: Ein Sicherheitsupdate, das die Zweckbestimmung nicht verändert,
    ist nach Erwägungsgrund 39 keine wesentliche Änderung. Wer den Vorgang
    später liest, sieht damit nicht nur die Antwort, sondern wovon sie
    ausgeht.
    """
    from cra import release_art as ra

    art = str((vorgang or {}).get("release_art") or "")
    if not art:
        # Bestandsvorgänge führen die Angabe nicht. Das als „unbestimmt" zu
        # drucken wäre ehrlicher als sie wegzulassen — und schweigen hieße,
        # der Leser müsste raten, ob die Frage gestellt wurde.
        doc.add_paragraph("Art des Releases: nicht erklärt.")
        return
    doc.add_paragraph(f"Art des Releases: {ra.label(art)} — {ra.ERKLAERUNG[art]}")


def _wesentliche_aenderung(doc, d: dict[str, Any], vorgang: dict[str, Any]) -> None:
    """Art. 3 Nr. 30 — die Verzweigung, die alles Weitere auslöst."""
    _ueberschrift(doc, "Wesentliche Änderung (Art. 3 Nr. 30)", level=1)
    _release_art(doc, vorgang)
    wert = str(vorgang.get("wesentliche_aenderung") or "").strip()
    grund = str(vorgang.get("wesentliche_aenderung_grund") or "").strip()
    if not wert:
        doc.add_paragraph(
            "Nicht beantwortet. Die Frage ist offen — sie wird hier ausdrücklich "
            "nicht als „nein“ gelesen, weil „nein“ eine Aussage mit Folgen ist "
            "(keine neue Konformitätsbewertung, keine neue Konformitätserklärung).")
    else:
        doc.add_paragraph("Antwort: " + ("ja" if wert == "ja" else "nein"))
        # Ohne Begründung bleibt die Antwort stehen, gilt aber nicht als Nachweis —
        # die Datenbank verlangt sie (#1744), Altbestand kann sie vermissen lassen.
        fehlt = "— (fehlt — ohne Begründung ist die Antwort nicht nachgewiesen)"
        doc.add_paragraph(f"Begründung: {grund or fehlt}")
        if wert == "ja":
            doc.add_paragraph(
                "Folge: neue Konformitätsbewertung, neue EU-Konformitätserklärung und "
                "aktualisierte technische Dokumentation.")
    # Umfang der Neubewertung — eigener Prüfpunkt des Katalogs.
    umfang = ((vorgang.get("release_checkliste") or {}).get("neubewertung_umfang") or {})
    text = str(umfang.get("begruendung") or "").strip()
    doc.add_paragraph(
        "Neu bewertete grundlegende Anforderungen: "
        + (text if text else "nicht dokumentiert"))


def _checkliste(doc, d: dict[str, Any], gate: dict[str, Any],
                vorgang: dict[str, Any], rc) -> None:
    """Die Checkliste in ihren drei Arten — getrennt, weil sie unterschiedlich
    zustande kommen (gerechnet / entschieden / bestätigt)."""
    _ueberschrift(doc, "Release-Checkliste", level=1)
    stand = vorgang.get("release_checkliste") or {}
    doc.add_paragraph(
        f"Gerechnete Punkte stammen aus dem Konformitäts-Gate ({d['gate_basis_quelle']}); "
        "Entscheidungen und Bestätigungen von Personen. Die Trennung bleibt sichtbar, "
        "weil eine Liste, die sich selbst abhaken ließe, als Nachweis wertlos wäre.")

    for art in ("gerechnet", "entscheidung", "gleichbleibend"):
        items = [i for i in rc.KATALOG if i.art == art]
        if not items:
            continue
        _ueberschrift(doc, ART_LABEL[art], level=2)
        if art == "gerechnet":
            kopf = ["Prüfpunkt", "Rechtsbezug", "Ergebnis"]
        elif art == "entscheidung":
            kopf = ["Prüfpunkt", "Rechtsbezug", "Ergebnis", "Begründung", "Person"]
        else:
            kopf = ["Angabe", "Rechtsbezug", "Stand", "Bestätigung"]
        tbl = _tabelle(doc, kopf, len(items))
        for ri, item in enumerate(items, start=1):
            tbl.cell(ri, 0).text = item.label + ("" if item.pflicht else " (optional)")
            tbl.cell(ri, 1).text = item.rechtsbezug
            if art == "gerechnet":
                erg, text = gerechnetes_ergebnis(item.quelle, gate)
                z = tbl.cell(ri, 2)
                _schattieren(z, ERG_FARBE.get(erg, "FFFFFF"))
                z.text = text
            elif art == "entscheidung":
                eintrag = stand.get(item.id) or {}
                if item.id == "wesentliche_aenderung":
                    erg, label, grund, person = wesentliche_zeile(vorgang, eintrag, item)
                else:
                    erg, label = entscheidungs_ergebnis(item, eintrag)
                    grund = str(eintrag.get("begruendung") or "").strip() or "—"
                    person = _person_am(eintrag.get("von"), eintrag.get("am"))
                z = tbl.cell(ri, 2)
                _schattieren(z, ERG_FARBE.get(erg, "FFFFFF"))
                z.text = label
                tbl.cell(ri, 3).text = grund
                tbl.cell(ri, 4).text = person
            else:
                tbl.cell(ri, 2).text = _GLEICHBLEIBEND_STAND.get(
                    item.id, lambda _g: "Stand nicht ausgewiesen")(gate)
                eintrag = stand.get(item.id) or {}
                bestaetigt = str(eintrag.get("status") or "offen") != "offen"
                z = tbl.cell(ri, 3)
                _schattieren(z, ERG_FARBE["bestaetigt" if bestaetigt else "offen"])
                z.text = (f"bestätigt — {_person_am(eintrag.get('von'), eintrag.get('am'))}"
                          if bestaetigt else "nicht bestätigt")
        doc.add_paragraph()


def _offene_punkte(doc, gate: dict[str, Any]) -> None:
    _ueberschrift(doc, "Offene Punkte", level=1)
    punkte = gate.get("offene_punkte") or []
    if not punkte:
        doc.add_paragraph("Keine offenen Punkte im zugrunde gelegten Rechenstand.")
        return
    tbl = _tabelle(doc, ["Phase", "Art", "Punkt", "Hinweis"], len(punkte))
    for ri, p in enumerate(punkte, start=1):
        z = tbl.cell(ri, 0)
        _schattieren(z, ERG_FARBE["ungeprueft" if p.get("ungeprueft") else "offen"])
        z.text = "Release" if str(p.get("phase")) == "release" else "Produkt"
        tbl.cell(ri, 1).text = str(p.get("typ") or "")
        tbl.cell(ri, 2).text = str(p.get("titel") or "")
        hinweis = str(p.get("hinweis") or "")
        if p.get("ungeprueft"):
            hinweis = f"{hinweis} [nicht geprüft — gilt als nicht erfüllt]"
        tbl.cell(ri, 3).text = hinweis
    doc.add_paragraph()


def _restrisiken_abschnitt(doc, d: dict[str, Any]) -> None:
    _ueberschrift(doc, "Akzeptierte Restrisiken", level=1)
    rr = d["restrisiken"]
    abw = rr["abweichungen"]
    if abw:
        doc.add_paragraph(
            "Abweichungen von grundlegenden Anforderungen, die über ein begründet "
            "akzeptiertes Restrisiko getragen werden (Anhang I Teil I).")
        tbl = _tabelle(doc, ["Anforderung", "Rechtsgrundlage", "Kategorie",
                             "Begründung", "akzeptiert von"], len(abw))
        for ri, a in enumerate(abw, start=1):
            tbl.cell(ri, 0).text = f"{a.get('requirement_id')} — {a.get('titel') or ''}"
            tbl.cell(ri, 1).text = str(a.get("rechtsgrundlage") or a.get("ref") or "")
            z = tbl.cell(ri, 2)
            _schattieren(z, ERG_FARBE["maengel"])
            z.text = str(a.get("kategorie") or "")
            tbl.cell(ri, 3).text = str(a.get("grund") or "—")
            tbl.cell(ri, 4).text = str(a.get("von") or "—")
        doc.add_paragraph()
    else:
        doc.add_paragraph("Keine begründet akzeptierte Abweichung von grundlegenden "
                          "Anforderungen.")

    if rr["offene_abweichungen"]:
        doc.add_paragraph(
            f"{len(rr['offene_abweichungen'])} Anforderung(en) sind weder erfüllt noch "
            "begründet akzeptiert. Sie stehen oben unter „Offene Punkte“ und schließen "
            "eine konforme Ausweisung dieses Releases aus.")

    if rr["weitere"]:
        _ueberschrift(doc, "Weitere akzeptierte Risiken der Risikobewertung "
                           f"„{rr['rb_projekt']}“", level=2)
        doc.add_paragraph(
            "Diese Risiken hängen an keiner grundlegenden Anforderung. Sie sind "
            "Teil der Risikolage, aber nicht Grundlage der Konformitätsaussage.")
        tbl = _tabelle(doc, ["Nr.", "Risiko", "Stufe", "Kategorie", "Begründung",
                             "akzeptiert von"], len(rr["weitere"]))
        for ri, r in enumerate(rr["weitere"], start=1):
            tbl.cell(ri, 0).text = str(r.get("nr") or "")
            tbl.cell(ri, 1).text = str(r.get("titel") or "")
            tbl.cell(ri, 2).text = str(r.get("level") or "")
            tbl.cell(ri, 3).text = str(r.get("kategorie") or "")
            tbl.cell(ri, 4).text = str(r.get("grund") or "—")
            tbl.cell(ri, 5).text = _person_am(r.get("von"), r.get("am"))
        doc.add_paragraph()


def _vier_augen(doc, d: dict[str, Any], vorgang: dict[str, Any]) -> None:
    """Wer eingereicht, geprüft, freigegeben und veröffentlicht hat."""
    _ueberschrift(doc, "Prüfung und Freigabe", level=1)
    zeilen = [
        ("Erstellt / eingereicht von", _person_am(vorgang.get("erstellt_von"),
                                                  vorgang.get("eingereicht_am"))),
        ("Geprüft von", _person_am(vorgang.get("geprueft_von"),
                                   vorgang.get("geprueft_am"))),
        ("Freigegeben von", _person_am(vorgang.get("freigegeben_von"),
                                       vorgang.get("freigegeben_am"))),
        ("Veröffentlicht am", _datum(vorgang.get("veroeffentlicht_am")) or "—"),
    ]
    tbl = _tabelle(doc, ["Rolle", "Person und Zeitpunkt"], len(zeilen))
    for ri, (k, v) in enumerate(zeilen, start=1):
        tbl.cell(ri, 0).paragraphs[0].add_run(k).bold = True
        tbl.cell(ri, 1).text = v
    doc.add_paragraph()
    grund = str(vorgang.get("zurueckweisung_grund") or "").strip()
    if grund:
        doc.add_paragraph(f"Zurückgewiesen mit der Begründung: {grund}")
    if d["ist_entwurf"]:
        doc.add_paragraph(
            f"{ENTWURF_MARKER}. Solange keine Freigabe durch eine zweite Person "
            "vorliegt, belegt dieser Bericht den Arbeitsstand und keine geprüfte "
            "Konformität.")


def _archiv(doc, d: dict[str, Any]) -> None:
    """Verweis auf die unveränderlichen Nachweise — mit SHA-256."""
    _ueberschrift(doc, "Archivierte Nachweise (Verweis mit SHA-256)", level=1)
    doc.add_paragraph(
        "Der Konformitätsnachweis liegt nach der Freigabe unveränderlich im Archiv. "
        "Dieser Bericht verweist darauf; die Beweislast trägt die archivierte Datei "
        "mit ihrem Hash, nicht dieser Bericht.")
    archiv = d["archiv"] or d["archiv_alle"]
    if not archiv:
        doc.add_paragraph(
            "Für dieses Projekt ist noch kein Nachweis archiviert — es gibt also "
            "keinen unveränderlichen Stand, auf den verwiesen werden kann.")
        return
    if not d["archiv"] and d["archiv_alle"]:
        doc.add_paragraph(
            f"Kein Archivstück zur Version {d['release_version'] or '—'}; aufgeführt "
            "sind die übrigen Nachweise dieses Projekts.")
    tbl = _tabelle(doc, ["Version", "Ampel", "SHA-256", "Freigegeben"], len(archiv))
    for ri, a in enumerate(archiv, start=1):
        tbl.cell(ri, 0).text = str(a.get("version") or "—")
        z = tbl.cell(ri, 1)
        ampel = str(a.get("ampel") or "")
        _schattieren(z, ERG_FARBE["erfuellt" if ampel == "gruen" else "maengel"])
        z.text = AMPEL_LABEL.get(ampel, ampel or "—")
        tbl.cell(ri, 2).text = str(a.get("sha256") or "—")
        tbl.cell(ri, 3).text = _person_am(a.get("freigegeben_von"),
                                          a.get("freigegeben_am"))
    doc.add_paragraph()


def _weitere_vorgaenge(doc, d: dict[str, Any]) -> None:
    """Kurze Reihe der übrigen Release-Vorgänge — für die Ablage, nicht als Bewertung."""
    andere = [v for v in (d["vorgaenge"] or [])
              if str(v.get("release_version") or "").strip()
              and v is not d["vorgang"]]
    if not andere:
        return
    _ueberschrift(doc, "Weitere Release-Vorgänge dieses Projekts", level=1)
    tbl = _tabelle(doc, ["Version", "Status", "wesentliche Änderung", "Freigegeben"],
                   len(andere))
    for ri, v in enumerate(andere, start=1):
        tbl.cell(ri, 0).text = str(v.get("release_version") or "")
        tbl.cell(ri, 1).text = STATUS_LABEL.get(str(v.get("freigabe_status")),
                                                str(v.get("freigabe_status") or ""))
        tbl.cell(ri, 2).text = str(v.get("wesentliche_aenderung") or "nicht beantwortet")
        tbl.cell(ri, 3).text = _person_am(v.get("freigegeben_von"),
                                          v.get("freigegeben_am"))
    doc.add_paragraph()


def build_pdf(projekt: str, *, version: str | None = None,
              db_path: Path = DB_PATH) -> bytes:
    """PDF über denselben DOCX-Weg (#1614) — kein zweiter Renderer."""
    from shared.reports.docx_pdf import docx_bytes_to_pdf

    return docx_bytes_to_pdf(build_docx(projekt, version=version, db_path=db_path))


def build(projekt: str, *, fmt: str = "docx", version: str | None = None,
          db_path: Path = DB_PATH) -> bytes:
    if (fmt or "docx").lower() == "pdf":
        return build_pdf(projekt, version=version, db_path=db_path)
    return build_docx(projekt, version=version, db_path=db_path)


# ── Kontext für die KI-Zusammenfassung ──────────────────────────────────────

def summary_kontext(projekt: str, version: str | None = None,
                    db_path: Path = DB_PATH) -> dict[str, Any]:
    """Zusatzschlüssel zum Vertrag von :mod:`shared.reports.summary`.

    ``reifegrad``/``offene`` liefert der CRA-Adapter; was hier dazukommt, ist der
    **Vorgang**. Zusatzschlüssel gehen als JSON-Block in den Prompt (#1725) —
    vorher verwarf der Prompt-Bauer alles außer den zwei Pflichtschlüsseln, und
    die KI berichtete „keine Daten vorhanden".

    Best effort: Ohne Vorgang bleibt ``release_vorgang`` leer, die
    Zusammenfassung läuft trotzdem.
    """
    try:
        d = sammeln(projekt, version, db_path=db_path)
    except Exception:  # noqa: BLE001
        return {}
    if not d.get("vorgang"):
        return {}
    gate = d["gate_basis"]
    vorgang = d["vorgang"]
    return {"release_vorgang": {
        "release_version": d["release_version"],
        "status": STATUS_LABEL.get(d["status"], d["status"]),
        "ist_entwurf": d["ist_entwurf"],
        "ampel": gate.get("ampel"),
        "rechenstand": d["gate_basis_quelle"],
        "wesentliche_aenderung": vorgang.get("wesentliche_aenderung") or "nicht beantwortet",
        "wesentliche_aenderung_grund": vorgang.get("wesentliche_aenderung_grund") or "",
        "phase1_festgestellt": (d["phase1"] or {}).get("bereit_status") == "bereit",
        "phase1_ueberholt": d["phase1_ueberholt"],
        "geprueft_von": vorgang.get("geprueft_von") or "",
        "freigegeben_von": vorgang.get("freigegeben_von") or "",
        "offene_punkte_anzahl": len(gate.get("offene_punkte") or []),
        "akzeptierte_restrisiken": len(d["restrisiken"]["abweichungen"]),
        "archivierte_nachweise": len(d["archiv"]),
        "datenlage_hinweise": d["hinweise"],
    }}
