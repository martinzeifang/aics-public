"""Kundenfassung des CRA-Berichts — zeigbar statt intern (#1758).

Alle übrigen Berichte des Moduls sind **interne** Dokumente: Sie enthalten
offene Punkte, Bewertungslücken, akzeptierte Restrisiken mit Begründung,
interne Verantwortliche und die Namen von Prüfern und Freigebern. Wer einem
Kunden „etwas zur CRA" geben wollte, hatte bisher die Wahl zwischen einem
internen Dokument und gar nichts — und griff im Zweifel zum internen.

Rechtlich muss der Kunde zweierlei bekommen: die Informationen nach Anhang II
(Art. 13 Abs. 18) und die EU-Konformitätserklärung (Art. 28). Die
**Risikoabschätzung gehört nicht dazu** — sie geht in die technische
Dokumentation und von dort an die Marktaufsicht (Art. 13 Abs. 1–3, Anhang VII
Nr. 3). Das ist kein Versehen des Gesetzgebers: Eine vollständige
Risikobewertung ist eine Landkarte der Schwachstellen.

WARUM ADDITIV UND NICHT SUBTRAKTIV
Dieses Modul sammelt **ausdrücklich erlaubte** Angaben ein, statt aus dem
internen Bericht etwas wegzulassen. Der Fehler ist hier asymmetrisch: Eine zu
knappe Kundenfassung kostet eine Rückfrage, eine zu offene kostet Vertrauen
oder liefert einem Angreifer Vorarbeit. Mit der Weglass-Variante landet beim
nächsten neuen internen Feld irgendwann etwas beim Kunden, das niemand dort
haben wollte — bei dieser Bauweise kann das nicht passieren, weil ein neues
Feld schlicht nicht abgeholt wird.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

DB_PATH = Path("data/db/cra.sqlite")

#: Begriffe, die in einer Kundenfassung nichts zu suchen haben. Der Test
#: ``tests/test_kundenfassung_1758.py`` sucht sie im erzeugten Dokument — nicht
#: weil die Liste vollständig wäre, sondern weil sie die typischen Rutschbahnen
#: abdeckt, falls jemand den Renderer später „nur kurz" erweitert.
VERBOTENE_BEGRIFFE: tuple[str, ...] = (
    "akzeptiert", "Restrisiko", "offene Punkte", "Freigeber", "geprüft von",
    "Maßnahmenplan", "Zieldatum", "verantwortlich", "Ampel", "Reifegrad",
)


def _s(wert: Any) -> str:
    return "" if wert is None else str(wert).strip()


def sammle(projekt_name: str, *, db_path: Path = DB_PATH,
           version: str | None = None) -> dict[str, Any]:
    """Die erlaubten Angaben — und nur die.

    Jeder Block ist einzeln gekapselt: Fällt eine Quelle aus, fehlt dieser
    Block, statt dass der ganze Bericht scheitert. Für eine Kundenfassung ist
    das richtig herum — ein fehlender Abschnitt ist eine Rückfrage, ein
    abgebrochener Export ist ein geplatzter Termin.
    """
    d: dict[str, Any] = {"erstellt_am": date.today().isoformat(),
                         "projekt_name": projekt_name}

    try:
        from cra.db import load_projekt
        from cra.requirements import PRODUKTKLASSEN

        p = load_projekt(db_path, projekt_name) or {}
        klasse = _s(p.get("produktklasse")) or "default"
        d["produkt"] = {
            "name": _s(p.get("produkt")) or _s(p.get("name")),
            "hersteller": _s(p.get("unternehmen")),
            "beschreibung": _s(p.get("beschreibung")),
            "klasse_label": _s((PRODUKTKLASSEN.get(klasse) or {}).get("label")),
        }
    except Exception:  # noqa: BLE001
        d["produkt"] = {}

    # ── Konformität: Verfahren, Erklärung, CE ───────────────────────────────
    try:
        from cra import doc_export

        dd = doc_export.doc_daten(db_path, projekt_name, version)
        rec = dd.get("record") or {}
        d["konformitaet"] = {
            "bewertungsweg": _s(rec.get("bewertungsweg")),
            "nb_name": _s((dd.get("doc") or {}).get("nb_name")),
            "nb_kennnummer": _s(rec.get("nb_kennnummer")),
            "doc_ausgestellt": bool(dd.get("doc_ausgestellt")) and not dd.get("fehlend"),
            "doc_ausgestellt_am": _s(dd.get("ausgestellt_am")),
            "doc_version": dd.get("doc_version"),
            "release_version": _s(dd.get("release_version")) or _s(version),
        }
        d["_record"] = rec          # nur für die CE-Sperre, nicht gedruckt
    except Exception:  # noqa: BLE001
        d["konformitaet"] = {}
        d["_record"] = {}

    # Hash der abgelegten Erklärung — der Kunde soll prüfen können, dass er
    # dieselbe Fassung hat.
    try:
        from cra import doc_export

        archiv = doc_export.list_archiv(db_path, projekt_name) or []
        d["erklaerung_archiv"] = [
            {"version": _s(a.get("doc_version") or a.get("version")),
             "am": _s(a.get("erstellt_am") or a.get("freigegeben_am"))[:10],
             "sha256": _s(a.get("sha256"))}
            for a in archiv[:5]]
    except Exception:  # noqa: BLE001
        d["erklaerung_archiv"] = []

    # ── Anhang II: was der Kunde wirklich braucht ───────────────────────────
    try:
        from cra.db import load_support_period

        sp = load_support_period(db_path, projekt_name) or {}
        d["support"] = {
            "markteintritt": _s(sp.get("markteintritt_datum"))[:10],
            "jahre": sp.get("support_jahre"),
            "eol": _s(sp.get("eol_datum"))[:10],
            "update_kanal": _s(sp.get("update_kanal")),
            # #1762: Die beiden Zusagen und die Kauf-Sichtbarkeit sind genau die
            # Art von Angabe, die hierher gehört — der Kunde soll sie schriftlich
            # haben und nicht erfragen müssen.
            "kostenlos": bool(sp.get("sicherheitsupdates_kostenlos")),
            "getrennt": bool(sp.get("kanaele_getrennt")),
            "sicherheitsupdate_kanal": _s(sp.get("sicherheitsupdate_kanal")),
            "funktionsupdate_kanal": _s(sp.get("funktionsupdate_kanal")),
            "zusagen_nachweis": _s(sp.get("update_zusagen_nachweis")),
            "kaufhinweis_url": _s(sp.get("support_end_kaufhinweis_url")),
        }
    except Exception:  # noqa: BLE001
        d["support"] = {}

    try:
        from cra.db import load_psirt

        from cra.psirt import als_text

        ps = load_psirt(db_path, projekt_name) or {}
        d["meldestelle"] = {
            "policy_url": _s(ps.get("disclosure_policy_url")),
            "security_md": _s(ps.get("security_md_url")),
            # #1796: `_s` einer Liste druckte `['security@…']` wörtlich in das
            # Dokument, das das Haus verlässt. Die Spalte trägt beides, also
            # entscheidet eine Stelle, wie Kontakte aussehen.
            "kontakt": als_text(ps),
        }
    except Exception:  # noqa: BLE001
        d["meldestelle"] = {}

    # ── Veröffentlichte Advisories ──────────────────────────────────────────
    # Nur **behobene** Schwachstellen mit veröffentlichter Advisory-Adresse.
    # Offene Befunde erscheinen hier unter keinen Umständen: Sie sind die
    # Angriffsfläche von heute, nicht die Historie von gestern.
    try:
        from cra.db import list_vuln

        d["advisories"] = [
            {"id": _s(v.get("cve_id")), "behoben_in": _s(v.get("fixed_in_version")),
             "url": _s(v.get("advisory_url"))}
            for v in (list_vuln(db_path, projekt_name) or [])
            if _s(v.get("advisory_url")) and _s(v.get("status")) in ("fixed", "resolved")
        ][:50]
    except Exception:  # noqa: BLE001
        d["advisories"] = []

    # ── SBOM: Stand, nicht Inhalt ───────────────────────────────────────────
    try:
        from cra.db import list_sbom

        sbom = list_sbom(db_path, projekt_name) or []
        neueste = sbom[0] if sbom else {}
        d["sbom"] = {"vorhanden": bool(sbom),
                     "format": _s(neueste.get("sbom_format")),
                     "stand": _s(neueste.get("sbom_datum"))[:10],
                     "version": _s(neueste.get("release_version"))}
    except Exception:  # noqa: BLE001
        d["sbom"] = {"vorhanden": False}

    # ── Sicherheitstests: Art und Stand, keine Befunde (#1824) ──────────────
    # Anhang I Teil II Nr. 3 ist eine Pflicht gegenüber dem Markt, und was der
    # Hersteller dazu tut — oder bewusst nicht tut —, ist eine Aussage über das
    # Produkt und keine interne Notiz. Ein Dokument, das den Verzicht
    # verschweigt, ist gegenüber dem Empfänger irreführend.
    #
    # Gedruckt wird **Art und Zeitpunkt**, nie ein Befund: Welche Lücke ein
    # Test gefunden hat, ist die Angriffsfläche von heute und bleibt — wie die
    # offenen Schwachstellen oben — außen vor.
    try:
        from cra import sicherheitstests as st

        b = st.befund(db_path, projekt_name, _s(version))
        v = b.get("verzicht") or {}
        d["tests"] = {
            "art": _s(b.get("art")),
            "art_label": _s(b.get("art_label")),
            "juengster": _s(b.get("juengster")),
            "anzahl": int(b.get("nachweis_count") or 0),
            "verzicht": bool(v.get("erklaert")),
            "verzicht_begruendung": _s(v.get("begruendung")),
            "verzicht_review": _s(v.get("review")),
        }
    except Exception:  # noqa: BLE001
        d["tests"] = {}

    # ── Risikomanagement: Methode und Stand, keine Risiken ──────────────────
    try:
        from cra.db import load_projekt
        from risikobewertung.db import load_projekt as rb_projekt

        meta = (load_projekt(db_path, projekt_name) or {}).get("meta") or {}
        rb_name = _s(meta.get("linked_risk_projekt"))
        rb = rb_projekt(Path("data/db/risikobewertung.sqlite"), rb_name) if rb_name else {}
        d["risikomanagement"] = {
            "gefuehrt": bool(rb_name),
            "methode": _s((rb or {}).get("framework")),
            "stand": _s((rb or {}).get("updated_at"))[:10],
        }
    except Exception:  # noqa: BLE001
        d["risikomanagement"] = {"gefuehrt": False}

    return d


def build_docx(projekt_name: str, *, db_path: Path = DB_PATH,
               version: str | None = None) -> bytes:
    """Kundenfassung als DOCX."""
    import io

    from docx import Document
    from docx.shared import Pt

    from cra import ce_mark
    from shared.reports.docx_pdf import apply_base_font, set_a4

    d = sammle(projekt_name, db_path=db_path, version=version)
    doc = Document()
    set_a4(doc)

    prod = d.get("produkt") or {}
    titel = doc.add_heading("Cybersicherheit — Information für Kunden", level=0)
    titel.runs[0].font.size = Pt(18)
    doc.add_paragraph(
        f"{prod.get('name') or projekt_name} · {prod.get('hersteller') or ''}"
        f" · Stand {d['erstellt_am']}")
    doc.add_paragraph(
        "Dieses Dokument fasst zusammen, was für Sie als Anwender dieses "
        "Produkts mit digitalen Elementen von Bedeutung ist. Die vollständige "
        "technische Dokumentation einschließlich der Risikoabschätzung wird "
        "der Marktaufsichtsbehörde auf Verlangen vorgelegt "
        "(Verordnung (EU) 2024/2847, Art. 13 Abs. 1–3 i.V.m. Anhang VII).")

    k = d.get("konformitaet") or {}
    doc.add_heading("Konformität", level=1)
    if k.get("doc_ausgestellt"):
        doc.add_paragraph(
            f"Für dieses Produkt liegt eine EU-Konformitätserklärung nach "
            f"Art. 28 i.V.m. Anhang V vor (Fassung {k.get('doc_version')}, "
            f"ausgestellt am {k.get('doc_ausgestellt_am')}).")
    else:
        doc.add_paragraph("Die EU-Konformitätserklärung ist noch nicht ausgestellt.")
    if k.get("bewertungsweg"):
        satz = f"Angewandtes Konformitätsbewertungsverfahren: Modul {k['bewertungsweg']} (Anhang VIII)."
        if k.get("nb_kennnummer"):
            satz += (f" Beteiligte notifizierte Stelle: {k.get('nb_name') or ''} "
                     f"(Kennnummer {k['nb_kennnummer']}).")
        doc.add_paragraph(satz)
    for a in d.get("erklaerung_archiv") or []:
        if a.get("sha256"):
            doc.add_paragraph(
                f"Hinterlegte Fassung {a['version']} vom {a['am']} — "
                f"SHA-256 {a['sha256'][:16]}…")

    doc.add_heading("CE-Kennzeichnung", level=1)
    ce_mark.einfuegen(doc, d.get("_record") or {}, bool(k.get("doc_ausgestellt")))

    s = d.get("support") or {}
    doc.add_heading("Unterstützungszeitraum und Updates", level=1)
    if s.get("eol"):
        doc.add_paragraph(
            f"Sicherheitsupdates werden bis zum {s['eol']} bereitgestellt"
            + (f" (Markteintritt {s['markteintritt']}, {s['jahre']} Jahre)."
               if s.get("markteintritt") else "."))
    else:
        doc.add_paragraph("Der Unterstützungszeitraum ist noch nicht festgelegt.")
    # #1762 — zwei Zusagen, die der Kunde schwarz auf weiß haben soll. Fehlt
    # eine, wird das gesagt und nicht verschwiegen: Eine Kundenfassung, die eine
    # nicht dokumentierte Zusage behauptet, ist schlimmer als eine, die ihre
    # Lücke zeigt. Die Angabe der Verordnungsstelle fehlt hier absichtlich — das
    # Issue markiert sie als unsicher (Art. 13 Abs. 12-14, nicht nachgeprüft),
    # und ein falsches Zitat in einem Kundendokument ist schlechter als keines.
    if s.get("kostenlos"):
        doc.add_paragraph(
            "Sicherheitsupdates werden während des Unterstützungszeitraums "
            "kostenlos bereitgestellt.")
    else:
        doc.add_paragraph(
            "Zur Entgeltfreiheit der Sicherheitsupdates liegt keine dokumentierte "
            "Zusage vor.")
    if s.get("getrennt"):
        doc.add_paragraph(
            "Sicherheitsupdates werden getrennt von Funktionsupdates "
            "bereitgestellt: Sie können Sicherheitsupdates beziehen, ohne "
            "Funktionsänderungen übernehmen zu müssen."
            + (f" Kanal für Sicherheitsupdates: {s['sicherheitsupdate_kanal']}."
               if s.get("sicherheitsupdate_kanal") else ""))
    elif s.get("sicherheitsupdate_kanal"):
        doc.add_paragraph(
            f"Kanal für Sicherheitsupdates: {s['sicherheitsupdate_kanal']}")
    if s.get("funktionsupdate_kanal"):
        doc.add_paragraph(f"Kanal für Funktionsupdates: {s['funktionsupdate_kanal']}")
    if s.get("update_kanal"):
        doc.add_paragraph(f"Bezug der Updates: {s['update_kanal']}")
    if s.get("zusagen_nachweis"):
        doc.add_paragraph(f"Diese Zusagen sind veröffentlicht unter: {s['zusagen_nachweis']}")
    if s.get("kaufhinweis_url"):
        doc.add_paragraph(
            "Das Ende des Unterstützungszeitraums ist vor dem Kauf einsehbar "
            f"unter: {s['kaufhinweis_url']}")

    m = d.get("meldestelle") or {}
    doc.add_heading("Schwachstellen melden", level=1)
    if m.get("policy_url") or m.get("security_md") or m.get("kontakt"):
        if m.get("kontakt"):
            doc.add_paragraph(f"Kontakt: {m['kontakt']}")
        for url in (m.get("policy_url"), m.get("security_md")):
            if url:
                doc.add_paragraph(url)
    else:
        doc.add_paragraph("Eine Meldestelle ist noch nicht veröffentlicht.")

    r = d.get("risikomanagement") or {}
    doc.add_heading("Risikomanagement", level=1)
    if r.get("gefuehrt"):
        satz = "Für dieses Produkt wird eine Cybersicherheits-Risikoabschätzung geführt"
        if r.get("methode"):
            satz += f" (Methode: {r['methode']})"
        if r.get("stand"):
            satz += f", zuletzt fortgeschrieben am {r['stand']}"
        doc.add_paragraph(satz + ".")
    else:
        doc.add_paragraph("Die Risikoabschätzung wird außerhalb dieses Systems geführt.")
    doc.add_paragraph(
        "Die Abschätzung selbst ist Teil der technischen Dokumentation und "
        "wird nicht veröffentlicht: Sie beschreibt Angriffswege und "
        "Schutzmaßnahmen im Einzelnen.")

    # #1824 — Anhang I Teil II Nr. 3. Zwei Dinge stehen hier bewusst **nicht**:
    # kein Befund (die Lücke von heute gehört nicht in ein Dokument, das das
    # Haus verlässt — dieselbe Regel wie bei den Advisories oben) und nicht die
    # interne Begründung des Verzichts. Dass verzichtet wird, ist eine Aussage
    # über das Produkt; *warum*, ist die Abwägung des Hauses und steht im
    # Konformitätsnachweis (`cra/konformitaet_report.py`), der auf Verlangen
    # vorgelegt wird. Verschwiegen wird der Verzicht aber nicht: Ein Dokument,
    # das Tests offenlässt, obwohl der Hersteller sich dagegen entschieden hat,
    # führt den Empfänger in die Irre.
    t = d.get("tests") or {}
    doc.add_heading("Sicherheitstests", level=1)
    if t.get("verzicht"):
        satz = ("Der Hersteller führt für dieses Produkt derzeit keine "
                "gesonderten Sicherheitstests durch. Diese Entscheidung ist "
                "dokumentiert")
        if t.get("verzicht_review"):
            satz += f" und zum {t['verzicht_review']} erneut zu treffen"
        doc.add_paragraph(satz + ".")
    elif t.get("anzahl"):
        satz = f"Die Sicherheit dieses Produkts wird getestet: {t['art_label']}"
        if t.get("juengster"):
            satz += f"; jüngster Nachweis vom {t['juengster']}"
        doc.add_paragraph(
            satz + ". Die Testberichte sind Teil der technischen "
            "Dokumentation (Anhang VII Nr. 6) und werden nicht veröffentlicht.")
    else:
        doc.add_paragraph(
            "Zu Sicherheitstests liegt in dieser Akte derzeit kein Nachweis vor.")

    doc.add_heading("Veröffentlichte Sicherheitshinweise", level=1)
    adv = d.get("advisories") or []
    if adv:
        doc.add_paragraph(
            "Zu folgenden behobenen Schwachstellen wurden Hinweise "
            "veröffentlicht (Anhang I Teil II Nr. 4):")
        for a in adv:
            doc.add_paragraph(
                f"{a['id']} — behoben in {a['behoben_in'] or '—'} — {a['url']}",
                style="List Bullet")
    else:
        doc.add_paragraph("Es wurden bislang keine Sicherheitshinweise veröffentlicht.")

    sb = d.get("sbom") or {}
    doc.add_heading("Software-Stückliste (SBOM)", level=1)
    doc.add_paragraph(
        (f"Eine SBOM im Format {sb.get('format') or '—'} liegt vor "
         f"(Stand {sb.get('stand') or '—'}"
         + (f", Version {sb['version']}" if sb.get("version") else "")
         + "). Sie wird auf Anfrage bereitgestellt.")
        if sb.get("vorhanden") else
        "Eine Software-Stückliste liegt derzeit nicht vor.")

    apply_base_font(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pdf(projekt_name: str, *, db_path: Path = DB_PATH,
              version: str | None = None) -> bytes:
    from shared.reports.docx_pdf import docx_bytes_to_pdf

    return docx_bytes_to_pdf(build_docx(projekt_name, db_path=db_path, version=version))


def build(projekt_name: str, *, fmt: str = "docx", db_path: Path = DB_PATH,
          version: str | None = None) -> bytes:
    if (fmt or "docx").lower() == "pdf":
        return build_pdf(projekt_name, db_path=db_path, version=version)
    return build_docx(projekt_name, db_path=db_path, version=version)


def summary_context(projekt_name: str, *, db_path: Path = DB_PATH) -> dict[str, Any]:
    """Kontext für die KI-Zusammenfassung (#1725) — ebenfalls ohne Internas."""
    d = sammle(projekt_name, db_path=db_path)
    k = d.get("konformitaet") or {}
    return {
        "reifegrad": None,
        "offene": [],
        "kundenfassung": {
            "produkt": (d.get("produkt") or {}).get("name"),
            "erklaerung_ausgestellt": bool(k.get("doc_ausgestellt")),
            "support_bis": (d.get("support") or {}).get("eol"),
            "updates_kostenlos": bool((d.get("support") or {}).get("kostenlos")),
            "updates_getrennt": bool((d.get("support") or {}).get("getrennt")),
            "advisories": len(d.get("advisories") or []),
            "sbom": bool((d.get("sbom") or {}).get("vorhanden")),
            # #1824: Der Verzicht gehört auch in eine Zusammenfassung — sonst
            # formuliert die KI um ihn herum.
            "sicherheitstests": (d.get("tests") or {}).get("art_label"),
            "ohne_sicherheitstests": bool((d.get("tests") or {}).get("verzicht")),
        },
    }
