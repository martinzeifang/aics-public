"""Bausteine für Word-Vorlagen im Cyberwoks-Design (#1790).

WOHER DAS DESIGN KOMMT
Nicht nachgebaut, sondern **geerbt**: Jede erzeugte Vorlage entsteht aus einer
der vorhandenen Cyberwoks-Dateien als Rohling. Damit bleiben Formatvorlagen,
Theme, Nummerierung und der Aufbau der Kopf- und Fußzeile erhalten — also genau
das, was ein Nachbau aus Einzelangaben nie vollständig trifft. Ersetzt wird nur
der Textkörper.

Die Farb- und Schriftwerte unten sind aus dem Rohling ausgelesen und stehen
hier, weil neue Absätze sie brauchen; sie sind eine Kopie, kein zweiter
Ursprung. Ändert sich das Briefpapier, wird aus dem neuen Rohling erzeugt.
"""

from __future__ import annotations

import contextlib
import copy
from pathlib import Path
from typing import Any, Callable, Iterable

from docx import Document
from docx.enum.section import WD_HEADER_FOOTER, WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# Aus dem Rohling gelesen.
BLAU = "0F2A44"        # Überschriften, Titel
TUERKIS = "00A3C4"     # Akzent über dem Titel
GRAU = "5B6B7A"        # Sekundärtext
HELL = "EEF3F7"        # Zellhintergrund für Label- und Kennzahlfelder
SCHRIFT = "Calibri"


def _rgb(hex_wert: str) -> RGBColor:
    return RGBColor.from_string(hex_wert)


# ── Rohling vorbereiten ─────────────────────────────────────────────────────

def oeffne_rohling(quelle: Path) -> Document:
    """Den Rohling laden und seinen Textkörper leeren.

    Der Abschnittsabschluss (``sectPr``) bleibt stehen: An ihm hängen
    Seitengröße, Ränder und die Verknüpfung zu Kopf- und Fußzeile. Wird er
    mitgelöscht, verliert das Dokument sein Briefpapier.
    """
    doc = Document(str(quelle))
    body = doc.element.body
    for kind in list(body):
        if kind.tag == qn("w:sectPr"):
            continue
        body.remove(kind)
    doc._cyberwoks_layout = _Layout(doc)
    return doc


# ── Hoch- und Querformat im selben Dokument ─────────────────────────────────
#
# Breite Auflistungen (Schwachstellen, Lücken, Prüffragen, Risiken) haben im
# Hochformat keinen Platz: Bei sechs Spalten mit Freitext bleiben je Spalte
# knapp drei Zentimeter, und die Tabelle wird zur Buchstabensäule. Word löst das
# über **Abschnitte** — die Ausrichtung gehört nicht dem Dokument, sondern dem
# Abschnitt. Eine Vorlage kann also hochformatig sein und für eine Tabelle ins
# Querformat wechseln.
#
# Zwei Dinge wechseln dabei mit und werden leicht übersehen:
#  * „Abweichende erste Seite" (``w:titlePg``) gilt **je Abschnitt**. Bleibt die
#    Einstellung in den neuen Abschnitten stehen, bekommt deren erste Seite die
#    Titelseiten-Kopf-/Fußzeile des Rohlings (leerer Kopf, Anschrift im Fuß).
#  * Der rechte Tabulator in Kopf- und Fußzeile steht **absolut** (im Rohling
#    9638 Twips = 17 cm Textbreite). Im Querformat sind es 25,7 cm; mit dem
#    geerbten Tabulator stünde „Seite x von y" mitten auf der Seite.

#: EMU je Twip — python-docx rechnet in EMU, Tabulatoren stehen in Twips.
_TWIP = 635


class _Layout:
    """Welche Ausrichtung gerade gilt und welche gelten soll.

    Der Wechsel wird **aufgeschoben** und erst beim nächsten Inhalt vollzogen.
    Das hat zwei Gründe: Zwei aufeinanderfolgende Querformat-Blöcke werden so zu
    *einem* Abschnitt (kein Seitenumbruch zwischen zwei breiten Tabellen), und
    ein Block, auf den nichts mehr folgt, hinterlässt keinen leeren Abschnitt am
    Dokumentende — also keine leere letzte Seite.
    """

    def __init__(self, doc: Document):
        sectPr = doc.sections[0]._sectPr
        kopf = sectPr.get_headerReference(WD_HEADER_FOOTER.PRIMARY)
        fuss = sectPr.get_footerReference(WD_HEADER_FOOTER.PRIMARY)
        self.hoch = (kopf.rId if kopf is not None else None,
                     fuss.rId if fuss is not None else None)
        self.quer: tuple[str | None, str | None] | None = None
        self.ist_quer = False
        self.soll_quer = False


@contextlib.contextmanager
def querformat(doc: Document):
    """Alles in diesem Block steht im Querformat, alles davor und danach hoch.

    Der Block gehört um **Überschrift und Tabelle** gelegt, nicht nur um die
    Tabelle: Sonst bleibt die Überschrift allein am Fuß der Hochformat-Seite
    zurück und die Tabelle beginnt beziehungslos auf der nächsten.

    Eine Schlussnotiz braucht den Block nicht: ``hinweiszeile`` mit
    ``layout_wechsel=False`` bleibt im laufenden Abschnitt, damit das Dokument
    für eine Zeile Kleindruck nicht noch eine eigene Seite anfängt.
    """
    zustand = _layout(doc)
    vorher = zustand.soll_quer
    zustand.soll_quer = True
    try:
        yield
    finally:
        zustand.soll_quer = vorher


def _layout(doc: Document) -> _Layout:
    zustand = getattr(doc, "_cyberwoks_layout", None)
    if zustand is None:
        zustand = _Layout(doc)
        doc._cyberwoks_layout = zustand
    return zustand


def _inhalt_beginnt(doc: Document) -> None:
    """Vor jedem sichtbaren Bestandteil: aufgeschobenen Wechsel nachholen."""
    zustand = _layout(doc)
    if zustand.soll_quer == zustand.ist_quer:
        return
    _abschnitt_wechseln(doc, quer=zustand.soll_quer)
    zustand.ist_quer = zustand.soll_quer


def _abschnitt_wechseln(doc: Document, *, quer: bool) -> None:
    neu = doc.add_section(WD_SECTION.NEW_PAGE)
    neu.different_first_page_header_footer = False
    kurz, lang = sorted((neu.page_width, neu.page_height))
    neu.orientation = WD_ORIENT.LANDSCAPE if quer else WD_ORIENT.PORTRAIT
    neu.page_width, neu.page_height = (lang, kurz) if quer else (kurz, lang)
    _hdrftr_verweisen(doc, neu, quer=quer)


def _hdrftr_verweisen(doc: Document, abschnitt, *, quer: bool) -> None:
    """Kopf- und Fußzeile des Abschnitts ausdrücklich benennen.

    ``add_section`` entfernt im neuen Abschnitt **alle** Kopf- und
    Fußzeilenverweise; er „erbt" damit vom vorhergehenden. Für den
    Querformat-Abschnitt ist das falsch (er erbt den Hochformat-Tabulator), und
    für den Hochformat-Abschnitt **nach** einem Querformat-Abschnitt erst recht:
    Er würde die breite Fassung erben. Jeder Abschnitt sagt deshalb selbst,
    welche Fassung gilt.
    """
    zustand = _layout(doc)
    if zustand.hoch == (None, None):
        return  # Rohling ohne Kopf- und Fußzeile: nichts nachzuziehen
    kopf_id, fuss_id = _hdrftr_quer(doc, abschnitt) if quer else zustand.hoch
    sectPr = abschnitt._sectPr
    for holen, anlegen, rid in (
            (sectPr.get_headerReference, sectPr.add_headerReference, kopf_id),
            (sectPr.get_footerReference, sectPr.add_footerReference, fuss_id)):
        if not rid:
            continue
        verweis = holen(WD_HEADER_FOOTER.PRIMARY)
        if verweis is None:
            anlegen(WD_HEADER_FOOTER.PRIMARY, rid)
        else:
            verweis.rId = rid


def _hdrftr_quer(doc: Document, abschnitt) -> tuple[str | None, str | None]:
    """Kopf- und Fußzeile fürs Querformat anlegen — einmal je Dokument.

    Der Inhalt wird **kopiert**, nicht neu gebaut: In der Fußzeile steckt das
    Feld „Seite x von y", und ein Neuaufbau verliert Feldcodes (Lehre aus
    #1007). Verschoben wird nur der rechte Tabulator auf die neue Textbreite.
    """
    zustand = _layout(doc)
    if zustand.quer is not None:
        return zustand.quer

    breite = int(abschnitt.page_width - abschnitt.left_margin
                 - abschnitt.right_margin) // _TWIP
    ids: list[str | None] = []
    for anlegen, quelle in ((doc.part.add_header_part, abschnitt.header),
                            (doc.part.add_footer_part, abschnitt.footer)):
        # Erst kopieren, dann den Teil anlegen: ``abschnitt.header`` löst die
        # geerbte Hochformat-Fassung auf, solange der Verweis noch auf sie zeigt.
        inhalt = copy.deepcopy(quelle._element)
        teil, rid = anlegen()
        for kind in list(teil.element):
            teil.element.remove(kind)
        for kind in list(inhalt):
            teil.element.append(kind)
        _tabulatoren_setzen(teil.element, breite)
        ids.append(rid)
    zustand.quer = (ids[0], ids[1])
    return zustand.quer


def _tabulatoren_setzen(element, pos_twips: int) -> None:
    """Rechte Tabulatoren auf die Textbreite setzen.

    Nur ``w:tab`` mit ``w:val="right"`` — dasselbe Element steht in einem
    Textlauf auch für das Tabulator*zeichen* und hat dort kein ``w:val``.
    """
    for tab in element.iter(qn("w:tab")):
        if tab.get(qn("w:val")) == "right":
            tab.set(qn("w:pos"), str(pos_twips))


def stil_finden(doc: Document, name: str):
    """Eine Formatvorlage am Objekt suchen, nicht über ihren Namen.

    ``add_paragraph(style="Heading 1")`` schlägt fehl, obwohl die Vorlage
    existiert: Word legt die eingebauten Namen klein ab (``heading 1``),
    python-docx zeigt beim Lesen die normalisierte Form und sucht beim
    Zuweisen die gespeicherte. Über das Style-Objekt gibt es diese Lücke nicht.
    """
    for s in doc.styles:
        if s.name == name:
            return s
    return None


def _absatz(doc: Document, text: str = "", *, stil: str | None = None,
            groesse: float | None = None, farbe: str | None = None,
            fett: bool = False, kursiv: bool = False,
            ausrichtung=None, abstand_vor: float = 0, abstand_nach: float = 4,
            layout_wechsel: bool = True):
    if layout_wechsel:
        _inhalt_beginnt(doc)
    p = doc.add_paragraph()
    if stil:
        gefunden = stil_finden(doc, stil)
        if gefunden is not None:
            p.style = gefunden
    if ausrichtung is not None:
        p.alignment = ausrichtung
    p.paragraph_format.space_before = Pt(abstand_vor)
    p.paragraph_format.space_after = Pt(abstand_nach)
    if text:
        r = p.add_run(text)
        r.font.name = SCHRIFT
        if groesse:
            r.font.size = Pt(groesse)
        if farbe:
            r.font.color.rgb = _rgb(farbe)
        r.font.bold = fett
        r.font.italic = kursiv
    return p


def _zelle_faerben(zelle, hex_wert: str) -> None:
    shd = zelle._tc.get_or_add_tcPr().makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_wert)
    zelle._tc.get_or_add_tcPr().append(shd)


def _zellen_text(zelle, text: str, *, fett: bool = False, farbe: str | None = None,
                 groesse: float = 9.5) -> None:
    zelle.text = ""
    p = zelle.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.name = SCHRIFT
    r.font.size = Pt(groesse)
    r.font.bold = fett
    if farbe:
        r.font.color.rgb = _rgb(farbe)


# ── Sichtbare Bestandteile ──────────────────────────────────────────────────

def titelblock(doc: Document, *, dokumentart: str, titel_ausdruck: str,
               zeile_unten: str, untertitel: str) -> None:
    """Der Kopf der ersten Seite, wie ihn der Rohling zeigt."""
    _absatz(doc, "CYBERWOKS", groesse=20, farbe=BLAU, fett=True, abstand_nach=0)
    _absatz(doc, "Finkenweg 42 · 89150 Laichingen", groesse=8.5, farbe=GRAU,
            abstand_nach=18)
    _absatz(doc, dokumentart.upper(), groesse=11, farbe=TUERKIS, fett=True,
            abstand_nach=2)
    _absatz(doc, titel_ausdruck, groesse=30, farbe=BLAU, fett=True, abstand_nach=2)
    if zeile_unten:
        _absatz(doc, zeile_unten, groesse=15, farbe=BLAU, abstand_nach=2)
    if untertitel:
        _absatz(doc, untertitel, groesse=11, farbe=GRAU, abstand_nach=16)


def infotabelle(doc: Document, zeilen: Iterable[tuple[str, str]]) -> None:
    """Zweispaltige Angaben zum Dokument (Label hinterlegt, Wert daneben)."""
    zeilen = list(zeilen)
    _inhalt_beginnt(doc)
    t = doc.add_table(rows=len(zeilen), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = True
    for i, (label, wert) in enumerate(zeilen):
        links, rechts = t.rows[i].cells
        _zelle_faerben(links, HELL)
        _zellen_text(links, label, fett=True, farbe=BLAU)
        _zellen_text(rechts, wert)
    _absatz(doc, abstand_nach=10)


def kennzahlen(doc: Document, felder: Iterable[tuple[str, str]]) -> None:
    """Eine Reihe hinterlegter Kacheln: großer Wert, kleine Bezeichnung."""
    felder = list(felder)
    if not felder:
        return
    _inhalt_beginnt(doc)
    t = doc.add_table(rows=1, cols=len(felder))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for zelle, (wert, bezeichnung) in zip(t.rows[0].cells, felder):
        _zelle_faerben(zelle, HELL)
        zelle.text = ""
        p1 = zelle.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(wert)
        r1.font.name = SCHRIFT
        r1.font.size = Pt(18)
        r1.font.bold = True
        r1.font.color.rgb = _rgb(BLAU)
        p2 = zelle.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(0)
        r2 = p2.add_run(bezeichnung)
        r2.font.name = SCHRIFT
        r2.font.size = Pt(8)
        r2.font.color.rgb = _rgb(GRAU)
    _absatz(doc, abstand_nach=12)


def kapitel(doc: Document, nummer: int, titel: str) -> None:
    _inhalt_beginnt(doc)
    p = doc.add_paragraph()
    gefunden = stil_finden(doc, "Heading 1")
    if gefunden is not None:
        p.style = gefunden
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f"{nummer}  {titel}")
    r.font.name = SCHRIFT
    r.font.color.rgb = _rgb(BLAU)


def unterkapitel(doc: Document, titel: str) -> None:
    _inhalt_beginnt(doc)
    p = doc.add_paragraph()
    gefunden = stil_finden(doc, "Heading 2")
    if gefunden is not None:
        p.style = gefunden
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(titel)
    r.font.name = SCHRIFT
    r.font.color.rgb = _rgb(BLAU)


def zwischentitel(doc: Document, text: str) -> None:
    """Die kleine türkise Auszeichnung über einem Block (wie im Rohling)."""
    _absatz(doc, text.upper(), groesse=9, farbe=TUERKIS, fett=True,
            abstand_vor=8, abstand_nach=3)


def fliesstext(doc: Document, text: str) -> None:
    _absatz(doc, text, groesse=10.5, farbe="000000", abstand_nach=6)


def hinweiszeile(doc: Document, text: str, *, layout_wechsel: bool = True) -> None:
    """``layout_wechsel=False`` für Schlussnotizen: Eine Zeile Kleindruck ist
    keinen Abschnittswechsel und damit keine eigene Seite wert."""
    _absatz(doc, text, groesse=8.5, farbe=GRAU, kursiv=True, abstand_nach=8,
            layout_wechsel=layout_wechsel)


def datentabelle(doc: Document, *, schleife: str, spalten: list[tuple[str, str]],
                 leer_hinweis: str = "", bedingung: str = "") -> None:
    """Tabelle über eine Liste.

    Aufbau mit **vier** Zeilen: Kopf, eine Zeile nur mit ``{%tr for … %}``, die
    Datenzeile, eine Zeile nur mit ``{%tr endfor %}``. Beide Steuerzeilen
    verschwinden beim Rendern, die Datenzeile wird wiederholt.

    Beides in eine Zeile zu schreiben — der naheliegende Weg — ergibt eine
    Vorlage, die Jinja mit „Encountered unknown tag 'endfor'" abweist: Die
    Steuerzeile wird als Ganzes entfernt, und damit verschwindet der
    Schleifenanfang mitsamt seinem Ende. Die mitgelieferte Cyberwoks-Vorlage
    macht es ebenso in getrennten Zeilen.
    """
    _inhalt_beginnt(doc)
    if bedingung:
        _absatz(doc, "{%%p if %s %%}" % bedingung, abstand_nach=0)

    t = doc.add_table(rows=4, cols=len(spalten))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for zelle, (kopf, _) in zip(t.rows[0].cells, spalten):
        _zelle_faerben(zelle, BLAU)
        _zellen_text(zelle, kopf, fett=True, farbe="FFFFFF", groesse=9)

    _zellen_text(t.rows[1].cells[0], "{%%tr %s %%}" % schleife)
    for zelle, (_, ausdruck) in zip(t.rows[2].cells, spalten):
        _zellen_text(zelle, ausdruck)
    _zellen_text(t.rows[3].cells[0], "{%tr endfor %}")

    _absatz(doc, abstand_nach=8)
    if bedingung:
        _absatz(doc, "{%p else %}", abstand_nach=0)
        if leer_hinweis:
            hinweiszeile(doc, leer_hinweis)
        _absatz(doc, "{%p endif %}", abstand_nach=8)


def aufzaehlung(doc: Document, *, schleife: str, ausdruck: str,
                bedingung: str = "", leer_hinweis: str = "") -> None:
    if bedingung:
        _absatz(doc, "{%%p if %s %%}" % bedingung, abstand_nach=0)
    _absatz(doc, "{%%p %s %%}" % schleife, abstand_nach=0)
    p = doc.add_paragraph()
    gefunden = stil_finden(doc, "List Paragraph")
    if gefunden is not None:
        p.style = gefunden
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(ausdruck)
    r.font.name = SCHRIFT
    r.font.size = Pt(10.5)
    _absatz(doc, "{%p endfor %}", abstand_nach=6)
    if bedingung:
        _absatz(doc, "{%p else %}", abstand_nach=0)
        if leer_hinweis:
            hinweiszeile(doc, leer_hinweis)
        _absatz(doc, "{%p endif %}", abstand_nach=6)


def seitenumbruch(doc: Document) -> None:
    _inhalt_beginnt(doc)
    doc.add_page_break()


# ── Kopf- und Fußzeile ──────────────────────────────────────────────────────

def kopf_und_fuss(doc: Document, *, dokumentart: str, titel_ausdruck: str,
                  fuss_mitte: str) -> None:
    """Die Texte in Kopf- und Fußzeile auf dieses Dokument umschreiben.

    Der Rohling bringt beide samt Tabulatoren und Seitenzahlfeld mit. Ersetzt
    wird nur der Text der bestehenden Textläufe — das Feld „Seite x von y" ist
    ein Feldcode und würde beim Neuaufbau verloren gehen (Lehre aus #1007).
    """
    for abschnitt in doc.sections:
        for absatz in abschnitt.header.paragraphs:
            _lauf_text_ersetzen(absatz, {
                "CYBERWOKS": "CYBERWOKS",
                "__RECHTS__": f"{dokumentart} · {titel_ausdruck}",
            }, rechts_ab_tab=f"{dokumentart} · {titel_ausdruck}")
        for absatz in abschnitt.footer.paragraphs:
            _lauf_text_ersetzen(absatz, {}, links_bis_tab=fuss_mitte)


def _lauf_text_ersetzen(absatz, _ersatz: dict[str, str], *,
                        rechts_ab_tab: str | None = None,
                        links_bis_tab: str | None = None) -> None:
    """Text vor bzw. nach dem Tabulator ersetzen, Feldcodes unangetastet.

    Der Tabulator steht im Rohling als **Zeichen** im Textlauf, nicht als
    ``w:tab``-Element. Die erste Fassung suchte nur nach dem Element und fand
    deshalb keinen — mit zwei Folgen, die erst beim Nachrechnen der fertigen
    Datei auffielen:

    * Die rechte Kopfzeilenhälfte wurde **nie** ersetzt. In allen 30 Vorlagen
      stand dort die Dokumentart des Rohlings („CRA-Readiness-Bericht").
    * „Seite x von y" verschwand. Ohne Tabulator galt der ganze Absatz als
      linke Hälfte, und ``run.text = ""`` entfernt **alle** Kinder eines
      Textlaufs — auch ``fldChar`` und ``instrText``. Genau der Verlust, vor
      dem der Docstring von :func:`kopf_und_fuss` warnt (#1007).
    """
    laeufe = absatz.runs
    if not laeufe:
        return
    tab_index = next((i for i, r in enumerate(laeufe) if "\t" in r.text), None)

    if links_bis_tab is not None:
        if tab_index == 0:
            _, _, nach_tab = laeufe[0].text.partition("\t")
            laeufe[0].text = f"{links_bis_tab}\t{nach_tab}"
        else:
            laeufe[0].text = links_bis_tab
            for r in laeufe[1:tab_index if tab_index is not None else len(laeufe)]:
                _text_leeren(r)

    if rechts_ab_tab is not None and tab_index is not None:
        vor_tab, _, _ = laeufe[tab_index].text.partition("\t")
        laeufe[tab_index].text = f"{vor_tab}\t{rechts_ab_tab}"
        for r in laeufe[tab_index + 1:]:
            _text_leeren(r)


def _text_leeren(lauf) -> None:
    """Nur den sichtbaren Text eines Laufs leeren.

    Ein Lauf mit Feldzeichen bleibt ganz unberührt: Dort steht die Seitenzahl,
    und ``w:t`` zu leeren würde den Feldcode zwar stehen lassen, das Feld aber
    aus dem Zusammenhang reißen.
    """
    if (lauf._element.findall(qn("w:fldChar"))
            or lauf._element.findall(qn("w:instrText"))):
        return
    for t in lauf._element.findall(qn("w:t")):
        t.text = ""


def setze_titel(doc: Document, titel: str) -> None:
    doc.core_properties.title = titel
    doc.core_properties.author = "Cyberwoks"
    doc.core_properties.category = "Vorlage"
